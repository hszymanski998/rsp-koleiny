# -*- coding: utf-8 -*-
"""
RSP vs LCMS – System Analizy Kolein (DSN 2019 / GDDKiA)
=======================================================

Uruchomienie:   streamlit run app.py

Co potrafi (skrót):
  * wiele plików naraz (RSP, LCMS, GPS) – typ pliku rozpoznawany po zawartości,
  * pikietaż bezwzględny (km+m) dla wszystkich źródeł, rejestr plików z ręczną korektą,
  * RSP: surowe lasery (.rsp / .txt / .xlsx / .csv) ORAZ gotowe koleiny z XML KOLC (.2025),
  * LCMS: profile (CSV / XLSX, układ "okien" ID) + GPS z PIC XML,
  * profil poprzeczny RSP vs LCMS (jak dotychczas) + nawigacja, mapa cieplna profili,
  * przebieg wzdłużny, ocena odcinkowa DSN 2019 (Hm = E(h) + 2 Dh), klasy A–D,
  * zgodność statystyczna: Bland–Altman, RPD, regresja, CCC, ICC, kappa, macierz klas,
  * synchronizacja: offset globalny, offsety sekcji, dryf skali, mapa lokalnych przesunięć,
  * mapa GPS, automatyczne wnioski, eksport (Excel / CSV / raport HTML).

Struktura pliku:
  1. stałe i narzędzia          6. statystyki zgodności
  2. rozpoznawanie plików       7. synchronizacja
  3. parsery                    8. wnioski
  4. geometria (łata DSN)       9. wykresy
  5. zestawy danych            10. eksport
                               11. interfejs Streamlit
"""
from __future__ import annotations

import datetime
import glob
import hashlib
import io
import math
import os
import re
import time
from collections import Counter

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

APP_VERSION = "2.0"

# ==========================================================================
# 1. STAŁE I NARZĘDZIA
# ==========================================================================
# Dokładna geometria laserów RSP [m] (-1.610 m do +1.618 m) – bez zmian
DEFAULT_RSP_X_LIST = [
    -1.610, -1.444, -1.285, -1.145, -1.007, -0.850, -0.700, -0.550, -0.400,
    -0.200, 0.000, 0.200, 0.400, 0.550, 0.700, 0.850, 1.003, 1.140, 1.280,
    1.440, 1.618,
]
DEFAULT_RSP_X_STR = ", ".join(f"{x:.3f}" for x in DEFAULT_RSP_X_LIST)

# Klasy stanu koleiny wg DSN 2019 (progi dolne klas B, C, D)
DSN_THRESHOLDS = (11.0, 21.0, 31.0)
DSN_CLASSES = ("A", "B", "C", "D")
DSN_COLORS = {"A": "#28a745", "B": "#ffc107", "C": "#fd7e14", "D": "#dc3545"}
DSN_EMOJI = {"A": "🟢", "B": "🟡", "C": "🟠", "D": "🔴"}

# Typy plików rozpoznawane przez aplikację
TYPE_RSP_RAW = "RSP_RAW"
TYPE_RSP_KOLC = "RSP_KOLC"
TYPE_RSP_GPS = "RSP_GPS"
TYPE_LCMS_PROFILE = "LCMS_PROFILE"
TYPE_LCMS_GPS = "LCMS_GPS"
TYPE_SKIP = "SKIP"
TYPE_LABELS = {
    TYPE_RSP_RAW: "RSP – surowe lasery",
    TYPE_RSP_KOLC: "RSP – koleiny (KOLC XML / RTG)",
    TYPE_RSP_GPS: "RSP – GPS",
    TYPE_LCMS_PROFILE: "LCMS – profile",
    TYPE_LCMS_GPS: "LCMS – GPS (PIC XML)",
    TYPE_SKIP: "— pomiń —",
}
LABEL_TO_TYPE = {v: k for k, v in TYPE_LABELS.items()}

C_RSP = "royalblue"
C_RSP2 = "darkorange"
C_LCMS = "forestgreen"
C_LCMS2 = "limegreen"
C_DIFF = "#7b2cbf"


def get_dsn_class(h_val):
    """Zwraca klasę stanu koleiny wg DSN 2019 (A/B/C/D)."""
    if h_val is None or (isinstance(h_val, float) and math.isnan(h_val)):
        return "?"
    if h_val < DSN_THRESHOLDS[0]:
        return "A"
    if h_val < DSN_THRESHOLDS[1]:
        return "B"
    if h_val < DSN_THRESHOLDS[2]:
        return "C"
    return "D"


def get_dsn_color(klasa):
    """Kolory dla klas DSN 2019."""
    return DSN_COLORS.get(klasa, "#999999")


def dsn_class_index(values):
    """Wektorowo: 0..3 (A..D), -1 dla NaN."""
    v = np.asarray(values, dtype=float)
    out = np.full(v.shape, -1, dtype=int)
    ok = ~np.isnan(v)
    out[ok] = np.searchsorted(np.array(DSN_THRESHOLDS), v[ok], side="right")
    return out


def km_str(m):
    """Pikietaż w zapisie km+m, np. 114310 -> '114+310'."""
    if m is None or (isinstance(m, float) and math.isnan(m)):
        return "—"
    sign = "-" if m < 0 else ""
    mm = int(round(abs(float(m))))
    return f"{sign}{mm // 1000}+{mm % 1000:03d}"


def to_float(v, default=np.nan):
    """Bezpieczna konwersja komórki (liczba / tekst z przecinkiem) na float."""
    if v is None:
        return default
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().strip("\"'").replace(" ", "").replace(" ", "")
    if s == "":
        return default
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return default


_PL_MONTHS = {"sty": 1, "lut": 2, "mar": 3, "kwi": 4, "maj": 5, "cze": 6,
              "lip": 7, "sie": 8, "wrz": 9, "paź": 10, "paz": 10, "lis": 11, "gru": 12}


def parse_rut_cell(v):
    """Wartość koleiny z nagłówka profilu LCMS.

    Obsługuje liczby, tekst z przecinkiem oraz sytuację, w której Excel
    zamienił zapis typu '1,10' na datę (1 października -> 1.10).
    Wartości <= -90 (np. -99) oznaczają brak danych -> NaN.
    """
    if v is None:
        return np.nan
    if isinstance(v, (datetime.datetime, datetime.date)):
        val = float(f"{v.day}.{v.month:02d}")
    elif isinstance(v, (int, float, np.integer, np.floating)):
        val = float(v)
    else:
        s = str(v).strip().lower().replace(",", ".")
        if s == "":
            return np.nan
        for pl, num in _PL_MONTHS.items():
            if pl in s:
                s = s.replace(pl, f"{num:02d}")
        m = re.findall(r"[-+]?\d*\.\d+|[-+]?\d+", s)
        if not m:
            return np.nan
        val = float(m[0])
    if val <= -90:
        return np.nan
    return val


def robust_step(chain):
    """Typowy krok między kolejnymi próbkami (mediana dodatnich różnic)."""
    c = np.asarray(chain, dtype=float)
    if c.size < 2:
        return 1.0
    d = np.diff(c)
    d = d[d > 0]
    return float(np.median(d)) if d.size else 1.0


def haversine_m(lat1, lon1, lat2, lon2):
    """Odległość w metrach (wektorowo) – sfera o R = 6371 km."""
    R = 6371000.0
    lat1 = np.radians(np.asarray(lat1, dtype=float))
    lat2 = np.radians(np.asarray(lat2, dtype=float))
    dphi = lat2 - lat1
    dl = np.radians(np.asarray(lon2, dtype=float) - np.asarray(lon1, dtype=float))
    a = np.sin(dphi / 2.0) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin(dl / 2.0) ** 2
    return 2.0 * R * np.arctan2(np.sqrt(a), np.sqrt(1.0 - a))


def haversine_distance_m(lat1, lon1, lat2, lon2):
    """Zgodność wsteczna z poprzednią wersją (pojedyncza odległość)."""
    return float(haversine_m(lat1, lon1, lat2, lon2))


def nan_moving_average(y, window):
    """Średnia krocząca odporna na NaN (okno nieparzyste, wycentrowane)."""
    y = np.asarray(y, dtype=float)
    window = int(max(1, window))
    if window <= 1:
        return y.copy()
    ok = ~np.isnan(y)
    v = np.where(ok, y, 0.0)
    k = np.ones(window)
    num = np.convolve(v, k, mode="same")
    den = np.convolve(ok.astype(float), k, mode="same")
    with np.errstate(invalid="ignore", divide="ignore"):
        out = num / den
    out[den == 0] = np.nan
    return out


def decimate_minmax(x, ys, max_points=4000):
    """Zmniejsza liczbę punktów do wyświetlenia zachowując obwiednię (min/max w koszykach).

    x – tablica pikietaży, ys – lista tablic y (te same długości). Zwraca (x2, [y2...]).
    """
    x = np.asarray(x, dtype=float)
    n = len(x)
    if n <= max_points:
        return x, [np.asarray(y, dtype=float) for y in ys]
    nb = max_points // 2
    edges = np.linspace(0, n, nb + 1).astype(int)
    keep = []
    ref = np.asarray(ys[0], dtype=float)
    for a, b in zip(edges[:-1], edges[1:]):
        if b <= a:
            continue
        seg = ref[a:b]
        if np.all(np.isnan(seg)):
            keep.append(a)
            continue
        i0 = a + int(np.nanargmin(seg))
        i1 = a + int(np.nanargmax(seg))
        keep.extend(sorted({i0, i1}))
    keep = np.array(sorted(set(keep)), dtype=int)
    return x[keep], [np.asarray(y, dtype=float)[keep] for y in ys]


def safe_div(a, b, default=np.nan):
    return a / b if b not in (0, 0.0) and not (isinstance(b, float) and math.isnan(b)) else default


def _first_non_empty(*vals):
    for v in vals:
        if v is not None and v != "":
            return v
    return None


# ==========================================================================
# 2. PLIKI: ODCZYT I ROZPOZNAWANIE TYPU PO ZAWARTOŚCI
# ==========================================================================
class FileRef:
    """Plik z uploadu (bajty w pamięci) albo z dysku (czytany strumieniowo).

    `ident` – stabilny identyfikator zawartości (hash / ścieżka+rozmiar+mtime),
    używany jako klucz cache.
    """

    def __init__(self, name, size, ident, data=None, path=None):
        self.name = name
        self.size = size
        self.ident = ident
        self.data = data
        self.path = path

    def read(self):
        if self.data is not None:
            return self.data
        with open(self.path, "rb") as f:
            return f.read()

    def head(self, n=131072):
        if self.data is not None:
            return self.data[:n]
        with open(self.path, "rb") as f:
            return f.read(n)

    def xlsx_source(self):
        return self.path if self.path else io.BytesIO(self.data)

    def lines(self):
        """Iterator po liniach tekstu (UTF-8, BOM ignorowany), bez znaków końca linii."""
        bio = io.BytesIO(self.data) if self.data is not None else open(self.path, "rb")
        wrapper = io.TextIOWrapper(bio, encoding="utf-8-sig", errors="replace", newline=None)
        try:
            for line in wrapper:
                yield line.rstrip("\r\n")
        finally:
            wrapper.close()

    @property
    def ext(self):
        return os.path.splitext(self.name.lower())[1]


def file_ref_from_upload(uf):
    """FileRef z obiektu st.file_uploader (hash liczony raz na plik)."""
    data = uf.getvalue()
    key = f"{getattr(uf, 'file_id', '') or uf.name}|{len(data)}"
    cache = st.session_state.setdefault("_upload_hashes", {})
    if key not in cache:
        cache[key] = hashlib.md5(data).hexdigest()
    return FileRef(uf.name, len(data), "up:" + cache[key], data=data)


def file_ref_from_path(path):
    stt = os.stat(path)
    ident = f"fs:{os.path.abspath(path)}|{stt.st_size}|{stt.st_mtime_ns}"
    return FileRef(os.path.basename(path), stt.st_size, ident, path=path)


def _is_xlsx(ref):
    return ref.ext in (".xlsx", ".xlsm") or ref.head(4)[:2] == b"PK"


def _open_workbook(ref):
    import openpyxl
    return openpyxl.load_workbook(ref.xlsx_source(), read_only=True, data_only=True)


def _peek_xlsx_rows(ref, max_rows=10, max_sheets=3):
    wb = _open_workbook(ref)
    out = []
    try:
        for ws in wb.worksheets[:max_sheets]:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                rows.append(list(r))
                if i + 1 >= max_rows:
                    break
            out.append((ws.title, rows))
    finally:
        wb.close()
    return out


def _is_int_like(v):
    f = to_float(v)
    return (not np.isnan(f)) and abs(f - round(f)) < 1e-9


def _num_count(row):
    return sum(1 for v in row if not np.isnan(to_float(v)))


def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def _guess_sep(line):
    cnt = {";": line.count(";"), "\t": line.count("\t"), ",": line.count(",")}
    return max(cnt, key=cnt.get)


def _classify_rows(rows):
    """Rozpoznanie typu po pierwszych wierszach arkusza / pliku tekstowego."""
    if not rows:
        return TYPE_SKIP
    hdr = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    if "profil" in hdr[:3] and any(re.fullmatch(r"l0?1", h) for h in hdr):
        return TYPE_RSP_RAW
    codes = [to_float(r[0]) for r in rows if r]
    if any(c == 5401 for c in codes):
        return TYPE_RSP_RAW
    if any(c == 5280 for c in codes):
        return TYPE_RSP_GPS
    # LCMS: wiersz nagłówkowy (id; rl; rr; x...) + wiersz z-ek bez numeru
    for k in range(len(rows) - 1):
        a, b = rows[k], rows[k + 1]
        if len(a) > 40 and _is_int_like(a[0]) and _num_count(a) > 40 and _is_blank(b[0]) and _num_count(b) > 30:
            return TYPE_LCMS_PROFILE
    # GPS ogólny: kolumny o wartościach jak szer. / dł. geograficzna Polski
    mat = []
    for r in rows:
        vals = [to_float(v) for v in r]
        if sum(1 for v in vals if not np.isnan(v)) >= 2:
            mat.append(vals)
    if len(mat) >= 2:
        w = max(len(m) for m in mat)
        arr = np.full((len(mat), w), np.nan)
        for i, m in enumerate(mat):
            arr[i, :len(m)] = m
        has_lat = any(np.nanmean(arr[:, j]) > 48 and np.nanmean(arr[:, j]) < 56 for j in range(w) if np.any(~np.isnan(arr[:, j])))
        has_lon = any(np.nanmean(arr[:, j]) > 13 and np.nanmean(arr[:, j]) < 25 for j in range(w) if np.any(~np.isnan(arr[:, j])))
        if has_lat and has_lon and w <= 20:
            return TYPE_RSP_GPS
    return TYPE_SKIP


def classify_file(ref):
    """Zwraca (typ, uwaga). Rozpoznanie po zawartości, nie po nazwie."""
    try:
        if _is_xlsx(ref):
            for _title, rows in _peek_xlsx_rows(ref):
                t = _classify_rows(rows)
                if t != TYPE_SKIP:
                    return t, ""
            return TYPE_SKIP, "nie rozpoznano układu arkusza"
        if ref.ext == ".xls":
            return TYPE_SKIP, "format .xls – zapisz jako .xlsx"
        head = ref.head()
        txt = head.decode("utf-8-sig", errors="ignore")
        low = txt.lower()
        if "<kolc" in low or "<odcinek>" in low:
            return TYPE_RSP_KOLC, ""
        if "average rutting" in low and re.search(r"(?im)^\s*from\s*,\s*to\s*,\s*left", txt):
            return TYPE_RSP_KOLC, "RTG (RspWin)"
        if "<picindex" in low:
            return TYPE_LCMS_GPS, ""
        if "<gpx" in low:
            return TYPE_RSP_GPS, ""
        if re.search(r'(?m)^\s*"?5401"?\s*[,;\t]', txt):
            return TYPE_RSP_RAW, ""
        if re.search(r'(?m)^\s*"?5280"?\s*[,;\t]', txt):
            return TYPE_RSP_GPS, ""
        lines = [ln for ln in txt.splitlines() if ln.strip()][:60]
        if not lines:
            return TYPE_SKIP, "pusty plik"
        sep = _guess_sep(lines[0])
        rows = [[c.strip() for c in ln.split(sep)] for ln in lines]
        if len(rows) >= 2 and len(rows[0]) > 40 and _is_int_like(rows[0][0]) and _is_blank(rows[1][0]):
            return TYPE_LCMS_PROFILE, ""
        t = _classify_rows(rows)
        return t, ("" if t != TYPE_SKIP else "nie rozpoznano układu pliku")
    except Exception as exc:  # noqa: BLE001
        return TYPE_SKIP, f"błąd rozpoznawania: {exc}"


# ==========================================================================
# 3. PARSERY
# ==========================================================================
# ---------- 3.1 RSP: surowe lasery ----------
def _trim_row(row):
    row = list(row)
    while row and _is_blank(row[-1]):
        row.pop()
    return row


def parse_rsp_code_rows(rows, n_lasers=21):
    """Rekordy natywnego formatu RSP: 5201–5221 (geometria), 5401 (lasery), 5280 (GPS)."""
    sensors, lasers, kms, gps = [], [], [], []
    for parts in rows:
        parts = _trim_row(parts)
        if not parts:
            continue
        c = to_float(parts[0])
        if np.isnan(c):
            continue
        code = int(c)
        if 5201 <= code <= 5221:
            x_mm = to_float(parts[-1])
            if not np.isnan(x_mm):
                sensors.append((code, x_mm / 1000.0))
        elif code == 5401:
            vals = [to_float(v) for v in parts[3:3 + n_lasers]]
            if len(vals) == n_lasers and not any(np.isnan(v) for v in vals):
                lasers.append(vals)
                kms.append(to_float(parts[1]))
        elif code == 5280:
            lat, lon = to_float(parts[5]) if len(parts) > 5 else np.nan, to_float(parts[6]) if len(parts) > 6 else np.nan
            if not (np.isnan(lat) or np.isnan(lon)):
                gps.append({
                    "km": to_float(parts[1]),
                    "LAT": lat,
                    "LON": lon,
                    "ELE": to_float(parts[7]) if len(parts) > 7 else np.nan,
                })
    return {
        "lasers": np.array(lasers, dtype=float) if lasers else None,
        "km": np.array(kms, dtype=float),
        "sensors": sensors,
        "gps": gps,
    }


def _gps_frame_from_records(gps_records):
    if not gps_records:
        return None
    g = pd.DataFrame(gps_records)
    km = g["km"].values.astype(float)
    if np.all(np.isnan(km)):
        chain = np.full(len(g), np.nan)
    else:
        chain = km * 1000.0
    out = pd.DataFrame({"chain_m": chain, "LAT": g["LAT"].values, "LON": g["LON"].values, "ELE": g["ELE"].values})
    return out


def _finish_gps(df, notes, label="GPS"):
    """Uzupełnia pikietaż GPS (gdy brak: skumulowana odległość) i usuwa zera."""
    if df is None or df.empty:
        return None
    df = df.dropna(subset=["LAT", "LON"]).copy()
    df = df[(df["LAT"].abs() > 1e-6) & (df["LON"].abs() > 1e-6)]
    if df.empty:
        return None
    kind = "km"
    if df["chain_m"].isna().all():
        d = haversine_m(df["LAT"].values[:-1], df["LON"].values[:-1], df["LAT"].values[1:], df["LON"].values[1:])
        df["chain_m"] = np.r_[0.0, np.cumsum(d)]
        kind = "dist"
        notes.append(f"{label}: brak pikietażu w pliku – użyto skumulowanej odległości po śladzie")
    df = df.reset_index(drop=True)
    df.attrs["chain_kind"] = kind
    return df


def parse_rsp_raw(ref):
    """Surowe dane RSP: .rsp/.txt (rekordy 5401), CSV/XLSX z kolumnami Profil, L1..L21."""
    notes = []
    lasers = prof = chain = x_pos = None
    gps_df = None
    n_lasers = 21
    codes = None

    if _is_xlsx(ref):
        wb = _open_workbook(ref)
        table = None
        code_rows = []
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                if not rows:
                    continue
                hdr = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
                if table is None and "profil" in hdr[:3] and any(re.fullmatch(r"l0?1", h) for h in hdr):
                    table = (hdr, rows[1:])
                else:
                    for r in rows:
                        if r and not _is_blank(r[0]):
                            c = to_float(r[0])
                            if not np.isnan(c) and 5200 <= c <= 5499:
                                code_rows.append(r)
        finally:
            wb.close()
        if table is not None:
            hdr, body = table
            ip = hdr.index("profil")
            lcols = [i for i, h in enumerate(hdr) if re.fullmatch(r"l\d+", h)]
            n_lasers = len(lcols)
            data = []
            for r in body:
                if r is None or len(r) <= max(lcols + [ip]) or _is_blank(r[ip]):
                    continue
                data.append([to_float(r[ip])] + [to_float(r[i]) for i in lcols])
            arr = np.array(data, dtype=float)
            arr = arr[~np.isnan(arr[:, 0])]
            prof = arr[:, 0]
            lasers = arr[:, 1:]
        if code_rows:
            codes = parse_rsp_code_rows(code_rows, n_lasers=n_lasers)
    else:
        sep = ","
        for ln in ref.lines():
            if re.match(r'\s*"?5401"?\s*[,;\t]', ln):
                sep = _guess_sep(ln)
                break
        # natywny format RSP?
        def _rows():
            for ln in ref.lines():
                if not ln.strip():
                    continue
                yield [p.strip().strip('"') for p in ln.split(sep)]
        codes = parse_rsp_code_rows(_rows(), n_lasers=n_lasers)
        if codes["lasers"] is None:
            codes = None
            raw = ref.read()
            df = pd.read_csv(io.BytesIO(raw), sep=None, engine="python", encoding="utf-8-sig")
            df.columns = [str(c).strip() for c in df.columns]
            low = [c.lower() for c in df.columns]
            ip = low.index("profil") if "profil" in low else 0
            lcols = [i for i, c in enumerate(low) if i != ip and re.fullmatch(r"l\d+", c)] or [i for i in range(len(low)) if i != ip]
            prof = np.array([to_float(v) for v in df.iloc[:, ip]], dtype=float)
            lasers = np.array([[to_float(v) for v in df.iloc[:, i]] for i in lcols], dtype=float).T
            n_lasers = lasers.shape[1]
            keep = ~np.isnan(prof)
            prof, lasers = prof[keep], lasers[keep]

    if codes is not None:
        if codes["sensors"]:
            xs = [s[1] for s in sorted(codes["sensors"], key=lambda s: s[0])]
            if lasers is not None and len(xs) == lasers.shape[1]:
                x_pos = xs
            elif lasers is None and len(xs) == n_lasers:
                x_pos = xs
        if lasers is None and codes["lasers"] is not None:
            lasers = codes["lasers"]
            prof = np.arange(1, len(lasers) + 1, dtype=float)
            chain = codes["km"] * 1000.0 if not np.any(np.isnan(codes["km"])) else None
        elif lasers is not None and codes["lasers"] is not None and len(codes["lasers"]) == len(lasers):
            same = np.allclose(codes["lasers"][0], lasers[0], atol=0.051)
            if same and not np.any(np.isnan(codes["km"])):
                chain = codes["km"] * 1000.0
                notes.append("pikietaż RSP odczytany z rekordów 5401 (km)")
        if codes["gps"]:
            gps_df = _finish_gps(_gps_frame_from_records(codes["gps"]), notes, "GPS wbudowany w RSP")
            if gps_df is not None:
                notes.append(f"GPS wbudowany w plik RSP: {len(gps_df)} pkt")

    if lasers is None or len(lasers) == 0:
        raise ValueError("nie znaleziono odczytów laserów RSP")
    if chain is not None:
        # RSP jedzie często "rosnąco" – upewnij się, że pikietaż jest rosnący
        if chain[-1] < chain[0]:
            notes.append("pikietaż RSP malejący – zachowano kolejność profili")
    return {
        "type": TYPE_RSP_RAW, "name": ref.name, "ident": ref.ident,
        "prof": prof, "lasers": lasers, "chain": chain, "x_pos": x_pos,
        "gps": gps_df, "notes": notes, "n": int(len(lasers)),
    }


# ---------- 3.2 RSP: GPS ----------
def parse_rsp_gps(ref):
    """GPS profilografu: GPX, rekordy 5280 (.rsp/.txt/.xlsx) lub tabela z kolumnami lat/lon."""
    notes = []
    head = ref.head(4096).decode("utf-8-sig", errors="ignore").lower()
    df = None
    if "<gpx" in head:
        import xml.etree.ElementTree as ET
        root = ET.fromstring(ref.read())
        for el in root.iter():
            if "}" in el.tag:
                el.tag = el.tag.split("}", 1)[1]
        recs = []
        for tag in ("trkpt", "rtept", "wpt"):
            for pt in root.iter(tag):
                try:
                    ele = pt.findtext("ele")
                    recs.append({"km": np.nan, "LAT": float(pt.attrib["lat"]), "LON": float(pt.attrib["lon"]),
                                 "ELE": float(ele) if ele else np.nan})
                except (KeyError, ValueError, TypeError):
                    continue
            if recs:
                break
        df = _finish_gps(_gps_frame_from_records(recs), notes, "GPX")
    else:
        if _is_xlsx(ref):
            wb = _open_workbook(ref)
            try:
                rows = [list(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
            finally:
                wb.close()
        else:
            ls = [ln for ln in ref.lines() if ln.strip()]
            sep = _guess_sep(ls[0]) if ls else ","
            rows = [[p.strip().strip('"') for p in ln.split(sep)] for ln in ls]
        codes = parse_rsp_code_rows([r for r in rows if r], n_lasers=21)
        if codes["gps"]:
            df = _finish_gps(_gps_frame_from_records(codes["gps"]), notes, "GPS RSP")
        else:
            mat = [[to_float(v) for v in _trim_row(r)] for r in rows if r]
            mat = [m for m in mat if sum(1 for v in m if not np.isnan(v)) >= 2]
            if mat:
                w = max(len(m) for m in mat)
                arr = np.full((len(mat), w), np.nan)
                for i, m in enumerate(mat):
                    arr[i, :len(m)] = m
                lat_c = lon_c = None
                for j in range(w):
                    col = arr[:, j]
                    if np.all(np.isnan(col)):
                        continue
                    mu = np.nanmean(col)
                    if 48.0 <= mu <= 56.0 and lat_c is None:
                        lat_c = j
                    elif 13.0 <= mu <= 25.0 and lon_c is None:
                        lon_c = j
                if lat_c is None or lon_c is None:
                    raise ValueError("w pliku GPS nie znaleziono kolumn szerokość/długość")
                km = arr[:, 1] if w > 1 and 1 not in (lat_c, lon_c) else np.full(len(arr), np.nan)
                if np.all(np.isnan(km)) or not (np.nanmax(km) < 1000):
                    km = np.full(len(arr), np.nan)
                recs = [{"km": km[i], "LAT": arr[i, lat_c], "LON": arr[i, lon_c], "ELE": np.nan} for i in range(len(arr))]
                df = _finish_gps(_gps_frame_from_records(recs), notes, "GPS RSP")
    if df is None or df.empty:
        raise ValueError("nie znaleziono punktów GPS")
    return {"type": TYPE_RSP_GPS, "name": ref.name, "ident": ref.ident, "gps": df,
            "notes": notes, "n": int(len(df)), "chain_kind": df.attrs.get("chain_kind", "km")}


# ---------- 3.3 RSP: koleiny z XML KOLC (.2025 / .xml) ----------
_KOLC_TAGS = ["KM_P", "KM_K", "L_KOL", "PEL_KOL", "P_KOL", "MAX_L_KOL", "MAX_PEL_KOL", "MAX_P_KOL",
              "LAT", "LON", "IL_SAT", "POMIAR", "FOTO", "DATA_POM"]


def _kolc_comments(txt_head):
    out = []
    pat = re.compile(
        r"^\s*(?P<name>.*?)\s+km:\s*(?P<k1>[\d.]+)\s+(?:[\d.]+\s+)?(?:-\s*)?(?P<k2>[\d.]+)\s+"
        r"IL\.\s*EL\.\s*WPAS\.:\s*(?P<n>\d+)\s+WPASOWANIE:\s*(?P<w>-?[\d.]+)", re.I)
    for c in re.findall(r"<!--(.*?)-->", txt_head, re.S):
        m = pat.match(c.strip())
        if m:
            out.append({
                "plik_rsp": m.group("name"),
                "km_od": float(m.group("k1")), "km_do": float(m.group("k2")),
                "el_wpasowania": int(m.group("n")), "wpasowanie": float(m.group("w")),
            })
    return out


def parse_kolc(ref):
    """XML KOLC (GDDKiA): ODCINEK co 1 m z L_KOL / P_KOL / PEL_KOL, MAX_*, GPS, data."""
    txt = ref.read().decode("utf-8-sig", errors="replace")
    notes = []
    n_blocks = txt.count("<ODCINEK>")
    if n_blocks == 0:
        raise ValueError("brak elementów <ODCINEK>")
    cols = {}
    ok = True
    for tag in _KOLC_TAGS:
        vals = re.findall(rf"<{tag}>([^<]*)</{tag}>", txt)
        vals += [""] * 0
        cols[tag] = vals
        if len(vals) != n_blocks:
            ok = False
    if not ok:  # rzadkie: brakujące znaczniki w części rekordów – wolniejsza ścieżka blokowa
        blocks = re.findall(r"<ODCINEK>(.*?)</ODCINEK>", txt, re.S)
        cols = {t: [] for t in _KOLC_TAGS}
        for b in blocks:
            d = dict(re.findall(r"<(\w+)>([^<]*)</\1>", b))
            for t in _KOLC_TAGS:
                cols[t].append(d.get(t, ""))

    def farr(tag):
        return np.array([to_float(v) for v in cols[tag]], dtype=float)

    km_p = farr("KM_P")
    km_k = farr("KM_K")
    chain = np.round(km_p * 1000.0, 3)
    lat, lon = farr("LAT"), farr("LON")
    gps_ok = (~np.isnan(lat)) & (~np.isnan(lon)) & (np.abs(lat) > 1.0) & (np.abs(lon) > 1.0)
    gps_df = None
    if gps_ok.any():
        gps_df = pd.DataFrame({"chain_m": chain[gps_ok], "LAT": lat[gps_ok], "LON": lon[gps_ok], "ELE": np.nan})
        gps_df.attrs["chain_kind"] = "km"
    head_txt = txt[:60000]
    attrs = {}
    m = re.search(r"<KOLC\b([^>]*)>", txt)
    if m:
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
    comments = _kolc_comments(head_txt)
    pomiar = np.array(cols["POMIAR"], dtype=object)
    data = {
        "L": farr("L_KOL"), "P": farr("P_KOL"), "PEL": farr("PEL_KOL"),
        "ML": farr("MAX_L_KOL"), "MP": farr("MAX_P_KOL"), "MPEL": farr("MAX_PEL_KOL"),
    }
    n_gps = int(gps_ok.sum())
    notes.append(f"{n_blocks} odcinków, GPS poprawny w {n_gps} ({100.0 * n_gps / n_blocks:.1f}%)")
    if comments:
        notes.append(f"pliki źródłowe RSP w nagłówku: {len(comments)}")
    return {
        "type": TYPE_RSP_KOLC, "name": ref.name, "ident": ref.ident,
        "chain": chain, "chain_end": np.round(km_k * 1000.0, 3), "data": data,
        "il_sat": farr("IL_SAT"), "pomiar": pomiar,
        "dates": sorted(set(v for v in cols["DATA_POM"] if v)),
        "gps": gps_df, "attrs": attrs, "comments": comments, "notes": notes, "n": int(n_blocks),
    }


def parse_rtg(ref):
    """RTG z RspWin: koleiny co 1 m (From, To [km], Left, Full, Right, LMax, FMax, RMax) – jak KOLC."""
    txt = ref.read().decode("utf-8-sig", errors="replace")
    lines = txt.splitlines()
    head = lines[:14]
    pat = re.compile(r"^\s*(-?\d+\.\d+)\s*,\s*(-?\d+\.\d+)\s*,\s*" + r"\s*,\s*".join([r"(-?[\d.]+)"] * 6) + r"\s*$")
    rows = []
    n_sum = 0
    for ln in lines[12:]:
        m = pat.match(ln)
        if not m:
            continue
        v = [float(x) for x in m.groups()]
        if abs(v[1] - v[0]) > 0.0015:  # wiersz podsumowania całego pliku
            n_sum += 1
            continue
        rows.append(v)
    if not rows:
        raise ValueError("brak wierszy z koleinami w pliku RTG")
    a = np.array(rows)
    a[:, 2:][a[:, 2:] <= -90] = np.nan
    chain = np.round(a[:, 0] * 1000.0, 3)
    order = np.argsort(chain, kind="stable")
    a, chain = a[order], chain[order]
    notes = [f"RTG (RspWin): {len(a)} odcinków co 1 m, {chain[0] / 1000:.3f}–{chain[-1] / 1000 + 0.001:.3f} km"]
    comments = []
    for h in head[:3]:
        mm = re.search(r"File:\s*(.+?)\s*$", h)
        if mm:
            notes.append("plik źródłowy RSP: " + mm.group(1).strip())
    data = {"L": a[:, 2], "PEL": a[:, 3], "P": a[:, 4], "ML": a[:, 5], "MPEL": a[:, 6], "MP": a[:, 7]}
    return {
        "type": TYPE_RSP_KOLC, "name": ref.name, "ident": ref.ident,
        "chain": chain, "chain_end": np.round(chain + 1.0, 3), "data": data,
        "il_sat": np.full(len(a), np.nan), "pomiar": np.array([""] * len(a), dtype=object),
        "dates": [], "gps": None, "attrs": {"FORMAT": "RTG"}, "comments": comments, "notes": notes, "n": int(len(a)),
    }


# ---------- 3.4 LCMS: profile poprzeczne ----------
def _num_array(strings, dtype=np.float32):
    try:
        return np.array(strings, dtype=dtype)
    except (ValueError, TypeError):
        return np.array([to_float(s) for s in strings if not _is_blank(s)], dtype=dtype)


def _lcms_records_from_text(ref):
    recs = []
    sep = None
    pending = None
    for line in ref.lines():
        if not line.strip():
            continue
        if sep is None:
            sep = _guess_sep(line)
        dec = sep != ","
        parts = line.split(sep, 3)
        if len(parts) < 4:
            continue
        first = parts[0].strip()
        rest = parts[3].rstrip(sep)
        if dec:
            rest = rest.replace(",", ".")
        if first.isdigit():  # nagłówek połówki: id; rl; rr; x...
            pending = (int(first), parse_rut_cell(parts[1]), parse_rut_cell(parts[2]), _num_array(rest.split(sep)))
        elif first == "" and pending is not None:  # wiersz z
            z = _num_array(rest.split(sep))
            x = pending[3]
            n = min(len(x), len(z))
            if n >= 3:
                recs.append((pending[0], pending[1], pending[2], x[:n], z[:n]))
            pending = None
    return recs


def _lcms_records_from_xlsx(ref):
    wb = _open_workbook(ref)
    recs = []
    pending = None
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            if row is None or len(row) < 6:
                continue
            c0 = row[0]
            if _is_blank(c0):
                if pending is not None:
                    z = _num_array([to_float(v) for v in row[3:] if not _is_blank(v)])
                    x = pending[3]
                    n = min(len(x), len(z))
                    if n >= 3:
                        recs.append((pending[0], pending[1], pending[2], x[:n], z[:n]))
                    pending = None
                continue
            if _is_int_like(c0) and sum(1 for v in row[3:12] if not np.isnan(to_float(v))) >= 5:
                x = _num_array([to_float(v) for v in row[3:] if not _is_blank(v)])
                pending = (int(round(to_float(c0))), parse_rut_cell(row[1]), parse_rut_cell(row[2]), x)
    finally:
        wb.close()
    return recs


def lcms_assign_stations(ids):
    """Numer stacji (profilu) i strona (0 = lewa, 1 = prawa połówka) z numerów ID.

    Plik LCMS składa się z "okien": ID narasta o 1 na połówkę profilu, a po ok. 10
    profilach numeracja wraca do wartości o `krok` większej od początku poprzedniego
    okna (5,6…24 | 15,16…34 | 25…). Stację wyznacza się z pozycji okna, więc luki
    (brakujące okna / połówki) NIE przesuwają kolejnych profili.
    """
    ids = np.asarray(ids, dtype=np.int64)
    n = len(ids)
    # 1) granice okien: powrót numeracji do wartości <= poprzedniej
    new_win = np.zeros(n, dtype=bool)
    new_win[0] = True
    new_win[1:] = ids[1:] <= ids[:-1]
    f0 = int(ids[0])
    step = 10
    first_ids = ids[new_win]
    if len(first_ids) >= 3:
        d = np.diff(first_ids)
        d = d[d > 0]
        if d.size:
            step = int(Counter(d.tolist()).most_common(1)[0][0])
    # 2) skok numeracji o więcej niż `krok`, trafiający w początek siatki okien = brakujące okna
    jump = np.zeros(n, dtype=bool)
    jump[1:] = (ids[1:] - ids[:-1] > step) & (((ids[1:] - f0) % step) == 0)
    new_win |= jump
    win_idx = np.cumsum(new_win) - 1
    first_ids = ids[new_win]
    base_w = f0 + step * np.floor((first_ids - f0) / step).astype(np.int64)
    base = base_w[win_idx]
    station = (base - f0) + (ids - base) // 2
    side = (ids - base) % 2
    return station, side, int(len(first_ids)), int(step)


def parse_lcms_profiles(ref):
    """LCMS: pliki 'Profile_DSN_1m' (CSV / XLSX), pary wierszy [id; rl; rr; x…] + [;;; z…]."""
    notes = []
    recs = _lcms_records_from_xlsx(ref) if _is_xlsx(ref) else _lcms_records_from_text(ref)
    if not recs:
        raise ValueError("nie znaleziono profili LCMS (oczekiwane pary wierszy: nagłówek + z)")
    ids = np.array([r[0] for r in recs], dtype=np.int64)
    station, side, n_win, step = lcms_assign_stations(ids)
    # grupowanie połówek w profile
    by_st = {}
    dup = 0
    for i, (s, sd) in enumerate(zip(station.tolist(), side.tolist())):
        slot = by_st.setdefault(s, [None, None])
        if slot[sd] is None:
            slot[sd] = i
        else:
            dup += 1
    st_sorted = sorted(by_st)
    # jednostki: mm -> m
    probe = [np.max(np.abs(recs[i][3])) for i in range(0, len(recs), max(1, len(recs) // 200))]
    to_m = 0.001 if (np.median(probe) > 20) else 1.0
    xs, zs, lens = [], [], []
    rut_l, rut_r, complete = [], [], []
    n_partial = 0
    for s in st_sorted:
        li, ri = by_st[s]
        parts = [recs[j] for j in (li, ri) if j is not None]
        x = np.concatenate([p[3] for p in parts]) * np.float32(to_m)
        z = np.concatenate([p[4] for p in parts])
        if len(parts) == 2 and np.any(np.diff(x) < 0):
            o = np.argsort(x, kind="stable")
            x, z = x[o], z[o]
        xs.append(x.astype(np.float32))
        zs.append(z.astype(np.float32))
        lens.append(len(x))
        head = parts[0]
        rut_l.append(head[1])
        rut_r.append(head[2])
        full = (li is not None and ri is not None)
        complete.append(full)
        if not full:
            n_partial += 1
    ptr = np.zeros(len(lens) + 1, dtype=np.int64)
    ptr[1:] = np.cumsum(lens)
    station_arr = np.array(st_sorted, dtype=np.int64)
    rut_l = np.array(rut_l, dtype=float)
    rut_r = np.array(rut_r, dtype=float)
    # uwagi o jakości
    n_nan = int(np.sum(np.isnan(rut_l) | np.isnan(rut_r)))
    if n_nan:
        notes.append(f"{n_nan} profili z brakiem wartości koleiny (np. -99) w pliku")
    if n_partial:
        notes.append(f"{n_partial} profili z tylko jedną połówką (lewą lub prawą)")
    if dup:
        notes.append(f"{dup} zdublowanych połówek pominięto")
    gaps = np.where(np.diff(station_arr) > 1)[0]
    if gaps.size:
        missing = int(np.sum(np.diff(station_arr)[gaps] - 1))
        txt_g = ", ".join(f"{station_arr[g] + 1}–{station_arr[g + 1] - 1}" for g in gaps[:4])
        notes.append(f"luki w numeracji: {missing} brakujących profili ({txt_g}{' …' if gaps.size > 4 else ''})")
    if to_m != 1.0:
        notes.append("pozycje poprzeczne przeliczone z mm na m")
    notes.append(f"okien ID: {n_win} (krok {step})")
    return {
        "type": TYPE_LCMS_PROFILE, "name": ref.name, "ident": ref.ident,
        "station": station_arr, "rut_l": rut_l, "rut_r": rut_r,
        "ptr": ptr, "x": np.concatenate(xs) if xs else np.zeros(0, np.float32),
        "z": np.concatenate(zs) if zs else np.zeros(0, np.float32),
        "complete": np.array(complete, dtype=bool), "notes": notes, "n": int(len(station_arr)),
    }


def guess_start_from_filename(name):
    """'17-114310-140000_Profile…' -> 114310 (początek odcinka w metrach)."""
    m = re.search(r"(?<!\d)(\d{1,3})[-_](\d{4,7})[-_](\d{4,7})(?!\d)", name)
    if m:
        a, b = int(m.group(2)), int(m.group(3))
        if b > a:
            return float(a)
    return None


# ---------- 3.5 LCMS: GPS z PIC XML ----------
def parse_pic_xml(ref):
    """PIC XML (indeks zdjęć LCMS): Station [m], LAT, LON, kamera (Buchst), VNK/NNK."""
    txt = ref.read().decode("utf-8-sig", errors="replace")
    blocks = re.findall(r"<PicIndex>(.*?)</PicIndex>", txt, re.S)
    if not blocks:
        raise ValueError("brak rekordów <PicIndex>")
    rows = []
    for b in blocks:
        d = dict(re.findall(r"<(\w+)>([^<]*)</\1>", b))
        if "Station" not in d or "LAT" not in d or "LON" not in d:
            continue
        rows.append({
            "Buchst": d.get("Buchst", ""), "Station": to_float(d.get("Station")),
            "LAT": to_float(d.get("LAT")), "LON": to_float(d.get("LON")),
            "ALT": to_float(d.get("ALT")), "Heading": to_float(d.get("Heading")),
            "Filename": d.get("Filename", ""),
            "VNK": d.get("VNK", ""), "NNK": d.get("NNK", ""), "IDDROGI": d.get("IDDROGI", ""),
        })
    df = pd.DataFrame(rows)
    df = df.dropna(subset=["Station", "LAT", "LON"])
    df = df[(df["LAT"].abs() > 1e-6) & (df["LON"].abs() > 1e-6)].reset_index(drop=True)
    if df.empty:
        raise ValueError("brak poprawnych współrzędnych w PIC XML")
    notes = []
    vnk = to_float(df["VNK"].iloc[0])
    nnk = to_float(df["NNK"].iloc[0])
    smax = float(df["Station"].max())
    start_abs = 0.0
    if not np.isnan(vnk) and not np.isnan(nnk) and nnk > vnk and 0.5 * smax <= (nnk - vnk) <= 2.0 * max(smax, 1.0):
        start_abs = float(vnk)
        notes.append(f"VNK/NNK = {int(vnk)}/{int(nnk)} traktowane jako pikietaż początku/końca [m]")
    cam = df["Buchst"].value_counts()
    notes.append("kamery: " + ", ".join(f"{k}={v}" for k, v in cam.items()))
    return {
        "type": TYPE_LCMS_GPS, "name": ref.name, "ident": ref.ident, "df": df,
        "start_abs": start_abs, "notes": notes, "n": int(len(df)),
        "iddrogi": str(df["IDDROGI"].iloc[0]) if "IDDROGI" in df and len(df) else "",
    }


# ---------- 3.6 dyspozytor ----------
def parse_source(ref, typ):
    """Wczytuje plik wg typu (bez cache – cache opakowuje warstwa UI)."""
    if typ == TYPE_RSP_RAW:
        return parse_rsp_raw(ref)
    if typ == TYPE_RSP_KOLC:
        head = ref.head(4096).decode("utf-8-sig", errors="ignore").lower()
        if "<kolc" not in head and "<odcinek" not in head and "average rutting" in head:
            return parse_rtg(ref)
        return parse_kolc(ref)
    if typ == TYPE_RSP_GPS:
        return parse_rsp_gps(ref)
    if typ == TYPE_LCMS_PROFILE:
        return parse_lcms_profiles(ref)
    if typ == TYPE_LCMS_GPS:
        return parse_pic_xml(ref)
    raise ValueError(f"nieobsługiwany typ: {typ}")


# ==========================================================================
# 4. GEOMETRIA: WIRTUALNA ŁATA DSN 2019 (górna otoczka wypukła)
# ==========================================================================
def smooth_profile(z_vals, window_size=31):
    """Filtruje makrostrukturę nawierzchni (średnia krocząca wycentrowana, jak dotychczas)."""
    z = np.asarray(z_vals, dtype=float)
    n = len(z)
    w = int(window_size)
    if w % 2 == 0:
        w += 1
    if n < w or w <= 1:
        return z
    h = w // 2
    cs = np.concatenate([[0.0], np.cumsum(z)])
    idx = np.arange(n)
    lo = np.maximum(idx - h, 0)
    hi = np.minimum(idx + h + 1, n)
    return (cs[hi] - cs[lo]) / (hi - lo)


def _upper_hull_idx(xs, ys):
    """Indeksy wierzchołków górnej otoczki wypukłej (xs rosnące)."""
    h = []
    for i in range(len(xs)):
        xi, yi = xs[i], ys[i]
        while len(h) >= 2:
            a, o = h[-1], h[-2]
            if (xs[a] - xs[o]) * (yi - ys[o]) - (ys[a] - ys[o]) * (xi - xs[o]) >= 0:
                h.pop()
            else:
                break
        h.append(i)
    return h


def get_upper_hull_pts(x_sec, y_sec):
    """Górna otoczka wypukła (Wirtualna Łata) – lista punktów (x, y)."""
    x = np.asarray(x_sec, dtype=float)
    y = np.asarray(y_sec, dtype=float)
    if len(x) == 0:
        return []
    if np.any(np.diff(x) < 0):
        o = np.argsort(x, kind="stable")
        x, y = x[o], y[o]
    idx = _upper_hull_idx(x.tolist(), y.tolist())
    return [(float(x[i]), float(y[i])) for i in idx]


def get_upper_hull(x_sec, y_sec):
    """Wartości prostej łaty DSN 2019 interpolowane do punktów pomiarowych."""
    x = np.asarray(x_sec, dtype=float)
    y = np.asarray(y_sec, dtype=float)
    if len(x) == 0:
        return y.copy()
    if np.any(np.diff(x) < 0):
        o = np.argsort(x, kind="stable")
        inv = np.empty_like(o)
        inv[o] = np.arange(len(o))
        xs, ys = x[o], y[o]
        idx = _upper_hull_idx(xs.tolist(), ys.tolist())
        return np.interp(xs, xs[idx], ys[idx])[inv]
    idx = _upper_hull_idx(x.tolist(), y.tolist())
    return np.interp(x, x[idx], y[idx])


def rut_pair_from_profile(x, z, window=31):
    """Koleina lewa / prawa z jednego profilu LCMS (jak w poprzedniej wersji).

    Zwraca (rut_l, rut_r). NaN gdy profil ma < 10 punktów.
    """
    x = np.asarray(x, dtype=float)
    z = np.asarray(z, dtype=float)
    if len(x) < 10:
        return np.nan, np.nan
    zs = smooth_profile(z, window)
    mid = (x >= -0.5) & (x <= 0.5)
    if np.any(mid):
        ic = int(np.where(mid)[0][np.argmax(zs[mid])])
    else:
        ic = int(np.argmin(np.abs(x)))
    xl, zl = x[: ic + 1], zs[: ic + 1]
    xr, zr = x[ic:], zs[ic:]
    rl = float(np.max(get_upper_hull(xl, zl) - zl)) if len(zl) else np.nan
    rr = float(np.max(get_upper_hull(xr, zr) - zr)) if len(zr) else np.nan
    return rl, rr


def compute_lcms_calc_ruts(lcms, window=31, progress=None):
    """Koleiny wyliczone z profili LCMS metodą Wirtualnej Łaty (cały zestaw)."""
    n = len(lcms["chain"])
    out_l = np.full(n, np.nan)
    out_r = np.full(n, np.nan)
    ptr, X, Z = lcms["ptr"], lcms["x"], lcms["z"]
    complete = lcms.get("complete")
    for i in range(n):
        if complete is not None and not complete[i]:
            continue
        a, b = ptr[i], ptr[i + 1]
        out_l[i], out_r[i] = rut_pair_from_profile(X[a:b], Z[a:b], window)
        if progress is not None and (i % 1000 == 0):
            progress(i / max(1, n))
    if progress is not None:
        progress(1.0)
    return out_l, out_r


def compute_rsp_ruts(lasers, x_vals):
    """Koleiny RSP (lewa / prawa) z odczytów 21 laserów – Wirtualna Łata DSN 2019."""
    data = np.asarray(lasers, dtype=float)
    ile = data.shape[1]
    i_min_mid = int(ile * 0.3)
    i_max_mid = int(ile * 0.7)
    x_arr = np.asarray(x_vals, dtype=float)
    if len(x_arr) != ile:
        x_arr = np.array([round((i - (ile - 1) / 2) * 0.165, 3) for i in range(ile)])
    y_mat = -data
    n = len(y_mat)
    rl = np.full(n, np.nan)
    rr = np.full(n, np.nan)
    for k in range(n):
        row = y_mat[k]
        if np.any(np.isnan(row)):
            continue
        ic = i_min_mid + int(np.argmax(row[i_min_mid: i_max_mid + 1]))
        rl[k] = float(np.max(get_upper_hull(x_arr[: ic + 1], row[: ic + 1]) - row[: ic + 1]))
        rr[k] = float(np.max(get_upper_hull(x_arr[ic:], row[ic:]) - row[ic:]))
    return rl, rr


def rsp_cross_section(odczyty_raw, x_vals, zero_l11=True):
    """Przekrój RSP: profil, łaty lewa/prawa, głębokości (wszystko jak dotychczas).

    Obliczenia zawsze w domenie wysokości (h = -odczyt); tryb "surowy" tylko zmienia
    sposób wyświetlania (odczyt czujnika, oś odwrócona).
    """
    odczyty_raw = np.asarray(odczyty_raw, dtype=float)
    ile = len(odczyty_raw)
    x_vals = np.asarray(x_vals, dtype=float)
    h = -odczyty_raw
    if zero_l11:
        ref = odczyty_raw[ile // 2]
        y_vals = ref - odczyty_raw
        y_label = "Wysokość względna do L11 [mm]"
    else:
        y_vals = odczyty_raw
        y_label = "Odczyt surowy czujnika [mm]"
    sign = 1.0 if zero_l11 else -1.0
    i_min, i_max = int(ile * 0.3), int(ile * 0.7)
    ic = i_min + int(np.argmax(h[i_min: i_max + 1]))
    xl, hl = x_vals[: ic + 1], h[: ic + 1]
    xr, hr = x_vals[ic:], h[ic:]
    lata_l_h = get_upper_hull(xl, hl)
    lata_r_h = get_upper_hull(xr, hr)
    diff_l = lata_l_h - hl
    diff_r = lata_r_h - hr
    idx_l = int(np.argmax(diff_l))
    idx_r_rel = int(np.argmax(diff_r))
    idx_r = ic + idx_r_rel
    # łaty w układzie wyświetlania
    if zero_l11:
        off = odczyty_raw[ile // 2]
        lata_l = off + lata_l_h
        lata_r = off + lata_r_h
        # y = ref - raw = ref + h  -> łata (h) + ref
    else:
        lata_l = -lata_l_h
        lata_r = -lata_r_h
    return {
        "x": x_vals, "y": y_vals, "y_label": y_label, "idx_crown": ic,
        "x_left": xl, "x_right": xr, "lata_l": lata_l, "lata_r": lata_r,
        "rut_l": float(diff_l[idx_l]), "rut_r": float(diff_r[idx_r_rel]),
        "idx_l": idx_l, "idx_r": idx_r, "idx_r_rel": idx_r_rel,
        "reverse_y": (not zero_l11),
        "y_l_prof": float(y_vals[idx_l]), "y_l_lata": float(lata_l[idx_l]),
        "y_r_prof": float(y_vals[idx_r]), "y_r_lata": float(lata_r[idx_r_rel]),
    }


def lcms_cross_section(x_orig, z_raw, rsp_x_min, rsp_x_max, window=31,
                       skaluj_szerokosc=False, zaczep_prawa=True, manual_shift=0.0,
                       wyrownaj_spadek=True):
    """Przekrój LCMS w układzie RSP (zaczepienie krawędzi, wyzerowanie spadku) – jak dotychczas."""
    x_orig = np.asarray(x_orig, dtype=float)
    z_raw = np.asarray(z_raw, dtype=float)
    z_smooth = smooth_profile(z_raw, window)
    # wartości z łaty liczone na pozycjach oryginalnych
    rl_calc, rr_calc = rut_pair_from_profile(x_orig, z_raw, window)

    x = x_orig.copy()
    if skaluj_szerokosc and len(x) > 1 and (x.max() - x.min()) > 0:
        x = rsp_x_min + (x - x.min()) * ((rsp_x_max - rsp_x_min) / (x.max() - x.min()))
    elif zaczep_prawa and len(x) > 0:
        x = x + (rsp_x_max - np.max(x))
    x = x + manual_shift

    mid = (x >= -0.5) & (x <= 0.5)
    if np.any(mid):
        ic = int(np.where(mid)[0][np.argmax(z_smooth[mid])])
    else:
        ic = int(np.argmin(np.abs(x)))
    x_l, z_l_raw = x[: ic + 1], z_smooth[: ic + 1]
    x_r, z_r_raw = x[ic:], z_smooth[ic:]
    lata_raw_l = get_upper_hull(x_l, z_l_raw)
    lata_raw_r = get_upper_hull(x_r, z_r_raw)

    if wyrownaj_spadek and len(x) > 10:
        p_trend = np.polyfit(x, z_smooth, 1)
        trend = p_trend[0] * x
        z_lev = z_smooth - trend
        lata_l = lata_raw_l - trend[: ic + 1]
        lata_r = lata_raw_r - trend[ic:]
    else:
        z_lev = z_smooth
        lata_l, lata_r = lata_raw_l, lata_raw_r

    idx_mid = int(np.argmin(np.abs(x)))
    z_ref = z_lev[idx_mid]
    z_plot = z_lev - z_ref
    lata_l0 = lata_l - z_ref
    lata_r0 = lata_r - z_ref

    d_l = lata_l0 - z_plot[: ic + 1]
    d_r = lata_r0 - z_plot[ic:]
    i_l = int(np.argmax(d_l)) if len(d_l) else 0
    i_r_rel = int(np.argmax(d_r)) if len(d_r) else 0
    i_r = ic + i_r_rel
    return {
        "x": x, "z_plot": z_plot, "idx_crown": ic, "x_l": x_l, "x_r": x_r,
        "lata_l": lata_l0, "lata_r": lata_r0,
        "rut_l_calc": rl_calc, "rut_r_calc": rr_calc,
        "i_l": i_l, "i_r": i_r, "i_r_rel": i_r_rel,
        "x_kl": float(x_l[i_l]) if len(x_l) else 0.0, "x_kp": float(x[i_r]),
        "y_l_prof": float(z_plot[i_l]), "y_l_lata": float(lata_l0[i_l]),
        "y_r_prof": float(z_plot[i_r]), "y_r_lata": float(lata_r0[i_r_rel]),
    }


def level_profile_for_map(x, z, x_grid, window=31):
    """Profil LCMS wyrównany (bez spadku poprzecznego) i wyzerowany na osi – do map cieplnych."""
    x = np.asarray(x, dtype=float)
    z = np.asarray(z, dtype=float)
    zs = smooth_profile(z, window)
    if len(x) > 10:
        p = np.polyfit(x, zs, 1)
        zs = zs - p[0] * x
    zs = zs - zs[int(np.argmin(np.abs(x)))]
    return np.interp(x_grid, x, zs, left=np.nan, right=np.nan)


# ==========================================================================
# 5. ZESTAWY DANYCH: SCALANIE PLIKÓW, PIKIETAŻ, SIATKA 1 m, DOPASOWANIE LCMS → RSP
# ==========================================================================
def native_chain(src, guess_start=0.0, step_m=1.0):
    """Pikietaż [m] kolejnych próbek źródła w jego własnym układzie (przed korektą)."""
    t = src["type"]
    if t == TYPE_RSP_RAW:
        if src.get("chain") is not None:
            return np.asarray(src["chain"], dtype=float)
        p = np.asarray(src["prof"], dtype=float)
        return (p - p[0]) * step_m
    if t == TYPE_RSP_KOLC:
        return np.asarray(src["chain"], dtype=float)
    if t == TYPE_LCMS_PROFILE:
        return np.asarray(src["station"], dtype=float) * step_m + guess_start
    if t == TYPE_RSP_GPS:
        return src["gps"]["chain_m"].values.astype(float)
    if t == TYPE_LCMS_GPS:
        return src["df"]["Station"].values.astype(float) + src.get("start_abs", 0.0)
    raise ValueError(t)


def source_default_start(src):
    """Domyślny pikietaż pierwszej próbki źródła (z nazwy pliku / PIC / km)."""
    t = src["type"]
    guess = 0.0
    if t == TYPE_LCMS_PROFILE:
        g = guess_start_from_filename(src["name"])
        guess = g if g is not None else 0.0
    return guess


def source_extent(src, guess_start=0.0, step_m=1.0):
    c = native_chain(src, guess_start, step_m)
    if len(c) == 0:
        return np.nan, np.nan
    return float(np.nanmin(c)), float(np.nanmax(c))


def _sort_unique_idx(chain, tol=1e-6):
    order = np.argsort(chain, kind="stable")
    c = chain[order]
    keep = np.ones(len(c), dtype=bool)
    keep[1:] = np.diff(c) > tol
    return order[keep], int(len(c) - keep.sum())


def lcms_take(L, idx, mirror=False):
    """Nowy zestaw LCMS z podzbioru / w innej kolejności profili (idx)."""
    idx = np.asarray(idx, dtype=np.int64)
    ptr, X, Z = L["ptr"], L["x"], L["z"]
    lens = (ptr[1:] - ptr[:-1])[idx]
    new_ptr = np.zeros(len(idx) + 1, dtype=np.int64)
    new_ptr[1:] = np.cumsum(lens)
    xs, zs = [], []
    for i in idx:
        a, b = ptr[i], ptr[i + 1]
        if mirror:
            xs.append(-X[a:b][::-1])
            zs.append(Z[a:b][::-1])
        else:
            xs.append(X[a:b])
            zs.append(Z[a:b])
    out = dict(L)
    out.update({
        "chain": L["chain"][idx], "rut_l": L["rut_l"][idx], "rut_r": L["rut_r"][idx],
        "complete": L["complete"][idx], "src_idx": L["src_idx"][idx],
        "ptr": new_ptr,
        "x": np.concatenate(xs) if xs else np.zeros(0, np.float32),
        "z": np.concatenate(zs) if zs else np.zeros(0, np.float32),
    })
    return out


def merge_lcms(entries, reverse=False, notes=None):
    """Łączy wiele plików profili LCMS w jeden zestaw posortowany po pikietażu."""
    notes = notes if notes is not None else []
    parts = []
    names = []
    for e in entries:
        if not e.get("use", True) or e["type"] != TYPE_LCMS_PROFILE:
            continue
        s = e["src"]
        nat = native_chain(s, e.get("guess_start", 0.0), e.get("step_m", 1.0))
        shift = float(e.get("start_m", nat[0] if len(nat) else 0.0)) - float(nat[0])
        parts.append((s, nat + shift, len(names)))
        names.append(e["name"])
    if not parts:
        return None
    chain = np.concatenate([p[1] for p in parts])
    src_idx = np.concatenate([np.full(len(p[1]), p[2], dtype=np.int32) for p in parts])
    rut_l = np.concatenate([p[0]["rut_l"] for p in parts])
    rut_r = np.concatenate([p[0]["rut_r"] for p in parts])
    complete = np.concatenate([p[0]["complete"] for p in parts])
    if len(parts) == 1:
        ptr, X, Z = parts[0][0]["ptr"], parts[0][0]["x"], parts[0][0]["z"]
    else:
        lens = np.concatenate([np.diff(p[0]["ptr"]) for p in parts])
        ptr = np.zeros(len(lens) + 1, dtype=np.int64)
        ptr[1:] = np.cumsum(lens)
        X = np.concatenate([p[0]["x"] for p in parts])
        Z = np.concatenate([p[0]["z"] for p in parts])
    L = {"chain": chain, "rut_l": rut_l, "rut_r": rut_r, "complete": complete, "ptr": ptr,
         "x": X, "z": Z, "src_idx": src_idx, "src_names": names}
    order, ndup = _sort_unique_idx(chain)
    if ndup:
        notes.append(f"LCMS: pominięto {ndup} zdublowanych pikietaży (pliki zachodzą na siebie)")
    if ndup or len(order) != len(chain) or np.any(np.diff(chain) < 0):
        L = lcms_take(L, order)
    if reverse:
        n = len(L["chain"])
        c = L["chain"]
        new_chain = (c.min() + c.max()) - c
        idx = np.arange(n)[::-1]
        L = lcms_take(L, idx, mirror=True)
        L["chain"] = new_chain[idx]
        L["rut_l"], L["rut_r"] = L["rut_r"].copy(), L["rut_l"].copy()
        notes.append("LCMS: odwrócono kierunek (pikietaż, L↔P, lustrzane odbicie profili)")
    L["c_min"], L["c_max"] = float(L["chain"][0]), float(L["chain"][-1])
    L["step"] = robust_step(L["chain"])
    uid_src = "|".join(f"{e['ident']}@{e.get('start_m', '')}" for e in entries
                       if e.get("use", True) and e["type"] == TYPE_LCMS_PROFILE) + f"|rev={reverse}"
    L["uid"] = hashlib.md5(uid_src.encode()).hexdigest()
    return L


def merge_rsp(entries, notes=None):
    """Łączy pliki RSP: surowe lasery i/lub koleiny KOLC (osobno), pikietaż po korekcie."""
    notes = notes if notes is not None else []
    raws, kolcs, names = [], [], []
    for e in entries:
        if not e.get("use", True) or e["type"] not in (TYPE_RSP_RAW, TYPE_RSP_KOLC):
            continue
        s = e["src"]
        nat = native_chain(s, 0.0, e.get("step_m", 1.0))
        shift = float(e.get("start_m", nat[0] if len(nat) else 0.0)) - float(nat[0])
        k = len(names)
        names.append(e["name"])
        (raws if e["type"] == TYPE_RSP_RAW else kolcs).append((s, nat + shift, k, shift))
    if not raws and not kolcs:
        return None
    R = {"src_names": names, "raw": None, "kolc": None, "gps": None, "kolc_meta": []}
    if raws:
        n_las = Counter(p[0]["lasers"].shape[1] for p in raws).most_common(1)[0][0]
        raws = [p for p in raws if p[0]["lasers"].shape[1] == n_las]
        chain = np.concatenate([p[1] for p in raws])
        lasers = np.concatenate([p[0]["lasers"] for p in raws])
        prof = np.concatenate([p[0]["prof"] for p in raws])
        sidx = np.concatenate([np.full(len(p[1]), p[2], dtype=np.int32) for p in raws])
        order, ndup = _sort_unique_idx(chain)
        if ndup:
            notes.append(f"RSP: pominięto {ndup} zdublowanych pikietaży (surowe lasery)")
        x_pos = next((p[0]["x_pos"] for p in raws if p[0].get("x_pos")), None)
        R["raw"] = {"chain": chain[order], "lasers": lasers[order], "prof": prof[order],
                    "src_idx": sidx[order], "x_pos": x_pos}
    if kolcs:
        chain = np.concatenate([p[1] for p in kolcs])
        keys = ["L", "P", "PEL", "ML", "MP", "MPEL"]
        data = {k: np.concatenate([p[0]["data"][k] for p in kolcs]) for k in keys}
        sidx = np.concatenate([np.full(len(p[1]), p[2], dtype=np.int32) for p in kolcs])
        order, ndup = _sort_unique_idx(chain)
        if ndup:
            notes.append(f"RSP: pominięto {ndup} zdublowanych pikietaży (KOLC)")
        R["kolc"] = {"chain": chain[order], "src_idx": sidx[order],
                     **{k: v[order] for k, v in data.items()}}
        for p in kolcs:
            for c in p[0].get("comments", []):
                R["kolc_meta"].append({**c, "plik_kolc": p[0]["name"]})
    los, his = [], []
    for part in (R["raw"], R["kolc"]):
        if part is not None and len(part["chain"]):
            los.append(part["chain"][0])
            his.append(part["chain"][-1])
    R["c_min"], R["c_max"] = float(min(los)), float(max(his))
    uid_src = "|".join(f"{e['ident']}@{e.get('start_m', '')}" for e in entries
                       if e.get("use", True) and e["type"] in (TYPE_RSP_RAW, TYPE_RSP_KOLC))
    R["uid"] = hashlib.md5(uid_src.encode()).hexdigest()
    return R


def _overlap_len(a0, a1, b0, b1):
    return min(a1, b1) - max(a0, b0)


def merge_gps(entries, prof_range, kinds, camera="auto", notes=None, label="GPS"):
    """Zbiera punkty GPS z plików (i z danych wbudowanych) i dopasowuje pikietaż do profili.

    kinds – typy plików GPS tego systemu, np. (TYPE_RSP_GPS,) lub (TYPE_LCMS_GPS,).
    Zwraca DataFrame [chain_m, LAT, LON, src] lub None.
    """
    notes = notes if notes is not None else []
    parts = []  # (df, kind, name, priority)
    for e in entries:
        if not e.get("use", True):
            continue
        s = e["src"]
        if e["type"] in kinds:
            if e["type"] == TYPE_LCMS_GPS:
                df = s["df"]
                cam = camera
                if cam in (None, "", "auto"):
                    cam = df["Buchst"].value_counts().idxmax()
                d = df[df["Buchst"] == cam]
                if d.empty:
                    d = df
                nat = d["Station"].values.astype(float) + s.get("start_abs", 0.0)
                shift = float(e.get("start_m", nat.min())) - float(
                    (df["Station"].values.astype(float) + s.get("start_abs", 0.0)).min())
                g = pd.DataFrame({"chain_m": nat + shift, "LAT": d["LAT"].values, "LON": d["LON"].values})
                parts.append((g, "km", e["name"], 0))
            else:
                g = s["gps"][["chain_m", "LAT", "LON"]].copy()
                nat0 = float(g["chain_m"].min())
                shift = float(e.get("start_m", nat0)) - nat0
                g["chain_m"] = g["chain_m"] + shift
                parts.append((g, s["gps"].attrs.get("chain_kind", "km"), e["name"], 0))
        elif e["type"] in (TYPE_RSP_RAW, TYPE_RSP_KOLC) and TYPE_RSP_GPS in kinds and s.get("gps") is not None:
            g = s["gps"][["chain_m", "LAT", "LON"]].copy()
            # GPS wbudowany dziedziczy korektę pikietażu pliku nadrzędnego
            nat = native_chain(s, 0.0, e.get("step_m", 1.0))
            shift = float(e.get("start_m", nat[0])) - float(nat[0])
            g["chain_m"] = g["chain_m"] + shift
            parts.append((g, s["gps"].attrs.get("chain_kind", "km"), e["name"] + " (wbudowany)", 1))
    if not parts:
        return None
    out = []
    for g, kind, name, prio in sorted(parts, key=lambda p: p[3]):
        g = g.dropna(subset=["chain_m", "LAT", "LON"]).copy()
        if g.empty:
            continue
        lo, hi = float(g["chain_m"].min()), float(g["chain_m"].max())
        if prof_range is not None:
            ov = _overlap_len(lo, hi, prof_range[0], prof_range[1])
            if kind == "dist" or ov <= 0:
                sh = prof_range[0] - lo
                g["chain_m"] = g["chain_m"] + sh
                notes.append(f"{label}: '{name}' nie pokrywa się z pikietażem profili – "
                             f"zakotwiczono początek w {km_str(prof_range[0])}")
        g["src"] = name
        out.append(g)
    if not out:
        return None
    df = pd.concat(out, ignore_index=True)
    df = df.sort_values("chain_m", kind="stable").drop_duplicates("chain_m").reset_index(drop=True)
    return df


def shift_system(obj, delta, kind):
    """Przesuwa cały system (profile + GPS) o delta [m]."""
    if obj is None:
        return None
    if kind == "rsp":
        for part in ("raw", "kolc"):
            if obj.get(part) is not None:
                obj[part]["chain"] = obj[part]["chain"] + delta
        obj["c_min"] += delta
        obj["c_max"] += delta
    else:
        obj["chain"] = obj["chain"] + delta
        obj["c_min"] += delta
        obj["c_max"] += delta
    if obj.get("gps") is not None:
        g = obj["gps"].copy()
        g["chain_m"] = g["chain_m"] + delta
        obj["gps"] = g
    return obj


def assemble_datasets(entries, camera="auto", reverse_lcms=False, relative_mode="auto"):
    """Główny krok: scala pliki w systemy RSP i LCMS na wspólnej osi pikietażu."""
    notes = []
    rsp = merge_rsp(entries, notes)
    lcms = merge_lcms(entries, reverse=reverse_lcms, notes=notes)
    if rsp is not None:
        rsp["gps"] = merge_gps(entries, (rsp["c_min"], rsp["c_max"]), (TYPE_RSP_GPS,), notes=notes, label="GPS RSP")
    if lcms is not None:
        lcms["gps"] = merge_gps(entries, (lcms["c_min"], lcms["c_max"]), (TYPE_LCMS_GPS,),
                                camera=camera, notes=notes, label="GPS LCMS")
    rel = False
    if rsp is not None and lcms is not None:
        ov = _overlap_len(rsp["c_min"], rsp["c_max"], lcms["c_min"], lcms["c_max"])
        if relative_mode == "tak" or (relative_mode == "auto" and ov <= 0):
            rel = True
            d_r, d_l = -rsp["c_min"], -lcms["c_min"]
            rsp = shift_system(rsp, d_r, "rsp")
            lcms = shift_system(lcms, d_l, "lcms")
            notes.append("Pikietaże RSP i LCMS nie pokrywają się – włączono tryb względny "
                         "(oba systemy od 0 m; offset koryguje różnicę startu)")
        elif ov < 0.1 * min(rsp["c_max"] - rsp["c_min"], lcms["c_max"] - lcms["c_min"]):
            notes.append("Wspólny zakres pikietażu jest bardzo mały – sprawdź pikietaże startowe plików")
    return {"rsp": rsp, "lcms": lcms, "rel": rel, "notes": notes}


# ---------- siatka 1 m i seria na siatce ----------
def to_grid(chain, values, c0, n):
    """Średnia wartości w koszykach 1 m: indeks = round(pikietaż - c0)."""
    chain = np.asarray(chain, dtype=float)
    values = np.asarray(values, dtype=float)
    idx = np.rint(chain - c0).astype(np.int64)
    ok = (idx >= 0) & (idx < n) & ~np.isnan(values)
    s = np.bincount(idx[ok], weights=values[ok], minlength=n)
    c = np.bincount(idx[ok], minlength=n)
    out = np.full(n, np.nan)
    m = c > 0
    out[m] = s[m] / c[m]
    return out


class Grid:
    """Regularna siatka pikietażu o kroku 1 m."""

    def __init__(self, c_min, c_max):
        self.c0 = float(np.floor(c_min))
        self.n = int(np.ceil(c_max) - self.c0) + 1
        self.chain = self.c0 + np.arange(self.n, dtype=float)

    def index(self, chain):
        return np.rint(np.asarray(chain, dtype=float) - self.c0).astype(np.int64)


def lcms_series(L, lcms_calc=None):
    """Serie koleiny LCMS (z pliku i wyliczone) na siatce 1 m LCMS."""
    g = Grid(L["c_min"], L["c_max"])
    s = {"grid": g,
         "file_L": to_grid(L["chain"], L["rut_l"], g.c0, g.n),
         "file_R": to_grid(L["chain"], L["rut_r"], g.c0, g.n)}
    if lcms_calc is not None:
        s["calc_L"] = to_grid(L["chain"], lcms_calc[0], g.c0, g.n)
        s["calc_R"] = to_grid(L["chain"], lcms_calc[1], g.c0, g.n)
    # indeks profilu LCMS na siatce (-1 gdy brak)
    idx = np.full(g.n, -1, dtype=np.int64)
    gi = g.index(L["chain"])
    ok = (gi >= 0) & (gi < g.n)
    idx[gi[ok][::-1]] = np.arange(len(gi))[ok][::-1]
    s["prof_index"] = idx
    return s


def rsp_series(R, raw_ruts=None, kolc_param="avg"):
    """Serie koleiny RSP na siatce 1 m (z laserów i/lub z KOLC)."""
    g = Grid(R["c_min"], R["c_max"])
    s = {"grid": g, "raw_L": None, "raw_R": None, "kolc_L": None, "kolc_R": None, "kolc_PEL": None}
    if R.get("raw") is not None and raw_ruts is not None:
        s["raw_L"] = to_grid(R["raw"]["chain"], raw_ruts[0], g.c0, g.n)
        s["raw_R"] = to_grid(R["raw"]["chain"], raw_ruts[1], g.c0, g.n)
        idx = np.full(g.n, -1, dtype=np.int64)
        gi = g.index(R["raw"]["chain"])
        ok = (gi >= 0) & (gi < g.n)
        idx[gi[ok][::-1]] = np.arange(len(gi))[ok][::-1]
        s["raw_index"] = idx
    if R.get("kolc") is not None:
        k = R["kolc"]
        lk, rk, pk = ("ML", "MP", "MPEL") if kolc_param == "max" else ("L", "P", "PEL")
        s["kolc_L"] = to_grid(k["chain"], k[lk], g.c0, g.n)
        s["kolc_R"] = to_grid(k["chain"], k[rk], g.c0, g.n)
        s["kolc_PEL"] = to_grid(k["chain"], k[pk], g.c0, g.n)
    return s


def rsp_pick(s, source="auto"):
    """Wybór źródła RSP: 'auto' (lasery, luki uzupełnione z KOLC), 'raw', 'kolc'. Zwraca (L, R, opis)."""
    rawL, rawR, kL, kR = s["raw_L"], s["raw_R"], s["kolc_L"], s["kolc_R"]
    if source == "raw" and rawL is not None:
        return rawL, rawR, "RSP: z laserów (Wirtualna Łata)"
    if source == "kolc" and kL is not None:
        return kL, kR, "RSP: z pliku KOLC"
    if rawL is not None and kL is not None:
        L = np.where(np.isnan(rawL), kL, rawL)
        R = np.where(np.isnan(rawR), kR, rawR)
        return L, R, "RSP: lasery (luki uzupełnione z KOLC)"
    if rawL is not None:
        return rawL, rawR, "RSP: z laserów (Wirtualna Łata)"
    if kL is not None:
        return kL, kR, "RSP: z pliku KOLC"
    n = s["grid"].n
    return np.full(n, np.nan), np.full(n, np.nan), "RSP: brak danych"


# ---------- przesunięcia ----------
def make_offsets(chain, global_off=0.0, drift_per_km=0.0, ref=0.0, sec_edges=None, sec_offsets=None):
    """Przesunięcie [m] dla każdego pikietażu RSP: LCMS_chain = RSP_chain + off.

    off = offset_globalny + offset_sekcji + dryf [m/km] * (pikietaż - ref) / 1000
    """
    chain = np.asarray(chain, dtype=float)
    off = np.full(chain.shape, float(global_off))
    if drift_per_km:
        off = off + drift_per_km * (chain - ref) / 1000.0
    if sec_edges is not None and sec_offsets is not None and len(sec_offsets):
        si = np.clip(np.searchsorted(sec_edges, chain, side="right") - 1, 0, len(sec_offsets) - 1)
        off = off + np.asarray(sec_offsets, dtype=float)[si]
    return off


def build_pairs(rsp_s, lcms_s, rsp_source, lcms_source, a0, a1, off_fn, kolc_param="avg",
                swap_lr=False):
    """Tabela porównawcza RSP vs LCMS na siatce 1 m RSP w zakresie [a0, a1].

    off_fn(chain_array) -> przesunięcie dla każdego pikietażu RSP.
    """
    g = rsp_s["grid"]
    i0 = max(0, int(np.floor(a0 - g.c0)))
    i1 = min(g.n - 1, int(np.ceil(a1 - g.c0)))
    chain = g.chain[i0:i1 + 1]
    rL, rR, rdesc = rsp_pick(rsp_s, rsp_source)
    df = pd.DataFrame({"Pikietaz": chain, "RSP_L": rL[i0:i1 + 1], "RSP_R": rR[i0:i1 + 1]})
    df["RSP_Max"] = np.maximum(df["RSP_L"].values, df["RSP_R"].values)
    if rsp_s.get("kolc_PEL") is not None:
        df["RSP_PEL"] = rsp_s["kolc_PEL"][i0:i1 + 1]
    if lcms_s is not None:
        lg = lcms_s["grid"]
        off = off_fn(chain)
        target = chain + off
        j = lg.index(target)
        ok = (j >= 0) & (j < lg.n)
        jj = np.where(ok, j, 0)

        def take(arr):
            if arr is None:
                return np.full(len(chain), np.nan)
            v = arr[jj].astype(float)
            v[~ok] = np.nan
            return v

        fL, fR = take(lcms_s["file_L"]), take(lcms_s["file_R"])
        cL, cR = take(lcms_s.get("calc_L")), take(lcms_s.get("calc_R"))
        if swap_lr:
            fL, fR = fR, fL
            cL, cR = cR, cL
        if lcms_source == "calc":
            pL, pR = cL, cR
        else:
            pL, pR = fL, fR
        # brak koleiny w pliku -> uzupełnij z łaty
        if np.all(np.isnan(pL)) and not np.all(np.isnan(cL)):
            pL, pR = cL, cR
        df["Off"] = off
        df["LCMS_Pikietaz"] = np.where(ok, target, np.nan)
        df["LCMS_L"], df["LCMS_R"] = pL, pR
        df["LCMS_Max"] = np.maximum(pL, pR)
        if lcms_source == "both":
            df["LCMSc_L"], df["LCMSc_R"] = cL, cR
            df["LCMSc_Max"] = np.maximum(cL, cR)
        df["LCMS_idx"] = np.where(ok, lcms_s["prof_index"][jj], -1)
    return df, rdesc


# ==========================================================================
# 6. STATYSTYKI ZGODNOŚCI RSP vs LCMS
# ==========================================================================
def clean_pair(a, b):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    m = ~np.isnan(a) & ~np.isnan(b)
    return a[m], b[m]


def _pearson(a, b):
    if len(a) < 3 or np.std(a) == 0 or np.std(b) == 0:
        return np.nan
    return float(np.corrcoef(a, b)[0, 1])


def _lag1_autocorr(x):
    x = np.asarray(x, dtype=float)
    if len(x) < 10 or np.std(x) == 0:
        return 0.0
    x = x - x.mean()
    return float(np.sum(x[1:] * x[:-1]) / np.sum(x * x))


def icc_a1(a, b):
    """ICC(A,1): zgodność bezwzględna dwóch pomiarów (dwukierunkowy model losowy)."""
    n = len(a)
    if n < 3:
        return np.nan
    X = np.c_[a, b]
    k = 2
    gm = X.mean()
    ms_r = k * np.sum((X.mean(axis=1) - gm) ** 2) / (n - 1)
    ms_c = n * np.sum((X.mean(axis=0) - gm) ** 2) / (k - 1)
    sse = np.sum((X - gm) ** 2) - k * np.sum((X.mean(axis=1) - gm) ** 2) - n * np.sum((X.mean(axis=0) - gm) ** 2)
    ms_e = sse / ((n - 1) * (k - 1))
    den = ms_r + (k - 1) * ms_e + k * (ms_c - ms_e) / n
    return float((ms_r - ms_e) / den) if den else np.nan


def deming_fit(a, b):
    """Regresja ortogonalna (Deming, λ = 1): b = ic + sl·a – nie ulega "rozcieńczeniu" przez szum w a."""
    sxx = np.var(a)
    syy = np.var(b)
    sxy = np.mean((a - a.mean()) * (b - b.mean()))
    if abs(sxy) < 1e-12:
        return np.nan, np.nan
    sl = (syy - sxx + math.sqrt((syy - sxx) ** 2 + 4 * sxy ** 2)) / (2 * sxy)
    return float(sl), float(b.mean() - sl * a.mean())


def deming_block_bootstrap(a, b, n_boot=150, block=50, seed=7):
    """Przedział ufności nachylenia Deminga (bootstrap blokowy – uwzględnia autokorelację wzdłuż trasy)."""
    n = len(a)
    if n < 4 * block:
        block = max(2, n // 10)
    rng = np.random.default_rng(seed)
    nb = int(math.ceil(n / block))
    out = []
    for _ in range(n_boot):
        starts = rng.integers(0, max(1, n - block), size=nb)
        idx = np.concatenate([np.arange(s, s + block) for s in starts])[:n]
        sl, _ = deming_fit(a[idx], b[idx])
        if not np.isnan(sl):
            out.append(sl)
    if len(out) < 20:
        return np.nan, np.nan
    return float(np.percentile(out, 2.5)), float(np.percentile(out, 97.5))


def agreement_stats(a, b, tol=2.0, deming_ci=True):
    """Komplet miar zgodności. a = RSP (odniesienie), b = LCMS; różnica d = b − a."""
    a, b = clean_pair(a, b)
    n = len(a)
    nan = np.nan
    out = {"n": n}
    if n < 3:
        return out
    d = b - a
    m = (a + b) / 2.0
    bias = float(np.mean(d))
    sd = float(np.std(d, ddof=1))
    rho1 = _lag1_autocorr(d)
    n_eff = max(3.0, n * (1 - rho1) / (1 + rho1)) if rho1 > 0 else float(n)
    se_bias = sd / math.sqrt(n_eff)
    out.update({
        "mean_rsp": float(a.mean()), "mean_lcms": float(b.mean()),
        "bias": bias, "sd": sd, "loa_lo": bias - 1.96 * sd, "loa_hi": bias + 1.96 * sd,
        "bias_ci_lo": bias - 1.96 * se_bias, "bias_ci_hi": bias + 1.96 * se_bias, "n_eff": float(n_eff),
        "rmse": float(np.sqrt(np.mean(d ** 2))), "mae": float(np.mean(np.abs(d))),
        "medae": float(np.median(np.abs(d))), "p95_abs": float(np.percentile(np.abs(d), 95)),
        "max_abs": float(np.max(np.abs(d))), "median_diff": float(np.median(d)),
    })
    r = _pearson(a, b)
    rank_a = pd.Series(a).rank().values
    rank_b = pd.Series(b).rank().values
    out["pearson"] = r
    out["spearman"] = _pearson(rank_a, rank_b)
    out["r2"] = r * r if not np.isnan(r) else nan
    va, vb = np.var(a), np.var(b)
    cov = np.mean((a - a.mean()) * (b - b.mean()))
    den = va + vb + (a.mean() - b.mean()) ** 2
    out["ccc"] = float(2 * cov / den) if den else nan
    out["icc"] = icc_a1(a, b)
    # regresja LCMS = a0 + a1 * RSP
    if np.std(a) > 0:
        sl, ic = np.polyfit(a, b, 1)
        res = b - (ic + sl * a)
        sxx = np.sum((a - a.mean()) ** 2)
        se_sl = math.sqrt(np.sum(res ** 2) / max(1, n - 2) / sxx) if sxx else nan
        out.update({"slope": float(sl), "intercept": float(ic),
                    "slope_ci_lo": float(sl - 1.96 * se_sl), "slope_ci_hi": float(sl + 1.96 * se_sl)})
    out["sd_rsp"], out["sd_lcms"] = float(np.std(a, ddof=1)), float(np.std(b, ddof=1))
    out["sd_ratio"] = out["sd_lcms"] / out["sd_rsp"] if out["sd_rsp"] > 0 else nan
    d_sl, d_ic = deming_fit(a, b)
    out["deming_slope"], out["deming_intercept"] = d_sl, d_ic
    if deming_ci and n >= 50 and not np.isnan(d_sl):
        out["deming_ci_lo"], out["deming_ci_hi"] = deming_block_bootstrap(a, b)
    # błąd proporcjonalny (różnica vs średnia) i heteroskedastyczność
    if np.std(m) > 0:
        sl2, ic2 = np.polyfit(m, d, 1)
        res2 = d - (ic2 + sl2 * m)
        sxx2 = np.sum((m - m.mean()) ** 2)
        se2 = math.sqrt(np.sum(res2 ** 2) / max(1, n - 2) / sxx2) if sxx2 else nan
        out["prop_slope"] = float(sl2)
        out["prop_t"] = float(sl2 / se2) if se2 and not np.isnan(se2) else nan
        out["hetero_r"] = _pearson(m, np.abs(d))
        sl3, _ = np.polyfit(m, np.abs(d), 1)
        out["hetero_slope"] = float(sl3)
    ad = np.abs(d)
    out["within_tol"] = float(np.mean(ad <= tol) * 100)
    for t in (1, 2, 3, 5):
        out[f"within_{t}"] = float(np.mean(ad <= t) * 100)
    ra, rb = np.round(a, 1), np.round(b, 1)
    out["share_lcms_higher"] = float(np.mean(rb > ra) * 100)
    out["share_rsp_higher"] = float(np.mean(ra > rb) * 100)
    out["share_equal"] = float(np.mean(ra == rb) * 100)
    pos = m > 0
    out["rpd_median"] = float(np.median(ad[pos] / m[pos] * 100)) if pos.any() else nan
    return out


def bland_altman_stats(a, b):
    """Statystyki Blanda–Altmana: (średnie, różnice, bias, sd, LoA górna, LoA dolna)."""
    a, b = clean_pair(a, b)
    means = (a + b) / 2.0
    diffs = b - a
    md = float(np.mean(diffs)) if len(diffs) else np.nan
    sd = float(np.std(diffs, ddof=1)) if len(diffs) > 1 else 0.0
    return means, diffs, md, sd, md + 1.96 * sd, md - 1.96 * sd


def build_rpd_table_df(rsp_v, lcms_v):
    """Tabela RPD w układzie analizy GDDKiA (0–30% oraz >30%) – bez zmian."""
    rsp_v, lcms_v = clean_pair(rsp_v, lcms_v)
    mean_v = (rsp_v + lcms_v) / 2.0
    diff_abs = np.abs(lcms_v - rsp_v)
    rpd = np.where(mean_v > 0, (diff_abs / np.where(mean_v > 0, mean_v, 1)) * 100.0, 0.0)
    m_low = rpd <= 30.0
    m_high = rpd > 30.0
    n_tot = len(rsp_v)

    def sub_stat(m_grp, label):
        n_grp = int(np.sum(m_grp))
        if n_grp == 0:
            return {"zakres RPD": label, "n": 0, "%": "0%", "RSP > LCMS": 0, "Wyniki równe": 0,
                    "LCMS > RSP": 0, "%RSP > LCMS": "0%", "%Wyniki równe": "0%", "%LCMS > RSP": "0%"}
        r_sub = np.round(rsp_v[m_grp], 1)
        l_sub = np.round(lcms_v[m_grp], 1)
        rsp_gt = int(np.sum(r_sub > l_sub))
        eq = int(np.sum(r_sub == l_sub))
        lcms_gt = int(np.sum(l_sub > r_sub))
        return {"zakres RPD": label, "n": n_grp, "%": f"{(n_grp / max(1, n_tot)) * 100:.0f}%",
                "RSP > LCMS": rsp_gt, "Wyniki równe": eq, "LCMS > RSP": lcms_gt,
                "%RSP > LCMS": f"{(rsp_gt / n_grp) * 100:.0f}%",
                "%Wyniki równe": f"{(eq / n_grp) * 100:.0f}%",
                "%LCMS > RSP": f"{(lcms_gt / n_grp) * 100:.0f}%"}

    row_sum = {"zakres RPD": "suma", "n": n_tot, "%": "100%", "RSP > LCMS": "-", "Wyniki równe": "-",
               "LCMS > RSP": "-", "%RSP > LCMS": "-", "%Wyniki równe": "-", "%LCMS > RSP": "-"}
    return pd.DataFrame([sub_stat(m_low, "0-30%"), sub_stat(m_high, ">30%"), row_sum])


def rpd_fine_table(rsp_v, lcms_v, edges=(0, 10, 20, 30, 50, 100, np.inf)):
    """RPD w drobniejszych przedziałach (0–10, 10–20, 20–30, 30–50, 50–100, >100 %)."""
    a, b = clean_pair(rsp_v, lcms_v)
    m = (a + b) / 2.0
    pos = m > 0
    rpd = np.full(len(a), np.nan)
    rpd[pos] = np.abs(b - a)[pos] / m[pos] * 100
    rows = []
    for lo, hi in zip(edges[:-1], edges[1:]):
        sel = (rpd >= lo) & (rpd < hi) if np.isfinite(hi) else (rpd >= lo)
        lab = f"{lo:g}–{hi:g}%" if np.isfinite(hi) else f">{lo:g}%"
        rows.append({"RPD": lab, "n": int(sel.sum()), "udział %": round(100 * sel.sum() / max(1, pos.sum()), 1)})
    return pd.DataFrame(rows)


def depth_bin_table(a, b, edges=(0, 5, 10, 15, 20, 30, 1e9), tol=2.0):
    """Zgodność w funkcji głębokości (przedziały wg średniej z obu systemów)."""
    a, b = clean_pair(a, b)
    m = (a + b) / 2.0
    rows = []
    for lo, hi in zip(edges[:-1], edges[1:]):
        sel = (m >= lo) & (m < hi)
        if sel.sum() == 0:
            continue
        d = b[sel] - a[sel]
        lab = f"{lo:g}–{hi:g} mm" if hi < 1e8 else f"≥ {lo:g} mm"
        rows.append({
            "Przedział (średnia)": lab, "n": int(sel.sum()),
            "RSP śr. [mm]": round(float(a[sel].mean()), 2), "LCMS śr. [mm]": round(float(b[sel].mean()), 2),
            "Bias [mm]": round(float(d.mean()), 2), "SD różnic [mm]": round(float(d.std(ddof=1)) if sel.sum() > 1 else 0.0, 2),
            "RMSE [mm]": round(float(np.sqrt(np.mean(d ** 2))), 2),
            f"|Δ| ≤ {tol:g} mm [%]": round(float(np.mean(np.abs(d) <= tol) * 100), 1),
        })
    return pd.DataFrame(rows)


# ---------- klasy DSN ----------
def class_confusion(a, b):
    """Macierz pomyłek klas DSN (wiersze: RSP, kolumny: LCMS) + miary zgodności klas."""
    a, b = clean_pair(a, b)
    ca, cb = dsn_class_index(a), dsn_class_index(b)
    k = 4
    M = np.zeros((k, k), dtype=int)
    for i, j in zip(ca, cb):
        M[i, j] += 1
    n = int(M.sum())
    if n == 0:
        return M, {"n": 0}
    exact = float(np.trace(M) / n * 100)
    idx = np.arange(k)
    near = float(sum(M[i, j] for i in idx for j in idx if abs(i - j) <= 1) / n * 100)
    rows, cols = M.sum(axis=1), M.sum(axis=0)
    E = np.outer(rows, cols) / n
    W = np.abs(idx[:, None] - idx[None, :]) / (k - 1)
    den = float(np.sum(W * E))
    kappa_w = 1 - float(np.sum(W * M)) / den if den > 0 else np.nan
    po = np.trace(M) / n
    pe = float(np.sum(rows * cols)) / (n * n)
    kappa = (po - pe) / (1 - pe) if pe < 1 else np.nan
    lcms_higher = float(sum(M[i, j] for i in idx for j in idx if j > i) / n * 100)
    lcms_lower = float(sum(M[i, j] for i in idx for j in idx if j < i) / n * 100)
    return M, {"n": n, "exact": exact, "within1": near, "kappa": float(kappa), "kappa_w": float(kappa_w),
               "lcms_higher": lcms_higher, "lcms_lower": lcms_lower}


# ---------- odcinki 50 m (DSN 2019): Hm = E(h) + 2 Dh ----------
def hm_value(e, d):
    return int(np.round(e + 2 * d))


def hm_table(df, length=50.0, origin=None, min_cover=0.0, has_lcms=None):
    """Ocena odcinkowa Hm = E(h) + 2·Dh na odcinkach `length` m (jak dotychczas, ale po pikietażu).

    df – tabela z kolumnami Pikietaz, RSP_L/R/Max oraz (opcjonalnie) LCMS_L/R/Max.
    Kolumny wynikowe zachowują nazwy z poprzedniej wersji ("RSP L_KOL Hm" itd.).
    """
    if df is None or len(df) == 0:
        return pd.DataFrame()
    if has_lcms is None:
        has_lcms = "LCMS_L" in df.columns
    org = float(df["Pikietaz"].min()) if origin is None else float(origin)
    sec = np.floor((df["Pikietaz"].values - org) / length).astype(np.int64)
    work = df.copy()
    work["_sec"] = sec
    rows = []
    systems = [("RSP", "RSP")] + ([("LCMS", "LCMS")] if has_lcms else [])
    mapping = (("L", "L_KOL"), ("R", "P_KOL"), ("Max", "MAX_KOL"))
    for s_id, g in work.groupby("_sec", sort=True):
        od = org + s_id * length
        do = od + length - 1
        rec = {"Od [m]": od, "Do [m]": do,
               "Odcinek": f"{km_str(od)} – {km_str(do)}",
               "n RSP": int(g["RSP_Max"].notna().sum())}
        if has_lcms:
            rec["n LCMS"] = int(g["LCMS_Max"].notna().sum())
        rec["Pokrycie %"] = round(100.0 * rec["n RSP" if not has_lcms else "n LCMS"] / length, 0)
        for sysn, pre in systems:
            for key, label in mapping:
                v = g[f"{pre}_{key}"].values.astype(float)
                v = v[~np.isnan(v)]
                if len(v) == 0:
                    rec.update({f"{sysn} {label} E(h)": np.nan, f"{sysn} {label} Dh": np.nan,
                                f"{sysn} {label} Hm": np.nan, f"{sysn} {label} Klasa": "?"})
                    continue
                e, dd = float(np.mean(v)), float(np.std(v, ddof=0))
                hm = hm_value(e, dd)
                rec.update({f"{sysn} {label} E(h)": round(e, 2), f"{sysn} {label} Dh": round(dd, 2),
                            f"{sysn} {label} Hm": hm, f"{sysn} {label} Klasa": get_dsn_class(hm)})
        rows.append(rec)
    out = pd.DataFrame(rows)
    if min_cover > 0 and len(out):
        out = out[out["Pokrycie %"] >= min_cover].reset_index(drop=True)
    return out


# ---------- statystyki w oknach i odcinkach ----------
def rolling_agreement(df, win=500, hop=100, a_col="RSP_Max", b_col="LCMS_Max"):
    """Zgodność w oknach przesuwnych po pikietażu (bias, RMSE, korelacja)."""
    if df is None or len(df) == 0 or b_col not in df.columns:
        return pd.DataFrame()
    ch = df["Pikietaz"].values
    a = df[a_col].values.astype(float)
    b = df[b_col].values.astype(float)
    lo, hi = float(ch.min()), float(ch.max())
    rows = []
    start = lo
    while start <= hi:
        end = start + win
        sel = (ch >= start) & (ch < end)
        x, y = clean_pair(a[sel], b[sel])
        if len(x) >= max(10, 0.2 * win):
            d = y - x
            rows.append({"Pikietaz": start + win / 2.0, "n": len(x), "RSP": x.mean(), "LCMS": y.mean(),
                         "Bias": d.mean(), "RMSE": float(np.sqrt(np.mean(d ** 2))),
                         "r": _pearson(x, y), "SD": float(d.std(ddof=1))})
        start += hop
    return pd.DataFrame(rows)


def section_agreement(df, length=1000.0, origin=None, a_col="RSP_Max", b_col="LCMS_Max", tol=2.0):
    """Tabela zgodności na odcinkach o stałej długości (np. po 1 km)."""
    if df is None or len(df) == 0 or b_col not in df.columns:
        return pd.DataFrame()
    org = float(df["Pikietaz"].min()) if origin is None else float(origin)
    sec = np.floor((df["Pikietaz"].values - org) / length).astype(np.int64)
    rows = []
    for s_id in np.unique(sec):
        g = df[sec == s_id]
        x, y = clean_pair(g[a_col].values, g[b_col].values)
        if len(x) < 5:
            continue
        d = y - x
        od = org + s_id * length
        rows.append({
            "Od [m]": od, "Do [m]": od + length - 1, "Odcinek": f"{km_str(od)} – {km_str(od + length - 1)}",
            "n": len(x), "RSP śr. [mm]": round(float(x.mean()), 2), "LCMS śr. [mm]": round(float(y.mean()), 2),
            "Bias [mm]": round(float(d.mean()), 2), "SD [mm]": round(float(d.std(ddof=1)), 2),
            "RMSE [mm]": round(float(np.sqrt(np.mean(d ** 2))), 2),
            "r": round(_pearson(x, y), 3) if len(x) > 3 else np.nan,
            f"|Δ| ≤ {tol:g} mm [%]": round(float(np.mean(np.abs(d) <= tol) * 100), 1),
            "Hm RSP": hm_value(x.mean(), x.std()), "Hm LCMS": hm_value(y.mean(), y.std()),
        })
    out = pd.DataFrame(rows)
    if len(out):
        out["Klasa RSP"] = out["Hm RSP"].apply(get_dsn_class)
        out["Klasa LCMS"] = out["Hm LCMS"].apply(get_dsn_class)
    return out


def corr_label(r):
    if r is None or np.isnan(r):
        return "nieokreślona"
    ar = abs(r)
    if ar < 0.3:
        return "bardzo słaba"
    if ar < 0.5:
        return "słaba"
    if ar < 0.7:
        return "umiarkowana"
    if ar < 0.9:
        return "silna"
    return "bardzo silna"


# ==========================================================================
# 7. SYNCHRONIZACJA PIKIETAŻU RSP ↔ LCMS
# ==========================================================================
def highpass(y, win):
    """Odjęcie wolnozmiennego trendu (średnia krocząca) – wyostrza korelację krzyżową."""
    y = np.asarray(y, dtype=float)
    if not win or win < 3:
        return y.copy()
    return y - nan_moving_average(y, int(win) | 1)


def corr_vs_lag(a, b, a_c0, b_c0, lags, min_n=30):
    """Korelacja Pearsona między serią a (RSP, siatka 1 m od a_c0) i b (LCMS, od b_c0).

    Konwencja jak dotychczas: pikietaż LCMS = pikietaż RSP + lag.
    a, b – tablice lub listy tablic (korelacja łączna po wszystkich parach).
    """
    A = a if isinstance(a, (list, tuple)) else [a]
    B = b if isinstance(b, (list, tuple)) else [b]
    base = int(round(a_c0 - b_c0))
    out = np.full(len(lags), np.nan)
    for k, d in enumerate(lags):
        xs, ys = [], []
        for aa, bb in zip(A, B):
            idx = np.arange(len(aa)) + base + int(d)
            ok = (idx >= 0) & (idx < len(bb))
            x = aa[ok]
            y = bb[idx[ok]]
            v = ~np.isnan(x) & ~np.isnan(y)
            xs.append(x[v])
            ys.append(y[v])
        x = np.concatenate(xs)
        y = np.concatenate(ys)
        if len(x) >= min_n and np.std(x) > 0 and np.std(y) > 0:
            out[k] = np.corrcoef(x, y)[0, 1]
    return out


def best_lag(corr, lags, exclude=5):
    """Najlepsze przesunięcie + jakość (korelacja, drugi pik poza ±`exclude` m)."""
    c = np.asarray(corr, dtype=float)
    if np.all(np.isnan(c)):
        return np.nan, np.nan, np.nan
    k = int(np.nanargmax(c))
    lag = lags[k]
    mask = np.abs(np.asarray(lags) - lag) > exclude
    second = np.nanmax(c[mask]) if mask.any() and not np.all(np.isnan(c[mask])) else np.nan
    return float(lag), float(c[k]), float(second)


def sync_series(pairs_src, metric="Max", hp_win=101):
    """Przygotowuje serie (RSP, LCMS) do korelacji. pairs_src: dict z tablicami na siatkach.

    Zwraca (A_list, B_list) po filtracji górnoprzepustowej.
    """
    names = {"Max": ["Max"], "L": ["L"], "R": ["R"], "L+R": ["L", "R"]}[metric]
    A, B = [], []
    for nm in names:
        a = pairs_src["rsp"][nm]
        b = pairs_src["lcms"][nm]
        A.append(highpass(a, hp_win))
        B.append(highpass(b, hp_win))
    return A, B


def global_offset_search(pairs_src, a_range, max_search=150, metric="Max", hp_win=101):
    """Automatyczny offset globalny (jak dotychczas, ale na siatce 1 m i szybciej)."""
    A, B = sync_series(pairs_src, metric, hp_win)
    c0a, c0b = pairs_src["rsp_c0"], pairs_src["lcms_c0"]
    i0 = max(0, int(a_range[0] - c0a))
    i1 = min(len(A[0]), int(a_range[1] - c0a) + 1)
    As = [a[i0:i1] for a in A]
    lags = np.arange(-max_search, max_search + 1)
    corr = corr_vs_lag(As, B, c0a + i0, c0b, lags)
    lag, c1, c2 = best_lag(corr, lags)
    return lag, c1, c2, lags, corr


def local_lag_map(pairs_src, a_range, win=1000, hop=250, max_search=60, metric="Max", hp_win=101, min_n=200):
    """Lokalne przesunięcia: dla okien po trasie najlepszy lag i jakość dopasowania."""
    A, B = sync_series(pairs_src, metric, hp_win)
    c0a, c0b = pairs_src["rsp_c0"], pairs_src["lcms_c0"]
    lags = np.arange(-max_search, max_search + 1)
    rows = []
    start = float(a_range[0])
    while start + win * 0.5 <= a_range[1]:
        end = min(a_range[1], start + win)
        i0 = max(0, int(start - c0a))
        i1 = min(len(A[0]), int(end - c0a) + 1)
        if i1 - i0 > min_n * 0.5:
            As = [a[i0:i1] for a in A]
            corr = corr_vs_lag(As, B, c0a + i0, c0b, lags, min_n=min_n)
            lag, c1, c2 = best_lag(corr, lags)
            if not np.isnan(lag):
                rows.append({"Pikietaz": (start + end) / 2.0, "Od": start, "Do": end,
                             "Lag [m]": lag, "Korelacja": c1, "Drugi pik": c2,
                             "Wyrazistość": (c1 - c2) if not np.isnan(c2) else np.nan})
        start += hop
    return pd.DataFrame(rows)


def fit_drift(df_lags, ref, min_corr=0.25, min_prom=0.03):
    """Odporna regresja lag(pikietaż) → offset w punkcie ref oraz dryf [m/km]."""
    if df_lags is None or len(df_lags) < 3:
        return None
    d = df_lags.dropna(subset=["Lag [m]", "Korelacja"])
    d = d[(d["Korelacja"] >= min_corr)]
    if "Wyrazistość" in d:
        d = d[(d["Wyrazistość"].fillna(1.0) >= min_prom)]
    if len(d) < 3:
        return None
    x = (d["Pikietaz"].values - ref) / 1000.0
    y = d["Lag [m]"].values
    w = np.clip(d["Korelacja"].values, 0.05, 1.0)
    keep = np.ones(len(x), dtype=bool)
    sl = ic = 0.0
    for _ in range(4):
        if keep.sum() < 3:
            break
        sl, ic = np.polyfit(x[keep], y[keep], 1, w=w[keep])
        res = y - (ic + sl * x)
        mad = 1.4826 * np.median(np.abs(res[keep] - np.median(res[keep]))) or 1.0
        keep = np.abs(res) <= max(3 * mad, 2.0)
    res = y - (ic + sl * x)
    ss_tot = np.sum((y - y.mean()) ** 2)
    r2 = 1 - np.sum(res[keep] ** 2) / ss_tot if ss_tot > 0 and keep.any() else np.nan
    return {"offset_ref": float(ic), "drift_per_km": float(sl), "n": int(keep.sum()), "n_all": int(len(x)),
            "resid_mad": float(np.median(np.abs(res[keep]))) if keep.any() else np.nan,
            "r2": float(r2) if not np.isnan(r2) else np.nan, "ref": float(ref)}


def section_offsets_from_lags(df_lags, edges, default=0.0, min_corr=0.2):
    """Mediana lagów w oknach dla każdej sekcji (do tabeli offsetów sekcji)."""
    out = []
    if df_lags is None or len(df_lags) == 0:
        return [default] * (len(edges) - 1)
    for a, b in zip(edges[:-1], edges[1:]):
        d = df_lags[(df_lags["Pikietaz"] >= a) & (df_lags["Pikietaz"] < b) & (df_lags["Korelacja"] >= min_corr)]
        out.append(float(np.median(d["Lag [m]"])) if len(d) else default)
    return out


# ---------- synchronizacja z GPS ----------
def densify_track(gps, step=5.0):
    """Zagęszcza ślad GPS do punktów co `step` m pikietażu (interpolacja liniowa)."""
    if gps is None or len(gps) < 2:
        return None
    g = gps.sort_values("chain_m")
    c = g["chain_m"].values
    grid = np.arange(np.ceil(c.min()), np.floor(c.max()) + 1, step)
    return pd.DataFrame({"chain_m": grid, "LAT": np.interp(grid, c, g["LAT"].values),
                         "LON": np.interp(grid, c, g["LON"].values)})


def gps_nearest_offsets(rsp_gps, lcms_gps, step=10.0, a_range=None, max_dist=60.0):
    """Dla punktów śladu RSP szuka najbliższego punktu śladu LCMS.

    Zwraca DataFrame [Pikietaz, Dystans [m], Offset [m]] (offset = pikietaż LCMS − pikietaż RSP).
    Zakres wspólny: tylko tam, gdzie oba ślady istnieją (GPS bywa pusty w części trasy).
    """
    r = densify_track(rsp_gps, step)
    l = densify_track(lcms_gps, 5.0)
    if r is None or l is None:
        return pd.DataFrame()
    if a_range is not None:
        r = r[(r["chain_m"] >= a_range[0]) & (r["chain_m"] <= a_range[1])]
    if r.empty:
        return pd.DataFrame()
    lat0 = float(np.mean(l["LAT"]))
    kx = 111320.0 * math.cos(math.radians(lat0))
    ky = 110540.0
    rx, ry = r["LON"].values * kx, r["LAT"].values * ky
    lx, ly = l["LON"].values * kx, l["LAT"].values * ky
    lc = l["chain_m"].values
    out_c, out_d, out_o = [], [], []
    chunk = 400
    for s in range(0, len(r), chunk):
        dx = rx[s:s + chunk, None] - lx[None, :]
        dy = ry[s:s + chunk, None] - ly[None, :]
        dist = np.hypot(dx, dy)
        j = np.argmin(dist, axis=1)
        dm = dist[np.arange(len(j)), j]
        out_c.append(r["chain_m"].values[s:s + chunk])
        out_d.append(dm)
        out_o.append(lc[j] - r["chain_m"].values[s:s + chunk])
    res = pd.DataFrame({"Pikietaz": np.concatenate(out_c), "Dystans [m]": np.concatenate(out_d),
                        "Offset [m]": np.concatenate(out_o)})
    return res[res["Dystans [m]"] <= max_dist].reset_index(drop=True)


def gps_track_separation(rsp_gps, lcms_gps, off_fn, chains):
    """Odległość [m] między punktem RSP (pikietaż g) i LCMS (pikietaż g + offset) dla listy pikietaży."""
    if rsp_gps is None or lcms_gps is None or len(rsp_gps) < 2 or len(lcms_gps) < 2:
        return None
    rg = rsp_gps.sort_values("chain_m")
    lg = lcms_gps.sort_values("chain_m")
    chains = np.asarray(chains, dtype=float)
    ok = (chains >= rg["chain_m"].min()) & (chains <= rg["chain_m"].max())
    target = chains + off_fn(chains)
    ok &= (target >= lg["chain_m"].min()) & (target <= lg["chain_m"].max())
    if not ok.any():
        return None
    c = chains[ok]
    t = target[ok]
    d = haversine_m(np.interp(c, rg["chain_m"], rg["LAT"]), np.interp(c, rg["chain_m"], rg["LON"]),
                    np.interp(t, lg["chain_m"], lg["LAT"]), np.interp(t, lg["chain_m"], lg["LON"]))
    return pd.DataFrame({"Pikietaz": c, "Dystans [m]": d})


# ==========================================================================
# 8. AUTOMATYCZNE WNIOSKI
# ==========================================================================
def _f(v, nd=2, sign=False):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    return f"{v:+.{nd}f}" if sign else f"{v:.{nd}f}"


def agreement_grade(st, tol=2.0):
    """Ogólna ocena zgodności wg CCC Lina i odsetka różnic w tolerancji."""
    if not st or st.get("n", 0) < 3 or np.isnan(st.get("ccc", np.nan)):
        return "brak danych", "info"
    ccc, wt = st["ccc"], st.get("within_tol", 0.0)
    if ccc >= 0.95 and wt >= 90:
        return "bardzo dobra", "ok"
    if ccc >= 0.85 and wt >= 75:
        return "dobra", "ok"
    if ccc >= 0.65 and wt >= 50:
        return "umiarkowana", "warn"
    return "niska", "bad"


def build_conclusions(ctx):
    """Buduje wnioski (lista pozycji + werdykt + markdown) z wyników analizy.

    ctx – słownik przygotowany przez UI (patrz funkcja `make_context`).
    """
    items = []
    tol = ctx.get("tol", 2.0)
    stx = ctx.get("stats", {})
    smax, sl, sr = stx.get("Max", {}), stx.get("L", {}), stx.get("R", {})

    # --- 1. Dane i pokrycie
    cov = ctx.get("cover", {})
    if cov:
        txt = (f"Porównano **{cov['n_valid']:,}** punktów (co 1 m) na odcinku "
               f"**{km_str(cov['c0'])} – {km_str(cov['c1'])}** (~{cov['length_m'] / 1000:.2f} km). "
               f"Pokrycie wspólnymi danymi: **{cov['pct']:.0f}%** zakresu analizy. "
               f"{ctx.get('rsp_desc', '')}; LCMS: {ctx.get('lcms_desc', '')}.").replace(",", " ")
        kind = "ok" if cov["pct"] >= 90 else ("warn" if cov["pct"] >= 60 else "bad")
        items.append({"kind": kind, "title": "Dane i pokrycie", "text": txt})
        if cov.get("gaps_note"):
            items.append({"kind": "info", "title": "Luki w danych", "text": cov["gaps_note"]})

    # --- 2. Synchronizacja
    sync = ctx.get("sync", {})
    if sync:
        parts = [f"Aktualne przesunięcie LCMS względem RSP: **{sync.get('offset', 0):+.1f} m**"]
        if abs(sync.get("drift", 0)) > 1e-9:
            parts.append(f"dryf **{sync['drift']:+.2f} m/km**")
        if sync.get("n_sections"):
            parts.append("offsety sekcji aktywne")
        txt = ", ".join(parts) + "."
        kind = "info"
        if sync.get("auto_lag") is not None and not np.isnan(sync["auto_lag"]):
            cur = sync.get("offset", 0.0)
            if abs(sync["auto_lag"] - cur) >= 2 and sync.get("auto_gain", 0) > 0.02:
                kind = "warn"
                txt += (f" Automatyczna synchronizacja wskazuje **{sync['auto_lag']:+.0f} m** "
                        f"(korelacja {sync['auto_corr']:.2f} wobec {sync.get('cur_corr', np.nan):.2f} przy bieżącym) – "
                        "warto zastosować.")
            else:
                txt += f" Nie ma lepszego globalnego przesunięcia (maksimum korelacji przy {sync['auto_lag']:+.0f} m)."
        fit = sync.get("fit")
        if fit and fit.get("n", 0) >= 5 and fit.get("r2", 0) >= 0.5 and abs(fit["drift_per_km"]) >= 0.3:
            total = fit["drift_per_km"] * cov.get("length_m", 0) / 1000.0 if cov else np.nan
            txt += (f" Lokalne okna wskazują **dryf skali ≈ {fit['drift_per_km']:+.2f} m/km** "
                    f"(≈ {total:+.0f} m na całym odcinku, R² = {fit['r2']:.2f}) – różne kalibracje drogomierzy; "
                    "rozważ korektę dryfu w zakładce Synchronizacja.")
            kind = "warn"
        items.append({"kind": kind, "title": "Synchronizacja pikietażu", "text": txt})

    # --- 3. Zgodność ilościowa
    for lab, st_ in (("koleina maksymalna (MAX_KOL)", smax), ("lewa (L_KOL)", sl), ("prawa (P_KOL)", sr)):
        if not st_ or st_.get("n", 0) < 3:
            continue
        grade, kind = agreement_grade(st_, tol)
        sig = not (st_["bias_ci_lo"] <= 0 <= st_["bias_ci_hi"])
        who = "LCMS wskazuje **głębsze**" if st_["bias"] > 0 else "LCMS wskazuje **płytsze**"
        bias_txt = (f"{who} koleiny o średnio **{abs(st_['bias']):.2f} mm** (95% CI {_f(st_['bias_ci_lo'], 2, True)} … "
                    f"{_f(st_['bias_ci_hi'], 2, True)}; n_eff ≈ {st_['n_eff']:.0f})" if sig else
                    f"brak istotnego błędu systematycznego (bias {_f(st_['bias'], 2, True)} mm, 95% CI "
                    f"{_f(st_['bias_ci_lo'], 2, True)} … {_f(st_['bias_ci_hi'], 2, True)})")
        txt = (f"Zgodność **{grade}**. {bias_txt}. Granice zgodności (±1.96 SD): "
               f"**{_f(st_['loa_lo'], 1, True)} … {_f(st_['loa_hi'], 1, True)} mm**, RMSE {st_['rmse']:.2f} mm, "
               f"MAE {st_['mae']:.2f} mm. W tolerancji ±{tol:g} mm: **{st_['within_tol']:.0f}%** punktów "
               f"(±1 mm: {st_['within_1']:.0f}%, ±3 mm: {st_['within_3']:.0f}%, ±5 mm: {st_['within_5']:.0f}%). "
               f"Korelacja r = **{st_['pearson']:.2f}** ({corr_label(st_['pearson'])}), CCC Lina = {st_['ccc']:.2f}, "
               f"ICC(A,1) = {st_['icc']:.2f}.")
        items.append({"kind": kind, "title": f"Zgodność – {lab}", "text": txt})

    # --- 4. Charakter rozbieżności
    st_ = smax if smax and smax.get("n", 0) >= 3 else (sl if sl else {})
    if st_ and st_.get("n", 0) >= 30:
        notes = []
        if "sd_ratio" in st_ and not np.isnan(st_["sd_ratio"]):
            rr_ = st_["sd_ratio"]
            if rr_ > 1.15:
                notes.append(f"zmienność wzdłuż trasy: SD(LCMS) = {st_['sd_lcms']:.2f} mm wobec SD(RSP) = {st_['sd_rsp']:.2f} mm "
                             f"(stosunek {rr_:.2f}) – LCMS „mocniej faluje” (większy szum lub większa czułość na lokalne zmiany)")
            elif rr_ < 0.87:
                notes.append(f"zmienność wzdłuż trasy: SD(LCMS) = {st_['sd_lcms']:.2f} mm wobec SD(RSP) = {st_['sd_rsp']:.2f} mm "
                             f"(stosunek {rr_:.2f}) – LCMS jest gładszy od RSP")
            else:
                notes.append(f"zmienność wzdłuż trasy porównywalna (SD LCMS/RSP = {rr_:.2f})")
        if "slope" in st_ and "deming_slope" in st_ and not np.isnan(st_["deming_slope"]):
            b1, b2 = st_["slope"], st_["deming_slope"]
            lo_, hi_ = min(b1, b2), max(b1, b2)
            eq = (f"nachylenie regresji LCMS(RSP) zależy od założeń o szumie: MNK {b1:.2f} (RSP bez błędu), "
                  f"ortogonalna {b2:.2f} (równy szum obu systemów)")
            if hi_ < 1.0:
                notes.append(eq + " – oba < 1: LCMS **spłaszcza** zakres względem RSP")
            elif lo_ > 1.0:
                notes.append(eq + " – oba > 1: LCMS **rozciąga** zakres względem RSP")
            else:
                notes.append(eq + " – wartość 1 mieści się między nimi, więc nie ma podstaw, by twierdzić o błędzie proporcjonalnym samej skali")
        if "hetero_r" in st_ and st_["hetero_r"] is not None and not np.isnan(st_["hetero_r"]):
            if st_["hetero_r"] > 0.2 and st_.get("hetero_slope", 0) > 0.03:
                notes.append("rozrzut różnic **rośnie wraz z głębokością koleiny** (kształt „ryby/trójkąta” na wykresie Blanda–Altmana; "
                             f"korelacja |Δ| ze średnią = {st_['hetero_r']:.2f}) – stała tolerancja w mm faworyzuje płytkie koleiny, "
                             "dlatego warto patrzeć na RPD i na zgodność w przedziałach głębokości")
            else:
                notes.append("rozrzut różnic jest w przybliżeniu stały w całym zakresie głębokości")
        if "prop_t" in st_ and st_["prop_t"] is not None and not np.isnan(st_["prop_t"]) and abs(st_["prop_t"]) > 1.96:
            notes.append(f"istotny błąd proporcjonalny (nachylenie różnicy względem średniej {st_['prop_slope']:+.3f}, t = {st_['prop_t']:.1f})")
        if notes:
            items.append({"kind": "info", "title": "Charakter rozbieżności", "text": "; ".join(notes) + "."})
        rp = ctx.get("rpd_high_share")
        if rp is not None:
            items.append({"kind": "info", "title": "RPD (względna różnica procentowa)",
                          "text": f"W grupie wysokiego zróżnicowania (RPD > 30%) znajduje się **{rp:.0f}%** punktów. "
                                  "Wysoki klasyczny współczynnik zmienności przy małych głębokościach wynika z „efektu małego "
                                  "mianownika”, dlatego do oceny płytkich kolein stosuje się RPD i różnice bezwzględne."})

    # --- 5. Klasy DSN
    for lab, cs in (("co 1 m", ctx.get("cls_point")), ("na odcinkach 50 m (Hm)", ctx.get("cls_sec"))):
        if cs and cs.get("n", 0) > 0:
            kw = cs.get("kappa_w", np.nan)
            kind = "ok" if kw >= 0.75 else ("warn" if kw >= 0.5 else "bad")
            trend = ("LCMS częściej zalicza do **gorszej** klasy" if cs["lcms_higher"] > cs["lcms_lower"] + 2 else
                     ("LCMS częściej zalicza do **lepszej** klasy" if cs["lcms_lower"] > cs["lcms_higher"] + 2 else
                      "rozbieżności klas są symetryczne"))
            txt = (f"Zgodność klas A–D {lab}: **{cs['exact']:.0f}%** identycznych, {cs['within1']:.0f}% w granicach ±1 klasy; "
                   f"kappa ważona = **{cs['kappa_w']:.2f}**. LCMS gorsza klasa: {cs['lcms_higher']:.0f}%, lepsza: "
                   f"{cs['lcms_lower']:.0f}% – {trend}.")
            items.append({"kind": kind, "title": f"Klasy DSN {lab}", "text": txt})

    # --- 6. Odcinki do weryfikacji
    sec = ctx.get("sections")
    if sec is not None and len(sec) >= 3:
        bad = sec.reindex(sec["Bias [mm]"].abs().sort_values(ascending=False).index).head(3)
        lines = [f"{r['Odcinek']} (bias {r['Bias [mm]']:+.1f} mm, r = {r['r']:.2f})" for _, r in bad.iterrows()]
        lowr = sec.dropna(subset=["r"]).sort_values("r").head(3)
        lines2 = [f"{r['Odcinek']} (r = {r['r']:.2f})" for _, r in lowr.iterrows()]
        txt = ("Największy błąd systematyczny: " + "; ".join(lines) + ". Najniższa korelacja: " + "; ".join(lines2) + ".")
        items.append({"kind": "info", "title": "Odcinki do weryfikacji", "text": txt})

    # --- 7. Jakość danych
    qn = ctx.get("quality_notes") or []
    if qn:
        items.append({"kind": "info", "title": "Uwagi o jakości danych", "text": "; ".join(qn) + "."})

    # --- 8. Rekomendacje i werdykt
    grade, gk = agreement_grade(smax if smax else sl, tol)
    recs = []
    base = smax if smax and smax.get("n", 0) >= 3 else sl
    if base and base.get("n", 0) >= 3:
        sig = not (base["bias_ci_lo"] <= 0 <= base["bias_ci_hi"])
        if sig and abs(base["bias"]) >= 0.5:
            recs.append(f"skorygować/ocenić kalibrację – istnieje stały błąd systematyczny rzędu {base['bias']:+.1f} mm")
        if base["pearson"] < 0.7 and not sig:
            recs.append("niska korelacja przy braku biasu wskazuje na niedopasowanie przestrzenne (pikietaż, dryf, przesunięcie "
                        "poprzeczne pasa) lub różne definicje parametru – sprawdź synchronizację i geometrię poprzeczną")
        if base.get("hetero_r", 0) and base["hetero_r"] > 0.2:
            recs.append("raportować zgodność w przedziałach głębokości, nie jedną wartością dla całej trasy")
        fit = sync.get("fit") if sync else None
        if fit and fit.get("n", 0) >= 5 and abs(fit["drift_per_km"]) >= 0.3 and fit.get("r2", 0) >= 0.5:
            recs.append("zastosować korektę dryfu pikietażu (zakładka Synchronizacja)")
    if not recs:
        recs.append("brak zastrzeżeń krytycznych – wyniki obu systemów można traktować zamiennie w granicach podanych tolerancji")
    items.append({"kind": gk, "title": "Podsumowanie i rekomendacje",
                  "text": f"Ogólna zgodność RSP–LCMS (wg progów CCC/tolerancji): **{grade}**. Zalecenia: " + "; ".join(recs) + "."})
    return {"grade": grade, "kind": gk, "items": items, "markdown": conclusions_markdown(items, grade)}


_ICON = {"ok": "🟢", "info": "🔵", "warn": "🟡", "bad": "🔴"}


def conclusions_markdown(items, grade=None):
    lines = ["## Wnioski z analizy RSP vs LCMS", ""]
    if grade:
        lines += [f"**Ogólna zgodność: {grade}**", ""]
    for it in items:
        lines.append(f"- {_ICON.get(it['kind'], '•')} **{it['title']}.** {it['text']}")
    return "\n".join(lines)


# ==========================================================================
# 9. WYKRESY (Plotly)
# ==========================================================================
def km_ticks(lo, hi, n=8):
    """Podziałka osi pikietażu w zapisie km+m."""
    span = max(hi - lo, 1.0)
    raw = span / n
    steps = [1, 2, 5, 10, 20, 25, 50, 100, 200, 250, 500, 1000, 2000, 2500, 5000, 10000, 20000, 50000]
    step = next((s for s in steps if s >= raw), steps[-1])
    first = math.ceil(lo / step) * step
    vals = list(np.arange(first, hi + 1e-9, step))
    return vals, [km_str(v) for v in vals]


def apply_km_axis(fig, lo, hi, row=None, col=None):
    vals, txt = km_ticks(lo, hi)
    kw = dict(tickmode="array", tickvals=vals, ticktext=txt, range=[lo, hi])
    if row is not None:
        fig.update_xaxes(row=row, col=col, **kw)
    else:
        fig.update_xaxes(**kw)


def _class_bands(fig, ymax, opacity=0.07):
    edges = [0.0, DSN_THRESHOLDS[0], DSN_THRESHOLDS[1], DSN_THRESHOLDS[2], max(ymax, DSN_THRESHOLDS[2] + 5)]
    for cls, y0, y1 in zip(DSN_CLASSES, edges[:-1], edges[1:]):
        fig.add_hrect(y0=y0, y1=y1, fillcolor=DSN_COLORS[cls], opacity=opacity, line_width=0, layer="below")


def fig_longitudinal(df, traces, title, view=None, selected=None, bands=True, max_points=4000,
                     ytitle="Głębokość koleiny [mm]", height=430, uirev=None, extra_vlines=None):
    """Przebieg wzdłużny: traces = [dict(col, name, color, width, dash)]; zakres `view` w metrach."""
    d = df
    if view is not None:
        d = df[(df["Pikietaz"] >= view[0]) & (df["Pikietaz"] <= view[1])]
    fig = go.Figure()
    x = d["Pikietaz"].values
    ymax = 10.0
    for t in traces:
        if t["col"] not in d.columns:
            continue
        y = d[t["col"]].values.astype(float)
        if len(y) and not np.all(np.isnan(y)):
            ymax = max(ymax, float(np.nanmax(y)))
        hc = t.get("hover_col")
        if hc and hc in d.columns:
            xd, (yd, cd) = decimate_minmax(x, [y, d[hc].values.astype(float)], max_points)
            ht = ("%{x:.0f} m: %{y:.1f} mm<br>przesunięcie LCMS względem RSP: <b>%{customdata:+.1f} m</b>"
                  "<extra>" + t["name"] + "</extra>")
            extra = dict(customdata=cd)
        else:
            xd, (yd,) = decimate_minmax(x, [y], max_points)
            ht = "%{x:.0f} m: %{y:.1f} mm<extra>" + t["name"] + "</extra>"
            extra = {}
        fig.add_trace(go.Scattergl(
            x=xd, y=yd, mode="lines", name=t["name"], connectgaps=False,
            line=dict(color=t.get("color", "gray"), width=t.get("width", 1.6), dash=t.get("dash", "solid")),
            hovertemplate=ht, **extra))
    if bands:
        _class_bands(fig, ymax * 1.05)
    if selected is not None:
        fig.add_vline(x=selected, line_width=2, line_color="red", annotation_text=km_str(selected))
    for xv, lab in (extra_vlines or []):
        fig.add_vline(x=xv, line_width=1, line_dash="dot", line_color="gray", annotation_text=lab)
    lo, hi = (view if view is not None else (float(x.min()) if len(x) else 0.0, float(x.max()) if len(x) else 1.0))
    fig.update_layout(title=title, xaxis_title="Pikietaż (km+m)", yaxis_title=ytitle, height=height,
                      hovermode="x unified", margin=dict(l=50, r=20, t=50, b=40),
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                      yaxis=dict(range=[0, ymax * 1.08]))
    if uirev:
        fig.update_layout(uirevision=uirev)
    apply_km_axis(fig, lo, hi)
    return fig


def fig_diff_along(df, a_col="RSP_Max", b_col="LCMS_Max", view=None, roll=101, title="Różnica LCMS − RSP wzdłuż trasy",
                   max_points=4000, tol=2.0, uirev=None):
    d = df if view is None else df[(df["Pikietaz"] >= view[0]) & (df["Pikietaz"] <= view[1])]
    diff = (d[b_col] - d[a_col]).values.astype(float)
    x = d["Pikietaz"].values
    sm = nan_moving_average(diff, roll)
    fig = go.Figure()
    xd, (dd,) = decimate_minmax(x, [diff], max_points)
    fig.add_trace(go.Scattergl(x=xd, y=dd, mode="lines", name="Δ (co 1 m)", line=dict(color="rgba(123,44,191,0.35)", width=1),
                               hovertemplate="%{x:.0f} m: %{y:+.1f} mm<extra>Δ</extra>"))
    xs, (ss,) = decimate_minmax(x, [sm], max_points)
    fig.add_trace(go.Scattergl(x=xs, y=ss, mode="lines", name=f"średnia krocząca {roll} m", line=dict(color=C_DIFF, width=2.5)))
    fig.add_hline(y=0, line_width=1, line_color="gray")
    fig.add_hrect(y0=-tol, y1=tol, fillcolor="green", opacity=0.07, line_width=0, layer="below")
    lo, hi = (view if view is not None else (float(x.min()) if len(x) else 0.0, float(x.max()) if len(x) else 1.0))
    fig.update_layout(title=title, xaxis_title="Pikietaż (km+m)", yaxis_title="LCMS − RSP [mm]", height=360,
                      hovermode="x unified", margin=dict(l=50, r=20, t=50, b=40),
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    if uirev:
        fig.update_layout(uirevision=uirev)
    apply_km_axis(fig, lo, hi)
    return fig


def build_bland_altman_fig(rsp_v, lcms_v, title="PRAWA KOLEINA", max_points=15000):
    """Wykres Blanda–Altmana w stylu GDDKiA (jak dotychczas) + linia błędu proporcjonalnego."""
    means, diffs, mean_diff, sd_diff, loa_upper, loa_lower = bland_altman_stats(rsp_v, lcms_v)
    fig = go.Figure()
    if len(means) > max_points:
        sel = np.random.default_rng(1).choice(len(means), max_points, replace=False)
        mx, my = means[sel], diffs[sel]
    else:
        mx, my = means, diffs
    big = len(mx) > 3000
    fig.add_trace(go.Scattergl(
        x=mx, y=my, mode="markers", name="Pomiary",
        marker=dict(size=4 if big else 7, color="#ffffcc" if not big else "rgba(70,70,70,0.35)",
                    line=dict(width=0 if big else 1, color="#333333")),
        hovertemplate="Średnia: %{x:.2f} mm<br>Różnica (LCMS-RSP): %{y:.2f} mm<extra></extra>"))
    if len(means) > 10 and np.std(means) > 0:
        sl, ic = np.polyfit(means, diffs, 1)
        xx = np.array([means.min(), means.max()])
        fig.add_trace(go.Scatter(x=xx, y=ic + sl * xx, mode="lines", name=f"trend różnicy ({sl:+.3f}·średnia)",
                                 line=dict(color="gray", dash="dot", width=2)))
    fig.add_hline(y=0, line_width=1, line_dash="solid", line_color="lightgray")
    fig.add_hline(y=mean_diff, line_width=3, line_color="#0066cc", annotation_text=f"Średnia: {mean_diff:+.2f} mm",
                  annotation_position="right", annotation_font=dict(color="#0066cc", size=13, family="sans-serif"))
    fig.add_hline(y=loa_upper, line_width=2, line_dash="dash", line_color="red",
                  annotation_text=f"+1.96 SD: {loa_upper:+.2f} mm", annotation_position="right",
                  annotation_font=dict(color="red", size=12))
    fig.add_hline(y=loa_lower, line_width=2, line_dash="dash", line_color="red",
                  annotation_text=f"-1.96 SD: {loa_lower:+.2f} mm", annotation_position="right",
                  annotation_font=dict(color="red", size=12))
    fig.update_layout(title=dict(text=f"<b>{title}</b>", x=0.5, font=dict(size=18)),
                      xaxis_title="ŚREDNIA (LCMS, RSP) [mm]", yaxis_title="RÓŻNICA (LCMS - RSP) [mm]",
                      height=480, template="plotly_white", hovermode="closest",
                      margin=dict(l=50, r=110, t=60, b=50))
    return fig, mean_diff, sd_diff, loa_upper, loa_lower


def fig_scatter_xy(a, b, title="LCMS vs RSP", max_points=12000, st_=None):
    """Wykres rozrzutu z gęstością, linią 1:1, regresją MNK i ortogonalną."""
    a, b = clean_pair(a, b)
    fig = go.Figure()
    if len(a) == 0:
        return fig
    top = float(max(a.max(), b.max())) * 1.05
    fig.add_trace(go.Histogram2d(x=a, y=b, nbinsx=70, nbinsy=70, colorscale="Blues", name="gęstość",
                                 colorbar=dict(title=dict(text="liczba")), hovertemplate="RSP %{x:.1f}<br>LCMS %{y:.1f}<br>n=%{z}<extra></extra>"))
    fig.add_trace(go.Scatter(x=[0, top], y=[0, top], mode="lines", name="1:1", line=dict(color="black", dash="dash")))
    if st_ and "slope" in st_:
        xx = np.array([0.0, top])
        fig.add_trace(go.Scatter(x=xx, y=st_["intercept"] + st_["slope"] * xx, mode="lines",
                                 name=f"MNK: {st_['intercept']:.2f} + {st_['slope']:.2f}·x", line=dict(color="crimson", width=2)))
        if not np.isnan(st_.get("deming_slope", np.nan)):
            fig.add_trace(go.Scatter(x=xx, y=st_["deming_intercept"] + st_["deming_slope"] * xx, mode="lines",
                                     name=f"ortogonalna: {st_['deming_intercept']:.2f} + {st_['deming_slope']:.2f}·x",
                                     line=dict(color="darkorange", width=2, dash="dot")))
    fig.update_layout(title=title, xaxis_title="RSP [mm]", yaxis_title="LCMS [mm]", height=480, template="plotly_white",
                      xaxis=dict(range=[0, top]), yaxis=dict(range=[0, top]),
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    return fig


def fig_diff_hist(a, b, tol=2.0, title="Rozkład różnic LCMS − RSP"):
    a, b = clean_pair(a, b)
    d = b - a
    fig = go.Figure()
    if len(d) == 0:
        return fig
    fig.add_trace(go.Histogram(x=d, nbinsx=80, marker=dict(color="rgba(123,44,191,0.65)"), name="różnice"))
    mu, sd = float(d.mean()), float(d.std(ddof=1)) if len(d) > 1 else 0.0
    fig.add_vline(x=0, line_color="gray", line_width=1)
    fig.add_vline(x=mu, line_color="#0066cc", line_width=3, annotation_text=f"bias {mu:+.2f}")
    fig.add_vline(x=mu + 1.96 * sd, line_color="red", line_dash="dash")
    fig.add_vline(x=mu - 1.96 * sd, line_color="red", line_dash="dash")
    fig.add_vrect(x0=-tol, x1=tol, fillcolor="green", opacity=0.08, line_width=0, layer="below")
    fig.update_layout(title=title, xaxis_title="LCMS − RSP [mm]", yaxis_title="liczba punktów", height=380,
                      template="plotly_white", showlegend=False)
    return fig


def fig_ecdf(a, b, tol=2.0, title="Dystrybuanta |LCMS − RSP|"):
    a, b = clean_pair(a, b)
    ad = np.sort(np.abs(b - a))
    fig = go.Figure()
    if len(ad) == 0:
        return fig
    p = np.arange(1, len(ad) + 1) / len(ad) * 100.0
    step = max(1, len(ad) // 3000)
    fig.add_trace(go.Scatter(x=ad[::step], y=p[::step], mode="lines", line=dict(color=C_DIFF, width=2.5), name="ECDF"))
    for t in sorted({1, 2, 3, 5, tol}):
        share = float(np.mean(ad <= t) * 100)
        fig.add_vline(x=t, line_dash="dot", line_color="gray", annotation_text=f"{t:g} mm: {share:.0f}%")
    fig.update_layout(title=title, xaxis_title="|LCMS − RSP| [mm]", yaxis_title="% punktów z różnicą ≤ x", height=380,
                      template="plotly_white", showlegend=False, yaxis=dict(range=[0, 100.5]))
    return fig


def fig_box_by_depth(a, b, edges=(0, 5, 10, 15, 20, 30, 1e9), title="Różnice LCMS − RSP w przedziałach głębokości"):
    a, b = clean_pair(a, b)
    m = (a + b) / 2.0
    d = b - a
    fig = go.Figure()
    for lo, hi in zip(edges[:-1], edges[1:]):
        sel = (m >= lo) & (m < hi)
        if sel.sum() < 3:
            continue
        lab = f"{lo:g}–{hi:g}" if hi < 1e8 else f"≥{lo:g}"
        fig.add_trace(go.Box(y=d[sel], name=f"{lab} mm (n={int(sel.sum())})", boxpoints=False, marker_color=C_DIFF))
    fig.add_hline(y=0, line_color="gray", line_width=1)
    fig.update_layout(title=title, yaxis_title="LCMS − RSP [mm]", height=400, template="plotly_white", showlegend=False)
    return fig


def fig_confusion(M, title="Macierz klas DSN (wiersze: RSP, kolumny: LCMS)"):
    M = np.asarray(M)
    tot = M.sum()
    pct = M / tot * 100 if tot else M * 0.0
    text = [[f"{M[i, j]:,}<br>{pct[i, j]:.1f}%".replace(",", " ") for j in range(4)] for i in range(4)]
    fig = go.Figure(go.Heatmap(z=M, x=[f"LCMS {c}" for c in DSN_CLASSES], y=[f"RSP {c}" for c in DSN_CLASSES],
                               text=text, texttemplate="%{text}", colorscale="Blues", showscale=False,
                               hovertemplate="%{y} → %{x}: %{z}<extra></extra>"))
    fig.update_layout(title=title, height=380, yaxis=dict(autorange="reversed"), margin=dict(l=60, r=20, t=50, b=40))
    return fig


def fig_rolling(rdf, title="Zgodność w oknach przesuwnych"):
    from plotly.subplots import make_subplots
    fig = make_subplots(rows=3, cols=1, shared_xaxes=True, vertical_spacing=0.06,
                        subplot_titles=("Średnia koleina (RSP / LCMS) [mm]", "Bias i RMSE [mm]", "Korelacja r"))
    if len(rdf):
        fig.add_trace(go.Scatter(x=rdf["Pikietaz"], y=rdf["RSP"], name="RSP", line=dict(color=C_RSP)), row=1, col=1)
        fig.add_trace(go.Scatter(x=rdf["Pikietaz"], y=rdf["LCMS"], name="LCMS", line=dict(color=C_LCMS)), row=1, col=1)
        fig.add_trace(go.Scatter(x=rdf["Pikietaz"], y=rdf["Bias"], name="Bias (LCMS−RSP)", line=dict(color=C_DIFF)), row=2, col=1)
        fig.add_trace(go.Scatter(x=rdf["Pikietaz"], y=rdf["RMSE"], name="RMSE", line=dict(color="crimson", dash="dot")), row=2, col=1)
        fig.add_trace(go.Scatter(x=rdf["Pikietaz"], y=rdf["r"], name="r", line=dict(color="black")), row=3, col=1)
        fig.add_hline(y=0, line_color="gray", line_width=1, row=2, col=1)
        fig.update_yaxes(range=[-1, 1], row=3, col=1)
        apply_km_axis(fig, float(rdf["Pikietaz"].min()), float(rdf["Pikietaz"].max()), row=3, col=1)
    fig.update_layout(title=title, height=620, hovermode="x unified", margin=dict(l=50, r=20, t=70, b=40),
                      legend=dict(orientation="h", yanchor="bottom", y=1.04, xanchor="right", x=1))
    return fig


def fig_section_bars(sec, metric="Bias [mm]", title="Zgodność na odcinkach"):
    colors = ["#dc3545" if v > 0 else "#0066cc" for v in sec[metric].values] if metric == "Bias [mm]" else C_DIFF
    mid = (sec["Od [m]"].values + sec["Do [m]"].values) / 2.0
    width = (sec["Do [m]"].values - sec["Od [m]"].values + 1) * 0.9
    fig = go.Figure(go.Bar(x=mid, y=sec[metric], width=width, marker_color=colors, customdata=sec["Odcinek"],
                           hovertemplate="%{customdata}<br>" + metric + ": %{y:.2f}<extra></extra>"))
    fig.update_layout(title=title, yaxis_title=metric, xaxis_title="Pikietaż (km+m)", height=360, margin=dict(l=50, r=20, t=50, b=40))
    if len(sec):
        apply_km_axis(fig, float(sec["Od [m]"].min()), float(sec["Do [m]"].max()))
    return fig


# ---------- ocena odcinkowa DSN ----------
def make_50m_bar_chart(df_sec_dsn, param_prefix="P_KOL", title="Prawa Koleina (P_KOL)", length=50.0):
    """Wykres słupkowy oceny Hm dla odcinków (jak dotychczas, oś w pikietażu)."""
    fig = go.Figure()
    mid = (df_sec_dsn["Od [m]"].values + df_sec_dsn["Do [m]"].values) / 2.0
    wid = np.full(len(mid), length * (0.9 if "LCMS %s Hm" % param_prefix not in df_sec_dsn.columns else 0.45))
    col_rsp_hm = f"RSP {param_prefix} Hm"
    col_rsp_klasa = f"RSP {param_prefix} Klasa"
    fig.add_trace(go.Bar(
        x=mid - (wid / 2 if "LCMS %s Hm" % param_prefix in df_sec_dsn.columns else 0), y=df_sec_dsn[col_rsp_hm], width=wid,
        name=f"RSP {title} (Hm)", marker_color=[get_dsn_color(k) for k in df_sec_dsn[col_rsp_klasa]],
        customdata=df_sec_dsn["Odcinek"], hovertemplate="%{customdata}<br>RSP Hm: %{y} mm<extra></extra>",
        text=None))
    col_l_hm = f"LCMS {param_prefix} Hm"
    col_l_k = f"LCMS {param_prefix} Klasa"
    if col_l_hm in df_sec_dsn.columns:
        fig.add_trace(go.Bar(
            x=mid + wid / 2, y=df_sec_dsn[col_l_hm], width=wid, name=f"LCMS {title} (Hm)",
            marker_color=[get_dsn_color(k) for k in df_sec_dsn[col_l_k]], opacity=0.55,
            marker_pattern_shape="/", customdata=df_sec_dsn["Odcinek"],
            hovertemplate="%{customdata}<br>LCMS Hm: %{y} mm<extra></extra>"))
    fig.add_hline(y=11, line_dash="dash", line_color="#ffc107", annotation_text="Próg B (11mm)")
    fig.add_hline(y=21, line_dash="dash", line_color="#fd7e14", annotation_text="Próg C (21mm)")
    fig.add_hline(y=31, line_dash="dash", line_color="#dc3545", annotation_text="Próg D (31mm)")
    fig.update_layout(title=f"Ocena Hm dla odcinków {length:g} m – {title} (pełne słupki: RSP, kreskowane: LCMS)",
                      xaxis_title="Pikietaż (km+m)", yaxis_title="Ocena Hm [mm]", barmode="overlay", height=480,
                      margin=dict(l=50, r=20, t=60, b=40))
    if len(mid):
        apply_km_axis(fig, float(df_sec_dsn["Od [m]"].min()), float(df_sec_dsn["Do [m]"].max()))
    return fig


def fig_class_strip(df_sec, param_prefix="MAX_KOL", title="Pasek klas DSN (RSP / LCMS)"):
    """Dwa paski kolorów (RSP, LCMS) z klasą A–D wzdłuż trasy."""
    cols = [("RSP", f"RSP {param_prefix} Klasa"), ("LCMS", f"LCMS {param_prefix} Klasa")]
    cols = [(n, c) for n, c in cols if c in df_sec.columns]
    cmap = {"A": 0, "B": 1, "C": 2, "D": 3}
    z = [[cmap.get(k, np.nan) for k in df_sec[c]] for _, c in cols]
    mid = (df_sec["Od [m]"].values + df_sec["Do [m]"].values) / 2.0
    scale = [[0.0, DSN_COLORS["A"]], [0.25, DSN_COLORS["A"]], [0.25, DSN_COLORS["B"]], [0.5, DSN_COLORS["B"]],
             [0.5, DSN_COLORS["C"]], [0.75, DSN_COLORS["C"]], [0.75, DSN_COLORS["D"]], [1.0, DSN_COLORS["D"]]]
    fig = go.Figure(go.Heatmap(z=z, x=mid, y=[n for n, _ in cols], zmin=0, zmax=3, colorscale=scale, showscale=False,
                               customdata=[list(df_sec["Odcinek"])] * len(cols),
                               hovertemplate="%{customdata}<br>%{y}: klasa wg koloru<extra></extra>"))
    fig.update_layout(title=title, height=200, margin=dict(l=60, r=20, t=50, b=40), xaxis_title="Pikietaż (km+m)")
    if len(mid):
        apply_km_axis(fig, float(df_sec["Od [m]"].min()), float(df_sec["Do [m]"].max()))
    return fig


# ---------- synchronizacja ----------
def fig_corr_curve(lags, corr, best, title="Korelacja krzyżowa w funkcji przesunięcia"):
    fig = go.Figure(go.Scatter(x=lags, y=corr, mode="lines", line=dict(color=C_DIFF, width=2)))
    if best is not None and not np.isnan(best):
        fig.add_vline(x=best, line_color="red", annotation_text=f"max przy {best:+.0f} m")
    fig.update_layout(title=title, xaxis_title="Przesunięcie LCMS względem RSP [m]", yaxis_title="Korelacja r", height=340,
                      template="plotly_white")
    return fig


def fig_lag_map(df_lags, fit=None, current=None, title="Lokalne przesunięcia (okna przesuwne)"):
    fig = go.Figure()
    if df_lags is not None and len(df_lags):
        fig.add_trace(go.Scatter(
            x=df_lags["Pikietaz"], y=df_lags["Lag [m]"], mode="markers", name="okno",
            marker=dict(size=9, color=df_lags["Korelacja"], colorscale="Viridis", cmin=0, cmax=max(0.5, float(df_lags["Korelacja"].max())),
                        showscale=True, colorbar=dict(title=dict(text="r")), line=dict(width=0.5, color="black")),
            hovertemplate="%{x:.0f} m<br>lag %{y:+.0f} m<br>r=%{marker.color:.2f}<extra></extra>"))
        if fit:
            xx = np.array([df_lags["Pikietaz"].min(), df_lags["Pikietaz"].max()])
            yy = fit["offset_ref"] + fit["drift_per_km"] * (xx - fit["ref"]) / 1000.0
            fig.add_trace(go.Scatter(x=xx, y=yy, mode="lines", name=f"dopasowanie: {fit['drift_per_km']:+.2f} m/km",
                                     line=dict(color="red", width=2)))
    if current is not None:
        fig.add_trace(go.Scatter(x=current[0], y=current[1], mode="lines", name="bieżące przesunięcie",
                                 line=dict(color="black", dash="dot", width=2)))
    fig.update_layout(title=title, xaxis_title="Pikietaż RSP (km+m)", yaxis_title="Przesunięcie LCMS − RSP [m]", height=420,
                      template="plotly_white", legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    if df_lags is not None and len(df_lags):
        apply_km_axis(fig, float(df_lags["Pikietaz"].min()), float(df_lags["Pikietaz"].max()))
    return fig


# ---------- dane i jakość ----------
def fig_coverage(items, overlap=None, title="Zakres pikietażu plików"):
    """items: lista (nazwa, typ, od, do). Poziome paski + zaznaczony zakres wspólny."""
    fig = go.Figure()
    colors = {TYPE_RSP_RAW: C_RSP, TYPE_RSP_KOLC: "#4169e1", TYPE_RSP_GPS: "#87a5f0", TYPE_LCMS_PROFILE: C_LCMS,
              TYPE_LCMS_GPS: "#8fd18f"}
    for name, typ, lo, hi in items:
        fig.add_trace(go.Bar(y=[name], x=[hi - lo], base=[lo], orientation="h", marker_color=colors.get(typ, "gray"),
                             name=TYPE_LABELS.get(typ, typ), showlegend=False,
                             hovertemplate=f"{name}<br>{km_str(lo)} – {km_str(hi)}<extra></extra>"))
    if overlap is not None and overlap[1] > overlap[0]:
        fig.add_vrect(x0=overlap[0], x1=overlap[1], fillcolor="orange", opacity=0.12, line_width=0,
                      annotation_text="zakres wspólny", annotation_position="top left")
    fig.update_layout(title=title, barmode="overlay", height=max(220, 60 + 38 * len(items)), xaxis_title="Pikietaż (km+m)",
                      margin=dict(l=10, r=20, t=50, b=40), yaxis=dict(autorange="reversed"))
    if items:
        apply_km_axis(fig, min(i[2] for i in items), max(i[3] for i in items))
    return fig


def fig_hist_depths(rsp_v, lcms_v, title="Rozkład głębokości kolein"):
    fig = go.Figure()
    for v, nm, col in ((rsp_v, "RSP", C_RSP), (lcms_v, "LCMS", C_LCMS)):
        v = np.asarray(v, dtype=float)
        v = v[~np.isnan(v)]
        if len(v):
            fig.add_trace(go.Histogram(x=v, nbinsx=80, name=nm, opacity=0.55, marker=dict(color=col), histnorm="percent"))
    for t, c in zip(DSN_THRESHOLDS, ("#ffc107", "#fd7e14", "#dc3545")):
        fig.add_vline(x=t, line_dash="dash", line_color=c)
    fig.update_layout(title=title, barmode="overlay", xaxis_title="Głębokość koleiny [mm]", yaxis_title="% punktów",
                      height=360, template="plotly_white", legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    return fig


# ---------- przekrój poprzeczny (jak dotychczas) ----------
def fig_cross_section(rs, lc, title, show_rsp_hull=True, show_lcms_hull=True, label_lcms=None,
                      x_range=(-1.85, 1.85), profile_name="LCMS"):
    """Przekrój poprzeczny RSP + LCMS. rs / lc – wyniki rsp_cross_section / lcms_cross_section (lub None)."""
    fig = go.Figure()
    y_label = "Wysokość względna [mm]"
    reverse = False
    if rs is not None:
        y_label = rs["y_label"]
        reverse = rs["reverse_y"]
        fig.add_trace(go.Scatter(
            x=rs["x"], y=rs["y"], mode="lines+markers", name="Profil RSP (Spline)",
            line=dict(shape="spline", smoothing=1.2, color="royalblue", width=3), marker=dict(size=7, color="navy"),
            hovertemplate="Poz X: %{x} m<br>Wys Y: %{y:.1f} mm<extra></extra>"))
        if show_rsp_hull:
            fig.add_trace(go.Scatter(x=rs["x_left"], y=rs["lata_l"], mode="lines", name="Łata RSP Lewa",
                                     line=dict(color="crimson", width=2)))
            fig.add_trace(go.Scatter(x=rs["x_right"], y=rs["lata_r"], mode="lines", name="Łata RSP Prawa",
                                     line=dict(color="darkorange", width=2)))
            x_kl = rs["x"][rs["idx_l"]]
            fig.add_trace(go.Scatter(x=[x_kl, x_kl], y=[rs["y_l_lata"], rs["y_l_prof"]], mode="lines+markers",
                                     name=f"Głębokość RSP Lewa ({rs['rut_l']:.1f} mm)", line=dict(color="crimson", width=3),
                                     marker=dict(size=6, symbol="triangle-down")))
            fig.add_annotation(x=x_kl, y=(rs["y_l_lata"] + rs["y_l_prof"]) / 2, text=f" <b>RSP: {rs['rut_l']:.1f} mm</b>",
                               showarrow=False, font=dict(color="crimson", size=12), bgcolor="rgba(255, 255, 255, 0.95)",
                               bordercolor="crimson", borderwidth=1, xanchor="right")
            x_kp = rs["x"][rs["idx_r"]]
            fig.add_trace(go.Scatter(x=[x_kp, x_kp], y=[rs["y_r_lata"], rs["y_r_prof"]], mode="lines+markers",
                                     name=f"Głębokość RSP Prawa ({rs['rut_r']:.1f} mm)", line=dict(color="darkorange", width=3),
                                     marker=dict(size=6, symbol="triangle-down")))
            fig.add_annotation(x=x_kp, y=(rs["y_r_lata"] + rs["y_r_prof"]) / 2, text=f" <b>RSP: {rs['rut_r']:.1f} mm</b>",
                               showarrow=False, font=dict(color="darkorange", size=12), bgcolor="rgba(255, 255, 255, 0.95)",
                               bordercolor="darkorange", borderwidth=1, xanchor="right")
    if lc is not None:
        fig.add_trace(go.Scatter(x=lc["x"], y=lc["z_plot"], mode="lines", name=profile_name,
                                 line=dict(color="forestgreen", width=2)))
        if show_lcms_hull:
            fig.add_trace(go.Scatter(x=lc["x_l"], y=lc["lata_l"], mode="lines", name="Łata LCMS Lewa",
                                     line=dict(color="limegreen", width=2)))
            fig.add_trace(go.Scatter(x=lc["x_r"], y=lc["lata_r"], mode="lines", name="Łata LCMS Prawa",
                                     line=dict(color="lime", width=2)))
            vl, vr = (label_lcms or (None, None))
            vl = lc["rut_l_calc"] if vl is None or np.isnan(vl) else vl
            vr = lc["rut_r_calc"] if vr is None or np.isnan(vr) else vr
            fig.add_trace(go.Scatter(x=[lc["x_kl"], lc["x_kl"]], y=[lc["y_l_lata"], lc["y_l_prof"]], mode="lines+markers",
                                     name=f"Głębokość LCMS Lewa ({vl:.1f} mm)", line=dict(color="forestgreen", width=2.5),
                                     marker=dict(size=5, symbol="triangle-down")))
            fig.add_annotation(x=lc["x_kl"], y=(lc["y_l_lata"] + lc["y_l_prof"]) / 2, text=f"<b>LCMS: {vl:.1f} mm</b> ",
                               showarrow=False, font=dict(color="forestgreen", size=12), bgcolor="rgba(235, 255, 235, 0.95)",
                               bordercolor="forestgreen", borderwidth=1, xanchor="left")
            fig.add_trace(go.Scatter(x=[lc["x_kp"], lc["x_kp"]], y=[lc["y_r_lata"], lc["y_r_prof"]], mode="lines+markers",
                                     name=f"Głębokość LCMS Prawa ({vr:.1f} mm)", line=dict(color="forestgreen", width=2.5),
                                     marker=dict(size=5, symbol="triangle-down")))
            fig.add_annotation(x=lc["x_kp"], y=(lc["y_r_lata"] + lc["y_r_prof"]) / 2, text=f"<b>LCMS: {vr:.1f} mm</b> ",
                               showarrow=False, font=dict(color="forestgreen", size=12), bgcolor="rgba(235, 255, 235, 0.95)",
                               bordercolor="forestgreen", borderwidth=1, xanchor="left")
    fig.add_vline(x=0, line_width=1, line_dash="dot", line_color="gray", annotation_text="Oś L11 (0.0 m)")
    fig.update_layout(title=title, xaxis_title="Pozycja czujnika na belce [m]", yaxis_title=y_label, height=580,
                      hovermode="x unified", xaxis=dict(range=list(x_range), dtick=0.2),
                      yaxis=dict(autorange="reversed" if reverse else True),
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    return fig


def fig_profile_heatmap(chains, X, z_title, x_grid, title, colorscale="RdBu_r", zrange=None, selected=None):
    """Mapa cieplna profili poprzecznych wzdłuż odcinka (oś Y: pikietaż, oś X: pozycja poprzeczna)."""
    Z = np.asarray(X, dtype=float)
    if zrange is None:
        lim = float(np.nanpercentile(np.abs(Z), 98)) if np.any(~np.isnan(Z)) else 10.0
        zrange = (-lim, lim)
    fig = go.Figure(go.Heatmap(z=Z, x=x_grid, y=chains, colorscale=colorscale, zmin=zrange[0], zmax=zrange[1],
                               colorbar=dict(title=dict(text=z_title)),
                               hovertemplate="x=%{x:.2f} m<br>%{y:.0f} m<br>%{z:.1f} mm<extra></extra>"))
    if selected is not None:
        fig.add_hline(y=selected, line_color="black", line_dash="dot")
    fig.update_layout(title=title, xaxis_title="Pozycja poprzeczna [m]", yaxis_title="Pikietaż [m]", height=520,
                      margin=dict(l=60, r=20, t=50, b=40))
    return fig


def fig_surface(chains, Z, x_grid, title):
    fig = go.Figure(go.Surface(z=np.asarray(Z, dtype=float), x=x_grid, y=chains, colorscale="RdBu_r",
                               colorbar=dict(title=dict(text="mm"))))
    fig.update_layout(title=title, height=560, margin=dict(l=0, r=0, t=50, b=0),
                      scene=dict(xaxis_title="x [m]", yaxis_title="pikietaż [m]", zaxis_title="z [mm]",
                                 aspectmode="manual", aspectratio=dict(x=1.2, y=2.0, z=0.4)))
    return fig


# ---------- mapa ----------
def make_map_trace(lat, lon, mode="lines", marker=None, line=None, name="", text=None):
    """Bezpieczny kreator warstw mapy Plotly (Scattermap / Scattermapbox / Scattergeo)."""
    kwargs = {"lat": lat, "lon": lon, "mode": mode, "name": name}
    if marker is not None:
        kwargs["marker"] = marker
    if line is not None:
        clean_line = dict(line)
        clean_line.pop("dash", None)
        kwargs["line"] = clean_line
    if text is not None:
        kwargs["text"] = text
        kwargs["hoverinfo"] = "text"
    if hasattr(go, "Scattermap"):
        return go.Scattermap(**kwargs)
    if hasattr(go, "Scattermapbox"):
        return go.Scattermapbox(**kwargs)
    return go.Scattergeo(**kwargs)


def map_zoom_for(lat, lon):
    lat = np.asarray(lat, dtype=float)
    lon = np.asarray(lon, dtype=float)
    span = max(float(np.nanmax(lon) - np.nanmin(lon)), float(np.nanmax(lat) - np.nanmin(lat)) * 1.6, 1e-4)
    return float(np.clip(math.log2(1406.0 / span) - 0.5, 3, 17))


def finish_map(fig, center_lat, center_lon, zoom=13, height=560):
    cfg = dict(center=dict(lat=center_lat, lon=center_lon), zoom=zoom)
    try:
        if hasattr(go, "Scattermap"):
            fig.update_layout(map_style="open-street-map", map=cfg, margin=dict(l=0, r=0, t=20, b=0), height=height)
        else:
            fig.update_layout(mapbox_style="open-street-map", mapbox=cfg, margin=dict(l=0, r=0, t=20, b=0), height=height)
    except Exception:  # noqa: BLE001
        fig.update_layout(margin=dict(l=0, r=0, t=20, b=0), height=height)
    return fig


# ==========================================================================
# 10. EKSPORT
# ==========================================================================
def _strip_md(s):
    return str(s).replace("**", "").replace("`", "")


def build_excel_bytes(sheets, conclusions=None):
    """Skoroszyt xlsx: słownik {nazwa arkusza: DataFrame} (+ arkusz „Wnioski”)."""
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        if conclusions:
            rows = [{"Wniosek": it["title"], "Ocena": it["kind"], "Opis": _strip_md(it["text"])} for it in conclusions["items"]]
            pd.DataFrame([{"Wniosek": "WERDYKT", "Ocena": conclusions.get("kind", ""),
                           "Opis": _strip_md(conclusions.get("grade", ""))}] + rows).to_excel(xw, sheet_name="Wnioski", index=False)
        used = {"Wnioski"}
        for name, df in sheets.items():
            if df is None or len(df) == 0:
                continue
            nm = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:31]
            if nm in used:
                continue
            used.add(nm)
            df.to_excel(xw, sheet_name=nm, index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(48, max(9, width + 2))
            ws.freeze_panes = "A2"
    return bio.getvalue()


def build_html_report(title, conclusions_md, figs, tables):
    """Samodzielny raport HTML (wykresy interaktywne, biblioteka Plotly z CDN)."""
    parts = [f"<html><head><meta charset='utf-8'><title>{title}</title>"
             "<style>body{font-family:Segoe UI,Arial,sans-serif;max-width:1100px;margin:24px auto;padding:0 16px;color:#222}"
             "table{border-collapse:collapse;font-size:13px}td,th{border:1px solid #ccc;padding:3px 8px}"
             "h1,h2{border-bottom:1px solid #ddd;padding-bottom:4px}</style></head><body>",
             f"<h1>{title}</h1><p>Wygenerowano: {datetime.datetime.now():%Y-%m-%d %H:%M}</p>"]
    html_md = []
    for line in str(conclusions_md).splitlines():
        line = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line)
        html_md.append(f"<p>{line}</p>" if line.strip() else "")
    parts.append("<h2>Wnioski</h2>" + "\n".join(html_md))
    first = True
    for name, fig in figs:
        try:
            parts.append(f"<h2>{name}</h2>" + fig.to_html(full_html=False, include_plotlyjs="cdn" if first else False))
            first = False
        except Exception:  # noqa: BLE001
            continue
    for name, df in tables:
        if df is not None and len(df):
            parts.append(f"<h2>{name}</h2>" + df.to_html(index=False, float_format=lambda v: f"{v:.2f}"))
    parts.append("</body></html>")
    return "\n".join(parts).encode("utf-8")


# ==========================================================================
# 11. INTERFEJS – CZĘŚĆ A: wczytywanie, cache, pomocnicze
# ==========================================================================
SUPPORTED_EXT = (".xlsx", ".xlsm", ".csv", ".txt", ".rsp", ".gpx", ".xml", ".2025", ".2024", ".2023", ".2026", ".dat")


@st.cache_resource(show_spinner=False, max_entries=16)
def cached_parse(ident, typ, _ref):
    return parse_source(_ref, typ)


@st.cache_resource(show_spinner=False, max_entries=6)
def cached_assemble(sig, _entries, camera, reverse_lcms, rel_mode):
    return assemble_datasets(_entries, camera=camera, reverse_lcms=reverse_lcms, relative_mode=rel_mode)


@st.cache_resource(show_spinner=False, max_entries=6)
def cached_lcms_calc(uid, window, _L):
    return compute_lcms_calc_ruts(_L, window)


@st.cache_resource(show_spinner=False, max_entries=6)
def cached_rsp_raw_ruts(uid, xs, _R):
    return compute_rsp_ruts(_R["raw"]["lasers"], list(xs))


@st.cache_data(show_spinner=False, max_entries=64)
def cached_stats(a, b, tol):
    return agreement_stats(a, b, tol)


def wide(fn, *args, **kw):
    """Wywołanie widgetu na całą szerokość – zgodne ze starszymi i nowszymi wersjami Streamlit."""
    try:
        return fn(*args, width="stretch", **kw)
    except TypeError:
        return fn(*args, use_container_width=True, **kw)
    except Exception as exc:  # noqa: BLE001
        if "width" in str(exc).lower():
            return fn(*args, use_container_width=True, **kw)
        raise


def fmt_mb(n):
    return f"{n / 1048576:.1f}"


def list_folder_files(path):
    out = []
    if path and os.path.isdir(path):
        for f in sorted(os.listdir(path)):
            fp = os.path.join(path, f)
            if os.path.isfile(fp) and (f.lower().endswith(SUPPORTED_EXT) or "." in f):
                out.append(fp)
    return out


def collect_refs(uploaded, folder):
    refs = {}
    for uf in uploaded or []:
        r = file_ref_from_upload(uf)
        refs[r.ident] = r
    for p in list_folder_files(folder):
        try:
            r = file_ref_from_path(p)
            refs[r.ident] = r
        except OSError:
            continue
    return list(refs.values())


def classify_cached(ref):
    cache = st.session_state.setdefault("_cls", {})
    if ref.ident not in cache:
        cache[ref.ident] = classify_file(ref)
    return cache[ref.ident]


def parse_all(refs, reg):
    """Wczytuje wszystkie używane pliki wg zarejestrowanego typu. Zwraca (entries, errors)."""
    entries, errors = [], []
    for r in refs:
        cfg = reg[r.ident]
        if cfg["typ"] == TYPE_SKIP:
            continue
        try:
            src = cached_parse(r.ident, cfg["typ"], r)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{r.name}: {exc}")
            continue
        if cfg.get("start") is None or (isinstance(cfg["start"], float) and np.isnan(cfg["start"])) or cfg.get("start_for") != cfg["typ"]:
            gs = source_default_start(src)
            lo, _ = source_extent(src, gs, cfg.get("step", 1.0))
            cfg["start"] = float(lo) if not np.isnan(lo) else 0.0
            cfg["start_for"] = cfg["typ"]
        e = {"ident": r.ident, "name": r.name, "type": cfg["typ"], "use": bool(cfg["use"]), "src": src,
             "guess_start": source_default_start(src), "step_m": float(cfg.get("step", 1.0) or 1.0),
             "start_m": float(cfg["start"])}
        entries.append(e)
    return entries, errors


def entries_sig(entries):
    return "|".join(f"{e['ident']}:{e['type']}:{int(e['use'])}:{e['start_m']}:{e['step_m']}" for e in entries)


def sys_desc(kind, A):
    return ""


def get_rsp_xvals(A, n):
    xs = list(A.get("rsp_x") or DEFAULT_RSP_X_LIST)
    if len(xs) != n:
        xs = [round((i - (n - 1) / 2) * 0.165, 3) for i in range(n)]
    return np.array(xs, dtype=float)


def rsp_geometry_ui(R):
    """Pozycje X laserów RSP: z pliku / domyślne (edytowalne) albo stały odstęp – jak w poprzedniej wersji."""
    n = R["raw"]["lasers"].shape[1] if R is not None and R.get("raw") is not None else 21
    file_x = R["raw"].get("x_pos") if R is not None and R.get("raw") is not None else None
    base = file_x if file_x and len(file_x) == n else DEFAULT_RSP_X_LIST
    with st.sidebar.expander("Geometria laserów RSP"):
        mode = st.radio("Typ konfiguracji czujników RSP:", ["Z pliku / domyślna RSP", "Stały odstęp (równomierny)"], key="rsp_geo")
        if mode.startswith("Z pliku"):
            txt = st.text_area("Pozycje X [m] dla laserów:", ", ".join(f"{x:.3f}" for x in base), height=90, key=f"rsp_x_txt_{n}_{int(file_x is not None)}")
            try:
                xs = [float(v.strip().replace(",", ".")) for v in txt.split(",") if v.strip()]
            except Exception:  # noqa: BLE001
                xs = list(base)
        else:
            step = st.number_input("Odstęp [m]:", 0.05, 0.50, 0.165, 0.005, format="%.3f", key="rsp_geo_step")
            xs = [round((i - (n - 1) / 2) * step, 3) for i in range(n)]
        if len(xs) != n:
            st.warning(f"Podano {len(xs)} pozycji, a laserów jest {n} – użyto stałego odstępu 0.165 m.")
            xs = [round((i - (n - 1) / 2) * 0.165, 3) for i in range(n)]
    return xs


def get_rsp_profile(A, c):
    R, s = A["rsp"], A["rsp_s"]
    if R is None or R.get("raw") is None or s is None or s.get("raw_index") is None:
        return None
    i = int(s["grid"].index(c))
    if i < 0 or i >= s["grid"].n:
        return None
    k = int(s["raw_index"][i])
    if k < 0:
        return None
    return R["raw"]["lasers"][k], int(k)


def get_lcms_profile(A, c_lcms):
    L, s = A["lcms"], A["lcms_s"]
    if L is None or s is None:
        return None
    i = int(s["grid"].index(c_lcms))
    if i < 0 or i >= s["grid"].n:
        return None
    k = int(s["prof_index"][i])
    if k < 0:
        return None
    a, b = L["ptr"][k], L["ptr"][k + 1]
    return L["x"][a:b].astype(float), L["z"][a:b].astype(float), k


def cover_info(valid, pairs, a0, a1):
    n_rng = max(1, int(a1 - a0) + 1)
    gaps = ""
    if len(valid) >= 2:
        d = np.diff(valid["Pikietaz"].values)
        big = np.where(d > 20)[0]
        if big.size:
            tot = int(np.sum(d[big] - 1))
            gaps = (f"Wykryto {big.size} luk (>20 m) w danych wspólnych, łącznie ok. {tot} m "
                    f"(np. {km_str(valid['Pikietaz'].values[big[0]])} → {km_str(valid['Pikietaz'].values[big[0] + 1])}).")
    return {"n_valid": int(len(valid)), "c0": float(a0), "c1": float(a1), "length_m": float(a1 - a0),
            "pct": 100.0 * len(valid) / n_rng, "gaps_note": gaps}


def pending_apply():
    p = st.session_state.pop("_pending", None)
    if p:
        for k, v in p.items():
            st.session_state[k] = v


def set_pending(**kw):
    st.session_state["_pending"] = kw
    st.rerun()


def section_edges_df():
    return st.session_state.get("sec_off_df")


def current_offset_fn(a0):
    g = float(st.session_state.get("off_global", 0.0))
    dr = float(st.session_state.get("off_drift", 0.0))
    sdf = section_edges_df()
    edges = offs = None
    if sdf is not None and len(sdf):
        s2 = sdf.dropna(subset=["Od [m]"]).sort_values("Od [m]")
        if len(s2) and np.any(np.abs(s2["Offset [m]"].fillna(0).values) > 0):
            edges = s2["Od [m]"].values.astype(float)
            offs = s2["Offset [m]"].fillna(0).values.astype(float)

    ref = float(st.session_state.get("off_ref", a0))

    def fn(chain):
        return make_offsets(chain, g, dr, ref, edges, offs)
    return fn


def nice_number(v, nd=2):
    return "–" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.{nd}f}"


def stats_table(stx, tol):
    rows = []
    spec = [("n", "Liczba punktów (n)", 0), ("n_eff", "n efektywne (z autokorelacją)", 0),
            ("mean_rsp", "Średnia RSP [mm]", 2), ("mean_lcms", "Średnia LCMS [mm]", 2),
            ("bias", "Bias LCMS−RSP [mm]", 2), ("bias_ci_lo", "Bias – 95% CI dół", 2), ("bias_ci_hi", "Bias – 95% CI góra", 2),
            ("sd", "SD różnic [mm]", 2), ("loa_lo", "Dolna granica zgodności [mm]", 2), ("loa_hi", "Górna granica zgodności [mm]", 2),
            ("rmse", "RMSE [mm]", 2), ("mae", "MAE [mm]", 2), ("medae", "Mediana |Δ| [mm]", 2), ("p95_abs", "95. percentyl |Δ| [mm]", 2),
            ("max_abs", "Maks. |Δ| [mm]", 1), ("pearson", "Korelacja Pearsona r", 3), ("spearman", "Korelacja Spearmana", 3),
            ("ccc", "CCC Lina", 3), ("icc", "ICC(A,1)", 3), ("slope", "Nachylenie MNK (LCMS=a+b·RSP)", 3),
            ("intercept", "Wyraz wolny MNK [mm]", 2), ("deming_slope", "Nachylenie ortogonalne (Deming)", 3),
            ("sd_ratio", "SD(LCMS)/SD(RSP)", 2), ("within_tol", f"% w tolerancji ±{tol:g} mm", 1),
            ("within_1", "% w ±1 mm", 1), ("within_3", "% w ±3 mm", 1), ("within_5", "% w ±5 mm", 1),
            ("share_lcms_higher", "% punktów LCMS > RSP", 1), ("share_rsp_higher", "% punktów RSP > LCMS", 1),
            ("rpd_median", "Mediana RPD [%]", 1)]
    for key, lab, nd in spec:
        row = {"Miara": lab}
        for k in ("L", "R", "Max"):
            v = stx.get(k, {}).get(key)
            row[{"L": "Lewa", "R": "Prawa", "Max": "Maksymalna"}[k]] = (
                "–" if v is None or (isinstance(v, float) and np.isnan(v)) else (f"{v:.{nd}f}" if nd else f"{int(round(v))}"))
        rows.append(row)
    return pd.DataFrame(rows)


# ==========================================================================
# 12. INTERFEJS – CZĘŚĆ B: zakładki
# ==========================================================================
def plot(fig, key):
    wide(st.plotly_chart, fig, key=key)


def make_pairs_src(A):
    rs, ls = A["rsp_s"], A["lcms_s"]
    rL, rR, _ = rsp_pick(rs, A["rsp_source"])
    if A["lcms_source"] == "calc" and ls.get("calc_L") is not None:
        lL, lR = ls["calc_L"], ls["calc_R"]
    else:
        lL, lR = ls["file_L"], ls["file_R"]
        if np.all(np.isnan(lL)) and ls.get("calc_L") is not None:
            lL, lR = ls["calc_L"], ls["calc_R"]
    if A["swap_lr"]:
        lL, lR = lR, lL
    return rL, rR, lL, lR, rs["grid"].c0, ls["grid"].c0


@st.cache_data(show_spinner=False, max_entries=8)
def cached_global_search(rL, rR, lL, lR, rc0, lc0, a0, a1, metric, hp, max_search):
    ps = {"rsp": {"L": rL, "R": rR, "Max": np.maximum(rL, rR)},
          "lcms": {"L": lL, "R": lR, "Max": np.maximum(lL, lR)}, "rsp_c0": rc0, "lcms_c0": lc0}
    return global_offset_search(ps, (a0, a1), max_search, metric, hp)


@st.cache_data(show_spinner=False, max_entries=8)
def cached_lag_map(rL, rR, lL, lR, rc0, lc0, a0, a1, metric, hp, win, hop, max_search):
    ps = {"rsp": {"L": rL, "R": rR, "Max": np.maximum(rL, rR)},
          "lcms": {"L": lL, "R": lR, "Max": np.maximum(lL, lR)}, "rsp_c0": rc0, "lcms_c0": lc0}
    return local_lag_map(ps, (a0, a1), win, hop, max_search, metric, hp)


# ---------- WNIOSKI ----------
def get_stats(A):
    memo = A["memo"]
    if "stats" not in memo:
        v = A["valid"]
        memo["stats"] = {k: cached_stats(v["RSP_" + k].values, v["LCMS_" + k].values, A["tol"]) for k in ("L", "R", "Max")} \
            if len(v) >= 3 else {}
    return memo["stats"]


def get_hm(A):
    memo = A["memo"]
    if "hm" not in memo:
        memo["hm"] = hm_table(A["valid"], A["sec_len"], origin=A["a0"], min_cover=A["min_cover"], has_lcms=True) \
            if len(A["valid"]) else pd.DataFrame()
    return memo["hm"]


def get_conclusions(A):
    memo = A["memo"]
    if "concl" in memo:
        return memo["concl"]
    v = A["valid"]
    stx = get_stats(A)
    hm = get_hm(A)
    cs_pt = cs_sec = None
    if len(v) >= 3:
        _, cs_pt = class_confusion(v["RSP_Max"].values, v["LCMS_Max"].values)
    if len(hm) >= 3:
        _, cs_sec = class_confusion(hm["RSP MAX_KOL Hm"].values, hm["LCMS MAX_KOL Hm"].values)
    sync = {"offset": float(st.session_state.get("off_global", 0.0)), "drift": float(st.session_state.get("off_drift", 0.0)),
            "n_sections": 0}
    sdf = section_edges_df()
    if sdf is not None and len(sdf) and np.any(np.abs(sdf["Offset [m]"].fillna(0).values) > 0):
        sync["n_sections"] = int(len(sdf))
    try:
        rL, rR, lL, lR, rc0, lc0 = make_pairs_src(A)
        lag, c1, c2, lags, corr = cached_global_search(rL, rR, lL, lR, rc0, lc0, A["a0"], A["a1"], "Max", 101, 150)
        cur = float(st.session_state.get("off_global", 0.0))
        k = int(np.argmin(np.abs(lags - round(cur))))
        sync.update({"auto_lag": lag, "auto_corr": c1, "cur_corr": float(corr[k]) if not np.isnan(corr[k]) else np.nan,
                     "auto_gain": (c1 - float(corr[k])) if not np.isnan(corr[k]) else 0.0})
    except Exception:  # noqa: BLE001
        pass
    sec = section_agreement(v, 1000.0, origin=A["a0"], tol=A["tol"]) if len(v) else pd.DataFrame()
    rpd_share = None
    if len(v) >= 3:
        t = build_rpd_table_df(v["RSP_Max"].values, v["LCMS_Max"].values)
        try:
            rpd_share = float(t.loc[1, "n"]) / max(1, float(t.loc[2, "n"])) * 100
        except Exception:  # noqa: BLE001
            rpd_share = None
    ctx = {"tol": A["tol"], "stats": stx, "cover": cover_info(v, A["pairs"], A["a0"], A["a1"]),
           "rsp_desc": A["rdesc"], "lcms_desc": A["ldesc"], "sync": sync, "cls_point": cs_pt, "cls_sec": cs_sec,
           "sections": sec, "quality_notes": A["notes"], "rpd_high_share": rpd_share}
    memo["concl"] = build_conclusions(ctx)
    memo["sec_1km"] = sec
    return memo["concl"]


def tab_conclusions(A):
    v = A["valid"]
    if len(v) < 3:
        st.warning("Brak wspólnych punktów RSP–LCMS w wybranym zakresie. Sprawdź pikietaże początkowe plików, "
                   "zakres analizy w panelu bocznym oraz przesunięcie (zakładka „Synchronizacja”).")
        return
    c = get_conclusions(A)
    st.subheader("Wnioski z analizy")
    icon = {"ok": "🟢", "info": "🔵", "warn": "🟡", "bad": "🔴"}.get(c["kind"], "🔵")
    st.markdown(f"### {icon} {c['grade']}")
    stx = get_stats(A)
    cols = st.columns(4)
    s = stx.get("Max", {})
    cols[0].metric("Punkty porównania", f"{s.get('n', 0):,}".replace(",", " "))
    cols[1].metric("Bias LCMS − RSP [mm]", nice_number(s.get("bias"), 2))
    cols[2].metric("Korelacja r (Max)", nice_number(s.get("pearson"), 2))
    cols[3].metric(f"W ±{A['tol']:g} mm", f"{nice_number(s.get('within_tol'), 0)}%")
    for it in c["items"]:
        ic = {"ok": "🟢", "info": "🔵", "warn": "🟡", "bad": "🔴"}.get(it["kind"], "🔵")
        with st.container(border=True):
            st.markdown(f"**{ic} {it['title']}**\n\n{it['text']}")
    fig, *_ = build_bland_altman_fig(v["RSP_Max"].values, v["LCMS_Max"].values, "KOLEINA MAKSYMALNA")
    plot(fig, "ov_ba")
    plot(fig_longitudinal(A["pairs"], [
        {"col": "RSP_Max", "name": "RSP Max", "color": C_RSP},
        {"col": "LCMS_Max", "name": "LCMS Max", "color": C_LCMS}], "Przebieg koleiny maksymalnej – RSP i LCMS",
        view=(A["a0"], A["a1"]), uirev="ov"), "ov_long")


# ---------- DANE I PLIKI ----------
def tab_data(A):
    st.subheader("Zakresy plików i zestawów")
    items = []
    for e in A["entries"]:
        if not e["use"]:
            continue
        try:
            lo, hi = source_extent(e["src"], e["guess_start"], e["step_m"])
            lo += e["start_m"] - native_chain(e["src"], e["guess_start"], e["step_m"])[0]
            hi += e["start_m"] - native_chain(e["src"], e["guess_start"], e["step_m"])[0]
            items.append((e["name"], e["type"], lo, hi))
        except Exception:  # noqa: BLE001
            continue
    ov = None
    if A["rsp"] is not None and A["lcms"] is not None:
        ov = (max(A["rsp"]["c_min"], A["lcms"]["c_min"] - A["off0"]), min(A["rsp"]["c_max"], A["lcms"]["c_max"] - A["off0"]))
    if items:
        plot(fig_coverage(items, ov), "cov")
    rows = []
    if A["rsp"] is not None:
        R = A["rsp"]
        rows.append({"System": "RSP", "Od": km_str(R["c_min"]), "Do": km_str(R["c_max"]),
                     "Profile laserów": 0 if R.get("raw") is None else len(R["raw"]["chain"]),
                     "Rekordy KOLC": 0 if R.get("kolc") is None else len(R["kolc"]["chain"]),
                     "GPS": 0 if R.get("gps") is None else len(R["gps"])})
    if A["lcms"] is not None:
        L = A["lcms"]
        rows.append({"System": "LCMS", "Od": km_str(L["c_min"]), "Do": km_str(L["c_max"]),
                     "Profile laserów": len(L["chain"]), "Rekordy KOLC": 0,
                     "GPS": 0 if L.get("gps") is None else len(L["gps"])})
    if rows:
        wide(st.dataframe, pd.DataFrame(rows), hide_index=True)
    if A["notes"]:
        with st.expander("Uwagi z wczytywania i scalania", expanded=True):
            for n in A["notes"]:
                st.write("• " + n)
    for e in A["entries"]:
        ns = e["src"].get("notes") or []
        if ns:
            with st.expander(f"Uwagi: {e['name']}"):
                for n in ns:
                    st.write("• " + n)
    if A["rsp"] is not None and A["rsp"].get("kolc_meta"):
        with st.expander("Komentarze / metadane KOLC"):
            wide(st.dataframe, pd.DataFrame(A["rsp"]["kolc_meta"]), hide_index=True)
    p = A["pairs"]
    if len(p):
        st.subheader("Jakość i kompletność danych (w zakresie analizy)")
        q = []
        for col, lab in (("RSP_L", "RSP lewa"), ("RSP_R", "RSP prawa"), ("LCMS_L", "LCMS lewa"), ("LCMS_R", "LCMS prawa")):
            if col in p:
                x = p[col].values
                q.append({"Seria": lab, "Dostępne [%]": round(100 * np.mean(~np.isnan(x)), 1),
                          "Średnia [mm]": round(float(np.nanmean(x)), 2) if np.any(~np.isnan(x)) else np.nan,
                          "Mediana [mm]": round(float(np.nanmedian(x)), 2) if np.any(~np.isnan(x)) else np.nan,
                          "Maks. [mm]": round(float(np.nanmax(x)), 1) if np.any(~np.isnan(x)) else np.nan})
        wide(st.dataframe, pd.DataFrame(q), hide_index=True)
        plot(fig_hist_depths(p["RSP_Max"].values, p["LCMS_Max"].values), "hist_dep")


# ---------- SYNCHRONIZACJA ----------
def tab_sync(A):
    st.subheader("Synchronizacja pikietażu RSP ↔ LCMS")
    st.caption("Konwencja: pikietaż LCMS = pikietaż RSP + przesunięcie. Przesunięcie = offset globalny + offset sekcji + dryf [m/km]. "
               "Wartości ustawiasz w panelu bocznym; poniższe narzędzia je tylko proponują.")
    if A["rsp_s"] is None or A["lcms_s"] is None:
        st.info("Potrzebne są oba systemy (RSP i LCMS).")
        return
    rL, rR, lL, lR, rc0, lc0 = make_pairs_src(A)
    c1, c2, c3 = st.columns(3)
    metric = c1.selectbox("Seria do korelacji", ["Max", "L", "R", "L+R"], key="sy_metric")
    hp = c2.number_input("Okno filtru górnoprzepustowego [m] (0 = brak)", 0, 1001, 101, 10, key="sy_hp")
    msr = c3.number_input("Zakres szukania ± [m]", 10, 2000, 150, 10, key="sy_max")
    cur = float(st.session_state.get("off_global", 0.0))
    with st.container(border=True):
        st.markdown("**1. Offset globalny (korelacja krzyżowa)**")
        lag, cc, c2nd, lags, corr = cached_global_search(rL, rR, lL, lR, rc0, lc0, A["a0"], A["a1"], metric, int(hp), int(msr))
        k = int(np.argmin(np.abs(lags - round(cur))))
        cur_corr = float(corr[k]) if not np.isnan(corr[k]) else np.nan
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Najlepsze przesunięcie [m]", nice_number(lag, 0))
        m2.metric("Korelacja w maksimum", nice_number(cc, 3))
        m3.metric("Korelacja przy bieżącym", nice_number(cur_corr, 3))
        m4.metric("Drugi pik (poza ±5 m)", nice_number(c2nd, 3))
        plot(fig_corr_curve(lags, corr, lag), "sy_corr")
        if not np.isnan(lag) and st.button(f"Zastosuj offset globalny {lag:+.0f} m", key="sy_apply_g"):
            set_pending(off_global=float(lag), off_ref=float(A["a0"]))
    with st.container(border=True):
        st.markdown("**2. Lokalne przesunięcia i dryf (okna przesuwne)**")
        w1, w2, w3 = st.columns(3)
        win = w1.number_input("Okno [m]", 300, 5000, 1000, 100, key="sy_win")
        hop = w2.number_input("Krok okna [m]", 50, 2000, 250, 50, key="sy_hop")
        loc_rng = w3.number_input("Zakres lokalny ± [m] (wokół offsetu)", 10, 300, 40, 10, key="sy_locrng")
        if st.button("Policz mapę lokalnych przesunięć", key="sy_run_local"):
            with st.spinner("Liczenie okien…"):
                st.session_state["lagmap"] = cached_lag_map(rL, rR, lL, lR, rc0, lc0, A["a0"], A["a1"], metric,
                                                            int(hp), int(win), int(hop), int(max(msr, abs(cur) + loc_rng)))
        dfl = st.session_state.get("lagmap")
        if dfl is not None and len(dfl):
            fit = fit_drift(dfl, A["a0"])
            xx = np.array([A["a0"], A["a1"]], dtype=float)
            plot(fig_lag_map(dfl, fit, (xx, A["off_fn"](xx))), "sy_lagmap")
            if fit:
                st.write(f"Dopasowanie: offset w {km_str(A['a0'])} = **{fit['offset_ref']:+.1f} m**, dryf = **{fit['drift_per_km']:+.2f} m/km** "
                         f"(R² = {fit['r2']:.2f}, {fit['n']}/{fit['n_all']} okien, odchyłka mediana {fit['resid_mad']:.1f} m).")
                b1, b2 = st.columns(2)
                if b1.button("Zastosuj offset + dryf", key="sy_apply_fit"):
                    st.session_state["sec_off_df"] = None
                    set_pending(off_global=float(fit["offset_ref"]), off_drift=float(fit["drift_per_km"]), off_ref=float(A["a0"]))
                if b2.button("Wypełnij offsety sekcji (co 1 km) z lagów", key="sy_apply_sec"):
                    edges = list(np.arange(A["a0"], A["a1"], 1000.0))
                    offs = section_offsets_from_lags(dfl, edges + [A["a1"] + 1], default=cur)
                    st.session_state["sec_off_df"] = pd.DataFrame({"Od [m]": edges, "Offset [m]": offs})
                    set_pending(off_global=0.0, off_drift=0.0, off_ref=float(A["a0"]))
            wide(st.dataframe, dfl.round(3), hide_index=True, height=220)
        elif dfl is not None:
            st.info("Brak wiarygodnych okien (za mało wspólnych danych lub zbyt słaba korelacja).")
    with st.container(border=True):
        st.markdown("**3. Offsety sekcji (ręczna edycja)**")
        sdf = section_edges_df()
        if sdf is None:
            sdf = pd.DataFrame({"Od [m]": [float(A["a0"])], "Offset [m]": [0.0]})
        ed = wide(st.data_editor, sdf, num_rows="dynamic", hide_index=True, key="sy_sec_editor")
        st.session_state["sec_off_df"] = ed
        st.caption("Offset sekcji obowiązuje od podanego pikietażu do następnego wiersza i jest dodawany do offsetu globalnego.")
    with st.container(border=True):
        st.markdown("**4. Dopasowanie po GPS (najbliższy punkt śladu)**")
        rg = A["rsp"].get("gps") if A["rsp"] else None
        lg = A["lcms"].get("gps") if A["lcms"] else None
        if rg is None or lg is None or len(rg) < 2 or len(lg) < 2:
            st.info("Brak śladów GPS po obu stronach – nie można porównać położenia. "
                    "W danych KOLC pole GPS bywa puste dla części trasy; wtedy GPS nie pomoże w synchronizacji tego odcinka.")
        else:
            go_ = gps_nearest_offsets(rg, lg, 10.0, (A["a0"], A["a1"]))
            if go_ is None or len(go_) == 0:
                st.warning("Ślady GPS RSP i LCMS nie zbliżają się do siebie w zakresie analizy (lub ślad RSP nie pokrywa zakresu).")
            else:
                med = float(np.median(go_["Offset [m]"]))
                st.write(f"Mediana offsetu z GPS: **{med:+.1f} m** (n = {len(go_)}, mediana odległości śladów {np.median(go_['Dystans [m]']):.1f} m).")
                f = go.Figure(go.Scattergl(x=go_["Pikietaz"], y=go_["Offset [m]"], mode="markers", marker=dict(size=4)))
                f.update_layout(height=300, xaxis_title="Pikietaż RSP [m]", yaxis_title="Offset z GPS [m]", template="plotly_white")
                plot(f, "sy_gps")
                if st.button(f"Zastosuj offset z GPS {med:+.0f} m", key="sy_apply_gps"):
                    set_pending(off_global=med, off_ref=float(A["a0"]))


# ---------- PRZEKRÓJ ----------
def _xs_step(d):
    st.session_state["xs_chain"] = float(st.session_state.get("xs_chain", 0.0)) + d


def tab_cross(A):
    st.subheader("Przekrój poprzeczny (profil)")
    p = A["pairs"]
    lo, hi = float(A["a0"]), float(A["a1"])
    if "xs_chain" not in st.session_state or not (lo <= st.session_state["xs_chain"] <= hi):
        st.session_state["xs_chain"] = float(np.floor((lo + hi) / 2))
    c1, c2, c3, c4 = st.columns([1, 1, 1, 3])
    c1.button("◀ −10 m", on_click=_xs_step, args=(-10.0,), key="xs_m10")
    c2.button("◀ −1 m", on_click=_xs_step, args=(-1.0,), key="xs_m1")
    c3.button("+1 m ▶", on_click=_xs_step, args=(1.0,), key="xs_p1")
    chain = c4.number_input("Pikietaż RSP [m]", min_value=lo, max_value=hi, step=1.0, key="xs_chain")
    chain = float(np.round(chain))
    st.caption(f"Pikietaż RSP: **{km_str(chain)}**; pikietaż LCMS = {km_str(chain + float(A['off_fn'](np.array([chain]))[0]))}")
    o1, o2, o3, o4 = st.columns(4)
    zero_l11 = o1.checkbox("Zeruj do L11", True, key="xs_zero")
    show_rh = o2.checkbox("Łata RSP", True, key="xs_rh")
    show_lh = o3.checkbox("Łata LCMS", True, key="xs_lh")
    manual = o4.number_input("Przesunięcie poprzeczne LCMS [m]", -1.0, 1.0, 0.0, 0.01, key="xs_shift")
    p1, p2, p3 = st.columns(3)
    skal = p1.checkbox("Skaluj szerokość LCMS do RSP", False, key="xs_scale")
    zacz = p2.checkbox("Zaczep prawej krawędzi", True, key="xs_hook")
    wyr = p3.checkbox("Wyrównaj spadek poprzeczny", True, key="xs_level")
    off = float(A["off_fn"](np.array([chain]))[0])
    rs = lc = None
    rp = get_rsp_profile(A, chain)
    x_vals = None
    if rp is not None:
        vals, k = rp
        x_vals = get_rsp_xvals(A, len(vals))
        rs = rsp_cross_section(vals, x_vals, zero_l11)
    lp = get_lcms_profile(A, chain + off)
    label = (None, None)
    if lp is not None:
        x, z, kk = lp
        xmin, xmax = (float(x_vals.min()), float(x_vals.max())) if x_vals is not None else (-1.85, 1.85)
        lc = lcms_cross_section(x, z, xmin, xmax, A["win"], skal, zacz, manual, wyr)
        fl, fr = A["lcms"]["rut_l"][kk], A["lcms"]["rut_r"][kk]
        label = (fl, fr) if not A["swap_lr"] else (fr, fl)
        if A["lcms_source"] == "calc":
            label = (None, None)
        if not bool(A["lcms"]["complete"][kk]):
            st.warning("Profil LCMS ma tylko jedną połówkę (lewą lub prawą) – wyniki częściowe.")
    if rs is None and lc is not None and A["rsp"] is not None and A["rsp"].get("raw") is None:
        st.info("Pokazany jest tylko przekrój LCMS – plik KOLC nie zawiera odczytów laserów RSP (tylko gotowe głębokości kolein).")
    if rs is None and lc is None:
        if A["rsp"] is not None and A["rsp"].get("raw") is None:
            st.warning("Przekrój RSP nie jest dostępny: wczytano tylko plik KOLC, który zawiera gotowe głębokości kolein co 1 m "
                       "(L_KOL, P_KOL, PEL_KOL), ale nie odczyty 21 laserów. Aby zobaczyć przekrój RSP, wgraj surowe pliki RSP "
                       "(.RSP / xlsx / csv) – ich nazwy podaje nagłówek pliku KOLC.")
        else:
            st.warning("Brak profilu w tym pikietażu (RSP: luka w danych laserów; LCMS: brak profilu).")
    else:
        fig = fig_cross_section(rs, lc, f"Przekrój – RSP {km_str(chain)} / LCMS {km_str(chain + off)}", show_rh, show_lh, label,
                                (min(-1.85, (x_vals.min() if x_vals is not None else -1.85)),
                                 max(1.85, (x_vals.max() if x_vals is not None else 1.85))))
        plot(fig, "xs_fig")
    row = p[p["Pikietaz"] == chain]
    if len(row):
        r = row.iloc[0]
        for col_, (k, lab) in zip(st.columns(3), (("L", "lewa"), ("R", "prawa"), ("Max", "maksymalna"))):
            rv, lv = r.get("RSP_" + k), r.get("LCMS_" + k)
            col_.metric(f"RSP {lab}", nice_number(rv, 1) + " mm")
            col_.metric(f"LCMS {lab}", nice_number(lv, 1) + " mm", delta=nice_number(lv - rv, 1) + " mm")
    if rs is not None and lc is not None:
        st.caption(f"Z geometrii wykresu: RSP lewa {rs['rut_l']:.1f}, prawa {rs['rut_r']:.1f} mm; "
                   f"LCMS (wirtualna łata) lewa {lc['rut_l_calc']:.1f}, prawa {lc['rut_r_calc']:.1f} mm.")
    with st.expander("Mapa cieplna profili LCMS wokół wybranego pikietażu"):
        half = st.slider("Zakres ± [m]", 5, 150, 40, key="xs_hm_half")
        step = max(1, int(half * 2 / 120))
        cs = np.arange(chain - half, chain + half + 1, step)
        xg = np.linspace(-1.9, 1.9, 77)
        rows_, ch_ = [], []
        for c in cs:
            q = get_lcms_profile(A, c + float(A["off_fn"](np.array([c]))[0]))
            if q is None:
                continue
            rows_.append(level_profile_for_map(q[0], q[1], xg, A["win"]))
            ch_.append(c)
        if rows_:
            plot(fig_profile_heatmap(ch_, np.array(rows_), "mm", xg, "Profile LCMS wyrównane (0 = oś)", selected=chain), "xs_hm")
        else:
            st.info("Brak profili LCMS w tym zakresie.")


# ---------- PRZEBIEG WZDŁUŻNY ----------
@st.cache_data(show_spinner=False, max_entries=16)
def cached_segment_lags(ch, a, b, seg_len, max_lag, hp, min_corr):
    """Dla kolejnych segmentów RSP szuka dodatkowego przesunięcia LCMS (maks. korelacji)."""
    ah, bh = highpass(a, hp), highpass(b, hp)
    c0 = float(ch[0])
    lags = np.arange(-max_lag, max_lag + 1)
    rows = []
    for st_ in np.arange(c0, ch[-1] + 1, seg_len):
        i0 = int(st_ - c0)
        i1 = min(len(ch), int(st_ + seg_len - c0))
        if i1 - i0 < 30:
            continue
        corr = corr_vs_lag(ah[i0:i1], bh, st_, c0, lags)
        lag, c, _ = best_lag(corr, lags)
        before = float(corr[max_lag]) if not np.isnan(corr[max_lag]) else np.nan
        if np.isnan(lag) or np.isnan(c) or c < min_corr:
            lag = 0.0
        rows.append({"Od [m]": float(st_), "Do [m]": float(st_ + seg_len - 1), "Dodatkowe przesunięcie [m]": float(lag),
                     "r przed": before, "r po": c})
    return pd.DataFrame(rows)


def apply_segment_shifts(p, seg_df, seg_len):
    """Kopia tabeli porównawczej, w której LCMS jest przesunięty osobno w każdym segmencie."""
    q = p.copy()
    ch = p["Pikietaz"].values
    c0 = float(ch[0])
    for col in ("L", "R", "Max"):
        src = p["LCMS_" + col].values.astype(float)
        out = np.full(len(src), np.nan)
        for _, r in seg_df.iterrows():
            i0 = max(0, int(r["Od [m]"] - c0))
            i1 = min(len(ch), int(r["Od [m]"] + seg_len - c0))
            idx = np.arange(i0, i1) + int(round(r["Dodatkowe przesunięcie [m]"]))
            ok = (idx >= 0) & (idx < len(src))
            seg = np.full(i1 - i0, np.nan)
            seg[ok] = src[idx[ok]]
            out[i0:i1] = seg
        q["LCMS_" + col] = out
    off_tot = p["Off"].values.astype(float).copy() if "Off" in p.columns else np.zeros(len(p))
    extra = np.zeros(len(p))
    for _, r in seg_df.iterrows():
        i0 = max(0, int(r["Od [m]"] - c0))
        i1 = min(len(ch), int(r["Od [m]"] + seg_len - c0))
        extra[i0:i1] = float(int(round(r["Dodatkowe przesunięcie [m]"])))
    q["Off"] = off_tot + extra
    q["Off_seg"] = extra
    return q


def view_overlaps(x, A, v0, span):
    hi = v0 + span if span > 0 else A["a1"]
    return v0 <= x <= hi


def tab_long(A):
    st.subheader("Przebieg wzdłużny koleiny")
    p = A["pairs"]
    lo, hi = float(A["a0"]), float(A["a1"])
    c1, c2 = st.columns(2)
    v0 = c1.number_input("Widok od [m]", lo, hi, lo, 100.0, key="lg_v0b")
    span = c2.number_input("Szerokość widoku [m] (0 = cały odcinek)", 0.0, hi - lo, 0.0, 100.0, key="lg_span2")
    view = (v0, min(hi, v0 + span)) if span > 0 else (lo, hi)
    bands = st.checkbox("Pokaż klasy DSN (tło)", True, key="lg_bands")
    both = A["lcms_source"] == "both" and "LCMSc_L" in p.columns
    marks = []
    aligned = False
    with st.expander("Dopasowanie segmentów (lokalne przesunięcie LCMS na tym samym wykresie)"):
        st.caption("Trasa jest dzielona na segmenty; w każdym LCMS jest osobno przesuwany względem RSP (maksimum korelacji), "
                   "dzięki czemu na wykresie nie ma narastających przesunięć. Wartości możesz poprawić ręcznie.")
        on = st.checkbox("Włącz dopasowanie segmentów", False, key="sg_on")
        g1, g2, g3, g4, g5 = st.columns(5)
        seg_len = g1.number_input("Długość segmentu [m]", 50, 5000, 250, 50, key="sg_len")
        max_lag = g2.number_input("Maks. przesunięcie ± [m]", 2, 300, 30, 2, key="sg_max")
        col_ = g3.selectbox("Seria do dopasowania", ["Max", "L", "R"], key="sg_col")
        hp_ = g4.number_input("Filtr górnoprzepustowy [m]", 0, 1001, 101, 10, key="sg_hp")
        mc_ = g5.number_input("Min. korelacja", 0.0, 1.0, 0.2, 0.05, key="sg_minc")
        if on:
            sig = (int(seg_len), int(max_lag), col_, int(hp_), float(mc_), float(A["a0"]), float(A["a1"]),
                   round(float(st.session_state.get("off_global", 0.0)), 3), round(float(st.session_state.get("off_drift", 0.0)), 4))
            if st.session_state.get("sg_sig") != sig:
                auto = cached_segment_lags(p["Pikietaz"].values, p["RSP_" + col_].values, p["LCMS_" + col_].values,
                                           int(seg_len), int(max_lag), int(hp_), float(mc_))
                st.session_state["sg_df"] = auto
                st.session_state["sg_sig"] = sig
            sdf = st.session_state["sg_df"]
            ed = wide(st.data_editor, sdf, hide_index=True, key="sg_editor",
                      disabled=["Od [m]", "Do [m]", "r przed", "r po"])
            seg_df = ed
            if len(seg_df):
                p = apply_segment_shifts(p, seg_df, int(seg_len))
                aligned = True
                marks = [(float(x), "") for x in seg_df["Od [m]"].values if view_overlaps(x, A, v0, span)]
                st.write(f"Mediana r: przed **{np.nanmedian(seg_df['r przed']):.2f}**, po **{np.nanmedian(seg_df['r po']):.2f}**; "
                         f"przesunięcia od {seg_df['Dodatkowe przesunięcie [m]'].min():+.0f} do {seg_df['Dodatkowe przesunięcie [m]'].max():+.0f} m.")
                if st.button("Zastosuj jako offsety sekcji w całej analizie", key="sg_apply"):
                    mids = seg_df["Od [m]"].values + float(seg_len) / 2.0
                    tot = A["off_fn"](mids) + seg_df["Dodatkowe przesunięcie [m]"].values
                    st.session_state["sec_off_df"] = pd.DataFrame({"Od [m]": seg_df["Od [m]"].values, "Offset [m]": tot})
                    st.session_state["sg_sig"] = None
                    set_pending(off_global=0.0, off_drift=0.0, off_ref=float(A["a0"]))
    pv = p[(p["Pikietaz"] >= view[0]) & (p["Pikietaz"] <= view[1])]
    if "Off" in pv.columns and len(pv):
        o = pv["Off"].dropna()
        sdf_ = section_edges_df()
        has_sec = sdf_ is not None and len(sdf_) and bool(np.any(np.abs(sdf_["Offset [m]"].fillna(0).values) > 0))
        txt = (f"Przesunięcie LCMS względem RSP w tym widoku: od **{o.min():+.1f}** do **{o.max():+.1f} m** "
               f"(globalne {float(st.session_state.get('off_global', 0.0)):+.1f} m, dryf {float(st.session_state.get('off_drift', 0.0)):+.2f} m/km"
               f"{', offsety sekcji aktywne' if has_sec else ''}).")
        if aligned:
            txt += (f" Dopasowanie segmentów dodaje {pv['Off_seg'].min():+.0f} … {pv['Off_seg'].max():+.0f} m "
                    "– wartość widoczna po najechaniu na krzywą LCMS.")
        else:
            txt += " Bez dopasowania segmentów LCMS nie jest dodatkowo przesuwany. Wartość przesunięcia widać po najechaniu na krzywą LCMS."
        st.info(txt)
    for key, lab in (("Max", "maksymalna"), ("L", "lewa"), ("R", "prawa")):
        tr = [{"col": f"RSP_{key}", "name": f"RSP {lab}", "color": C_RSP},
              {"col": f"LCMS_{key}", "name": f"LCMS {lab}" + (" (dopasowany)" if aligned else ""),
               "color": C_LCMS, "hover_col": "Off"}]
        if both:
            tr.append({"col": f"LCMSc_{key}", "name": f"LCMS (łata) {lab}", "color": "#8fd18f", "dash": "dot"})
        plot(fig_longitudinal(p, tr, f"Koleina {lab} – RSP vs LCMS", view, bands=bands, uirev="lg" + key, extra_vlines=marks), "lg_" + key)
    roll = st.slider("Okno średniej kroczącej różnicy [m]", 11, 1001, 101, 10, key="lg_roll")
    plot(fig_diff_along(p, view=view, roll=roll, tol=A["tol"], uirev="lgd"), "lg_diff")
    with st.expander("Największe rozbieżności (|LCMS − RSP| dla koleiny maks., średnia z 20 m)"):
        d = (p["LCMS_Max"] - p["RSP_Max"]).abs()
        sm = pd.Series(d.values).rolling(21, center=True, min_periods=10).mean()
        t = pd.DataFrame({"Pikietaz": p["Pikietaz"].values, "Δ śr. [mm]": sm.values}).dropna()
        t = t.sort_values("Δ śr. [mm]", ascending=False)
        picked, rows = [], []
        for _, r in t.iterrows():
            if all(abs(r["Pikietaz"] - q) > 50 for q in picked):
                picked.append(r["Pikietaz"])
                rows.append({"Pikietaż": km_str(r["Pikietaz"]), "[m]": r["Pikietaz"], "Δ śr. [mm]": round(r["Δ śr. [mm]"], 2)})
            if len(rows) >= 15:
                break
        wide(st.dataframe, pd.DataFrame(rows), hide_index=True)


# ---------- OCENA ODCINKOWA ----------
def tab_dsn(A):
    st.subheader(f"Ocena odcinkowa DSN 2019 – odcinki {A['sec_len']:g} m  (Hm = E(h) + 2·Dh)")
    hm = get_hm(A)
    if len(hm) == 0:
        st.info("Brak odcinków do oceny.")
        return
    cfg = [("L_KOL", "Lewa koleina (L_KOL)"), ("P_KOL", "Prawa koleina (P_KOL)"), ("MAX_KOL", "Koleina maksymalna (MAX_KOL)")]
    for pre, title in cfg:
        plot(make_50m_bar_chart(hm, pre, title, A["sec_len"]), "dsn_bar_" + pre)
    plot(fig_class_strip(hm, "MAX_KOL"), "dsn_strip")
    M, cs = class_confusion(hm["RSP MAX_KOL Hm"].values, hm["LCMS MAX_KOL Hm"].values)
    if cs.get("n", 0):
        a, b = st.columns([1, 1])
        with a:
            plot(fig_confusion(M, "Zgodność klas odcinkowych (MAX_KOL)"), "dsn_conf")
        with b:
            st.metric("Zgodność klas (dokładna)", f"{cs['exact']:.0f}%")
            st.metric("Zgodność ±1 klasa", f"{cs['within1']:.0f}%")
            st.metric("Kappa ważona", nice_number(cs["kappa_w"], 2))
            st.write(f"LCMS wskazuje wyższą klasę (gorszy stan) na {cs['lcms_higher']:.0f}% odcinków, niższą na {cs['lcms_lower']:.0f}%.")
    show = hm.copy()
    for pre in ("L_KOL", "P_KOL", "MAX_KOL"):
        for s_ in ("RSP", "LCMS"):
            c = f"{s_} {pre} Klasa"
            show[c] = show[c].map(lambda k: f"{DSN_EMOJI.get(k, '')} {k}")
    keep = ["Odcinek", "n RSP", "n LCMS", "Pokrycie %"] + [f"{s_} {pre} {q}" for pre in ("L_KOL", "P_KOL", "MAX_KOL")
                                                              for s_ in ("RSP", "LCMS") for q in ("Hm", "Klasa")]
    wide(st.dataframe, show[[c for c in keep if c in show.columns]], hide_index=True)
    st.download_button("Pobierz tabelę odcinków (CSV)", hm.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
                       "ocena_odcinkowa_DSN.csv", "text/csv", key="dl_hm_csv")


# ---------- STATYSTYKI ----------
def tab_stats(A):
    st.subheader("Zgodność statystyczna RSP vs LCMS")
    v = A["valid"]
    if len(v) < 3:
        st.info("Za mało wspólnych punktów.")
        return
    stx = get_stats(A)
    wide(st.dataframe, stats_table(stx, A["tol"]), hide_index=True, height=560)
    sel = st.radio("Seria do wykresów", ["Max", "L", "R"], horizontal=True, key="st_sel",
                   format_func=lambda k: {"Max": "Maksymalna", "L": "Lewa", "R": "Prawa"}[k])
    a, b = v["RSP_" + sel].values, v["LCMS_" + sel].values
    lab = {"Max": "KOLEINA MAKSYMALNA", "L": "LEWA KOLEINA", "R": "PRAWA KOLEINA"}[sel]
    fig, md, sd, lu, ll = build_bland_altman_fig(a, b, lab)
    plot(fig, "st_ba")
    c1, c2 = st.columns(2)
    with c1:
        plot(fig_scatter_xy(a, b, f"LCMS vs RSP – {lab.lower()}", st_=stx.get(sel)), "st_sc")
    with c2:
        plot(fig_diff_hist(a, b, A["tol"]), "st_hist")
    c3, c4 = st.columns(2)
    with c3:
        plot(fig_ecdf(a, b, A["tol"]), "st_ecdf")
    with c4:
        plot(fig_box_by_depth(a, b), "st_box")
    st.markdown("**Zgodność w przedziałach głębokości**")
    wide(st.dataframe, depth_bin_table(a, b, tol=A["tol"]), hide_index=True)
    st.markdown("**Analiza RPD (układ GDDKiA: 0–30% i >30%)**")
    wide(st.dataframe, build_rpd_table_df(a, b), hide_index=True)
    with st.expander("RPD w drobniejszych przedziałach"):
        wide(st.dataframe, rpd_fine_table(a, b), hide_index=True)
    M, cs = class_confusion(a, b)
    if cs.get("n", 0):
        st.markdown("**Klasy DSN co 1 m (macierz pomyłek)**")
        plot(fig_confusion(M), "st_conf")
        st.write(f"Zgodność dokładna {cs['exact']:.1f}%, ±1 klasa {cs['within1']:.1f}%, kappa {cs['kappa']:.2f}, kappa ważona {cs['kappa_w']:.2f}.")
    st.markdown("**Zgodność wzdłuż trasy**")
    wcol = st.columns(2)
    rw = wcol[0].number_input("Okno [m]", 100, 5000, 500, 100, key="st_rw")
    rh = wcol[1].number_input("Krok [m]", 50, 2000, 100, 50, key="st_rh")
    roll = rolling_agreement(v, int(rw), int(rh), "RSP_" + sel, "LCMS_" + sel)
    if len(roll):
        plot(fig_rolling(roll), "st_roll")
    sec = section_agreement(v, 1000.0, origin=A["a0"], a_col="RSP_" + sel, b_col="LCMS_" + sel, tol=A["tol"])
    if len(sec):
        plot(fig_section_bars(sec, "Bias [mm]", "Bias na odcinkach 1 km"), "st_secbias")
        wide(st.dataframe, sec.drop(columns=["Od [m]", "Do [m]"]), hide_index=True)
        A["memo"]["sec_tbl"] = sec


# ---------- MAPA ----------
def tab_map(A):
    st.subheader("Mapa śladu GPS")
    rg = A["rsp"].get("gps") if A["rsp"] else None
    lg = A["lcms"].get("gps") if A["lcms"] else None
    p = A["pairs"]
    fig = go.Figure()
    lat_all, lon_all = [], []
    if rg is not None and len(rg):
        g = rg[(rg["chain_m"] >= A["a0"]) & (rg["chain_m"] <= A["a1"])]
        if len(g):
            fig.add_trace(make_map_trace(g["LAT"].values, g["LON"].values, "lines", line=dict(width=4, color=C_RSP), name="ślad RSP"))
            lat_all += list(g["LAT"]); lon_all += list(g["LON"])
        else:
            st.warning("GPS RSP nie obejmuje wybranego zakresu pikietażu (np. brak współrzędnych w pliku KOLC dla tego odcinka).")
    else:
        st.warning("Brak GPS RSP – pochodzi z pliku RSP (gpx / rekordy 5280 / KOLC); w wczytanych plikach go nie znaleziono.")
    if lg is not None and len(lg):
        tgt0, tgt1 = A["a0"] + A["off0"], A["a1"] + A["off0"]
        g = lg[(lg["chain_m"] >= tgt0) & (lg["chain_m"] <= tgt1)]
        if len(g):
            fig.add_trace(make_map_trace(g["LAT"].values, g["LON"].values, "lines", line=dict(width=3, color=C_LCMS), name="ślad LCMS"))
            lat_all += list(g["LAT"]); lon_all += list(g["LON"])
    if rg is not None and len(rg) > 1 and len(p):
        gg = rg.sort_values("chain_m")
        pp = p[(p["Pikietaz"] >= gg["chain_m"].min()) & (p["Pikietaz"] <= gg["chain_m"].max())].iloc[::10]
        pp = pp.dropna(subset=["RSP_Max"])
        if len(pp):
            lat = np.interp(pp["Pikietaz"], gg["chain_m"], gg["LAT"])
            lon = np.interp(pp["Pikietaz"], gg["chain_m"], gg["LON"])
            cols = [get_dsn_color(get_dsn_class(h)) for h in pp["RSP_Max"]]
            txt = [f"{km_str(c)}<br>RSP {r:.1f} mm<br>LCMS {l:.1f} mm" for c, r, l in zip(pp["Pikietaz"], pp["RSP_Max"], pp["LCMS_Max"])]
            fig.add_trace(make_map_trace(lat, lon, "markers", marker=dict(size=7, color=cols), name="klasa RSP (Max)", text=txt))
    if not lat_all:
        st.info("Brak punktów GPS do wyświetlenia w wybranym zakresie.")
        return
    finish_map(fig, float(np.mean(lat_all)), float(np.mean(lon_all)), map_zoom_for(lat_all, lon_all))
    plot(fig, "map")
    sep = gps_track_separation(rg, lg, A["off_fn"], p["Pikietaz"].values) if (rg is not None and lg is not None) else None
    if sep is not None and len(sep):
        st.write(f"Odległość śladów RSP i LCMS w tych samych (po przesunięciu) pikietażach: mediana **{sep['Dystans [m]'].median():.1f} m**, "
                 f"95. percentyl {sep['Dystans [m]'].quantile(0.95):.1f} m. Duże odległości sugerują błędny offset albo różne pasy ruchu.")
        f = go.Figure(go.Scattergl(x=sep["Pikietaz"], y=sep["Dystans [m]"], mode="lines"))
        f.update_layout(height=280, xaxis_title="Pikietaż [m]", yaxis_title="Odległość śladów [m]", template="plotly_white")
        plot(f, "map_sep")


# ---------- EKSPORT ----------
def tab_export(A):
    st.subheader("Eksport wyników")
    v = A["valid"]
    if len(v) < 3:
        st.info("Brak danych do eksportu.")
        return
    c = get_conclusions(A)
    stx = get_stats(A)
    hm = get_hm(A)
    sec = A["memo"].get("sec_1km")
    st.download_button("Tabela porównawcza co 1 m (CSV)", A["pairs"].to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
                       "porownanie_RSP_LCMS_1m.csv", "text/csv", key="dl_pairs")
    if st.button("Przygotuj raport Excel i HTML", key="ex_build"):
        with st.spinner("Budowanie raportu…"):
            sheets = {"Porównanie 1 m": A["pairs"], "Odcinki DSN": hm, "Statystyki": stats_table(stx, A["tol"]),
                      "Zgodność 1 km": sec, "Przedziały głębokości": depth_bin_table(v["RSP_Max"].values, v["LCMS_Max"].values, tol=A["tol"]),
                      "RPD": build_rpd_table_df(v["RSP_Max"].values, v["LCMS_Max"].values)}
            st.session_state["xlsx_bytes"] = build_excel_bytes(sheets, c)
            figs = [("Bland–Altman (Max)", build_bland_altman_fig(v["RSP_Max"].values, v["LCMS_Max"].values, "KOLEINA MAKSYMALNA")[0]),
                    ("Przebieg wzdłużny (Max)", fig_longitudinal(A["pairs"], [{"col": "RSP_Max", "name": "RSP", "color": C_RSP},
                                                                              {"col": "LCMS_Max", "name": "LCMS", "color": C_LCMS}],
                                                                  "Koleina maksymalna", (A["a0"], A["a1"]))),
                    ("Różnica wzdłuż trasy", fig_diff_along(A["pairs"], tol=A["tol"]))]
            for pre, title in (("L_KOL", "Lewa"), ("P_KOL", "Prawa"), ("MAX_KOL", "Maksymalna")):
                if len(hm):
                    figs.append((f"Ocena odcinkowa – {title}", make_50m_bar_chart(hm, pre, title, A["sec_len"])))
            st.session_state["html_bytes"] = build_html_report(
                "Analiza porównawcza RSP – LCMS", c["markdown"], figs,
                [("Statystyki", stats_table(stx, A["tol"])), ("Zgodność na odcinkach 1 km", sec), ("Ocena odcinkowa DSN", hm)])
    if st.session_state.get("xlsx_bytes"):
        st.download_button("Pobierz raport Excel (.xlsx)", st.session_state["xlsx_bytes"], "analiza_RSP_LCMS.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_xlsx")
    if st.session_state.get("html_bytes"):
        st.download_button("Pobierz raport HTML", st.session_state["html_bytes"], "analiza_RSP_LCMS.html", "text/html", key="dl_html")
    st.download_button("Wnioski (Markdown)", c["markdown"].encode("utf-8"), "wnioski.md", "text/markdown", key="dl_md")
    with st.expander("Podgląd wniosków (Markdown)"):
        st.markdown(c["markdown"])


# ==========================================================================
# 13. APLIKACJA
# ==========================================================================
SECTIONS = ["Wnioski", "Dane i pliki", "Synchronizacja", "Przekrój poprzeczny", "Przebieg wzdłużny",
            "Ocena odcinkowa DSN", "Statystyki zgodności", "Mapa", "Eksport"]


def file_registry_ui(refs):
    """Tabela plików (typ rozpoznany po zawartości, edytowalny) w formularzu."""
    reg = st.session_state.setdefault("reg", {})
    for r in refs:
        if r.ident not in reg:
            typ, note = classify_cached(r)
            reg[r.ident] = {"use": typ != TYPE_SKIP, "typ": typ, "note": note, "start": None, "step": 1.0, "start_for": None}
    # domyślne starty (wymagają wczytania) – tylko dla plików używanych
    for r in refs:
        cfg = reg[r.ident]
        if cfg["typ"] != TYPE_SKIP and (cfg["start"] is None or cfg.get("start_for") != cfg["typ"]):
            try:
                src = cached_parse(r.ident, cfg["typ"], r)
                gs = source_default_start(src)
                lo, _ = source_extent(src, gs, cfg["step"])
                cfg["start"] = float(lo) if not np.isnan(lo) else 0.0
                cfg["start_for"] = cfg["typ"]
            except Exception as exc:  # noqa: BLE001
                cfg["note"] = f"błąd wczytania: {exc}"
                cfg["use"] = False
    df = pd.DataFrame([{
        "Użyj": bool(reg[r.ident]["use"]), "Plik": r.name, "MB": float(fmt_mb(r.size)),
        "Typ": TYPE_LABELS.get(reg[r.ident]["typ"], reg[r.ident]["typ"]),
        "Pikietaż początku [m]": reg[r.ident]["start"], "Krok profili [m]": reg[r.ident]["step"],
        "Uwagi": reg[r.ident]["note"]} for r in refs])
    with st.form("files_form"):
        ed = wide(
            st.data_editor,
            df, hide_index=True, key="files_editor",
            disabled=["Plik", "MB", "Uwagi"],
            column_config={
                "Typ": st.column_config.SelectboxColumn("Typ", options=list(LABEL_TO_TYPE.keys()), required=True),
                "Pikietaż początku [m]": st.column_config.NumberColumn(format="%.1f",
                                                                      help="Pikietaż pierwszej próbki pliku (w metrach, np. 114310 = km 114+310)"),
                "Krok profili [m]": st.column_config.NumberColumn(format="%.3f", min_value=0.01)})
        ok = st.form_submit_button("Zastosuj zmiany w plikach")
    if ok or "files_applied" not in st.session_state:
        for r, (_, row) in zip(refs, ed.iterrows()):
            cfg = reg[r.ident]
            new_t = LABEL_TO_TYPE.get(row["Typ"], cfg["typ"])
            if new_t != cfg["typ"]:
                cfg["typ"], cfg["start"], cfg["start_for"] = new_t, None, None
            else:
                if row["Pikietaż początku [m]"] is not None and not pd.isna(row["Pikietaż początku [m]"]):
                    cfg["start"] = float(row["Pikietaż początku [m]"])
            cfg["use"] = bool(row["Użyj"])
            cfg["step"] = float(row["Krok profili [m]"]) if not pd.isna(row["Krok profili [m]"]) else 1.0
        st.session_state["files_applied"] = True
        if ok:
            st.rerun()
    return reg


def sidebar_settings():
    sb = st.sidebar
    sb.header("Ustawienia analizy")
    sb.number_input("Tolerancja zgodności ± [mm]", 0.5, 10.0, 2.0, 0.5, key="tol")
    sb.selectbox("Długość odcinka oceny DSN [m]", [25, 50, 100, 200], index=1, key="sec_len")
    sb.slider("Minimalne pokrycie odcinka [%]", 0, 100, 0, 10, key="min_cover")
    sb.selectbox("Źródło koleiny RSP", ["auto", "raw", "kolc"], key="rsp_source",
                 format_func=lambda k: {"auto": "auto (lasery, luki z KOLC)", "raw": "tylko lasery (Wirtualna Łata)", "kolc": "tylko KOLC"}[k])
    sb.selectbox("KOLC: wartość", ["avg", "max"], key="kolc_param",
                 format_func=lambda k: {"avg": "uśredniona (L/P)", "max": "maksymalna (ML/MP)"}[k])
    sb.selectbox("Źródło koleiny LCMS", ["file", "calc", "both"], key="lcms_source",
                 format_func=lambda k: {"file": "z pliku", "calc": "wyliczona z profili (Wirtualna Łata)", "both": "oba (plik + łata)"}[k])
    sb.number_input("Okno wygładzania profilu LCMS [pkt]", 1, 101, 31, 2, key="win")
    sb.checkbox("Zamień lewą/prawą LCMS", False, key="swap_lr")
    sb.checkbox("Odwróć kierunek LCMS", False, key="reverse_lcms")
    sb.selectbox("Pikietaże RSP i LCMS", ["auto", "tak", "nie"], key="rel_mode",
                 format_func=lambda k: {"auto": "auto (względne gdy się nie pokrywają)", "tak": "względne (od 0 m)", "nie": "bezwzględne"}[k])
    sb.text_input("Kamera PIC (GPS LCMS)", "auto", key="camera")
    sb.divider()
    sb.subheader("Przesunięcie LCMS względem RSP")
    st.session_state.setdefault("off_global", 0.0)
    st.session_state.setdefault("off_drift", 0.0)
    sb.number_input("Offset globalny [m]", step=1.0, key="off_global")
    sb.number_input("Dryf [m/km]", step=0.05, key="off_drift", format="%.2f")


def main():
    st.set_page_config(page_title="RSP vs LCMS – analiza kolein", layout="wide", page_icon="🛣️")
    pending_apply()
    st.title("🛣️ Analiza porównawcza kolein: RSP vs LCMS")
    st.caption(f"DSN 2019 · wersja {APP_VERSION} · wiele plików, synchronizacja pikietażu, statystyki zgodności, wnioski, eksport")

    with st.sidebar:
        st.header("Pliki")
        up = st.file_uploader("Wgraj pliki (RSP, KOLC, LCMS, GPX, PIC XML …)", accept_multiple_files=True, key="uploader")
        use_folder = st.checkbox("Czytaj pliki z folderu na dysku (duże pliki)", False, key="use_folder")
        folder = st.text_input("Ścieżka folderu", "", key="folder") if use_folder else ""
    refs = collect_refs(up, folder)
    if not refs:
        st.info("Wgraj pliki w panelu bocznym. Obsługiwane: surowe RSP (xlsx/csv/rsp), wyniki KOLC (XML), profile LCMS (csv/xlsx), "
                "ślady GPS (gpx, PIC XML). Typ pliku jest rozpoznawany po zawartości; dla wielu plików z kolejnych odcinków "
                "ustaw pikietaż początku każdego pliku.")
        return
    sidebar_settings()

    with st.expander(f"Pliki ({len(refs)})", expanded=("sections_seen" not in st.session_state)):
        reg = file_registry_ui(refs)
    entries, errors = parse_all(refs, reg)
    for e in errors:
        st.error("Błąd wczytywania – " + e)
    entries = [e for e in entries if e["use"]]
    if not entries:
        st.warning("Brak aktywnych plików.")
        return

    with st.spinner("Scalanie plików…"):
        data = cached_assemble(entries_sig(entries), entries, st.session_state.get("camera", "auto"),
                               st.session_state.get("reverse_lcms", False), st.session_state.get("rel_mode", "auto"))
    rsp, lcms, notes = data["rsp"], data["lcms"], list(data["notes"])
    if rsp is None and lcms is None:
        st.warning("Nie znaleziono danych koleiny (RSP/LCMS) wśród aktywnych plików.")
        return

    ss = st.session_state
    lcms_source = ss.get("lcms_source", "file")
    lcms_calc = None
    rsp_s = lcms_s = None
    if lcms is not None:
        need_calc = lcms_source in ("calc", "both") or bool(np.all(np.isnan(lcms["rut_l"])))
        if need_calc:
            with st.spinner("Liczenie kolein z profili LCMS (Wirtualna Łata)…"):
                lcms_calc = cached_lcms_calc(lcms["uid"] + f"|{ss.get('win', 31)}", int(ss.get("win", 31)), lcms)
        lcms_s = lcms_series(lcms, lcms_calc)
    rsp_x = None
    if rsp is not None:
        raw_ruts = None
        if rsp.get("raw") is not None:
            rsp_x = rsp_geometry_ui(rsp)
            with st.spinner("Liczenie kolein RSP z laserów…"):
                raw_ruts = cached_rsp_raw_ruts(rsp["uid"], tuple(rsp_x), rsp)
        rsp_s = rsp_series(rsp, raw_ruts, ss.get("kolc_param", "avg"))

    # domyślny zakres analizy
    uid = (rsp["uid"] if rsp else "-") + "|" + (lcms["uid"] if lcms else "-") + f"|{data['rel']}"
    off0 = float(ss.get("off_global", 0.0))
    if ss.get("range_uid") != uid:
        if rsp is not None and lcms is not None:
            lo = max(rsp["c_min"], lcms["c_min"] - off0)
            hi = min(rsp["c_max"], lcms["c_max"] - off0)
            if hi <= lo:
                lo, hi = rsp["c_min"], rsp["c_max"]
        else:
            o = rsp or lcms
            lo, hi = o["c_min"], o["c_max"]
        ss["a0"], ss["a1"] = float(np.ceil(lo)), float(np.floor(hi))
        ss["range_uid"] = uid
    with st.sidebar:
        st.divider()
        st.subheader("Zakres analizy")
        st.number_input("Od pikietażu [m]", step=100.0, key="a0")
        st.number_input("Do pikietażu [m]", step=100.0, key="a1")
        st.caption(f"{km_str(ss['a0'])} → {km_str(ss['a1'])}  (~{(ss['a1'] - ss['a0']) / 1000:.2f} km)")
    a0, a1 = float(ss["a0"]), float(ss["a1"])
    if a1 <= a0:
        st.error("Zakres analizy jest pusty (Do ≤ Od).")
        return

    if rsp_s is None or lcms_s is None:
        only = "RSP" if rsp_s is not None else "LCMS"
        st.warning(f"Wczytano tylko dane {only}. Porównanie wymaga obu systemów – pokazuję dostępne widoki.")
        tab_data({"entries": entries, "notes": notes, "rsp": rsp, "lcms": lcms, "off0": 0.0, "pairs": pd.DataFrame()})
        return

    off_fn = current_offset_fn(a0)
    pairs, rdesc = build_pairs(rsp_s, lcms_s, ss.get("rsp_source", "auto"), lcms_source, a0, a1, off_fn,
                               ss.get("kolc_param", "avg"), bool(ss.get("swap_lr", False)))
    need = ["RSP_L", "RSP_R", "LCMS_L", "LCMS_R"]
    valid = pairs.dropna(subset=need).reset_index(drop=True)
    ldesc = {"file": "koleiny z pliku", "calc": "koleiny wyliczone z profili", "both": "plik + wyliczone"}[lcms_source]
    for e in entries:
        for n in (e["src"].get("notes") or []):
            if ("brakiem" in n or "połówk" in n or "luki w numeracji" in n) and e["type"] == TYPE_LCMS_PROFILE:
                notes.append(f"{e['name']}: {n}")
    A = {"entries": entries, "rsp": rsp, "lcms": lcms, "rsp_s": rsp_s, "lcms_s": lcms_s, "pairs": pairs, "valid": valid,
         "rdesc": rdesc, "ldesc": ldesc, "a0": a0, "a1": a1, "off_fn": off_fn, "off0": off0, "notes": notes,
         "tol": float(ss.get("tol", 2.0)), "sec_len": float(ss.get("sec_len", 50)), "min_cover": float(ss.get("min_cover", 0)),
         "rsp_source": ss.get("rsp_source", "auto"), "lcms_source": lcms_source, "swap_lr": bool(ss.get("swap_lr", False)),
         "win": int(ss.get("win", 31)), "rsp_x": rsp_x, "memo": {}}
    st.session_state["sections_seen"] = True
    if len(valid) == 0:
        st.warning("Brak wspólnych punktów RSP–LCMS. Zwykle oznacza to złe pikietaże początkowe plików albo offset – "
                   "sprawdź zakładkę „Dane i pliki” (zakresy) i „Synchronizacja”.")

    sec = st.radio("Widok", SECTIONS, horizontal=True, key="section", label_visibility="collapsed")
    fn = {"Wnioski": tab_conclusions, "Dane i pliki": tab_data, "Synchronizacja": tab_sync, "Przekrój poprzeczny": tab_cross,
          "Przebieg wzdłużny": tab_long, "Ocena odcinkowa DSN": tab_dsn, "Statystyki zgodności": tab_stats,
          "Mapa": tab_map, "Eksport": tab_export}[sec]
    fn(A)


if __name__ == "__main__":
    main()
